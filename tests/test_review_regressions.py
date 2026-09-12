"""Focused independent regressions found in final contract review."""

import json
from dataclasses import replace
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from gerador_cnab_nox.errors import ValidationError
from gerador_cnab_nox.matching import prepare_batch
from gerador_cnab_nox.models import CESSAO, PfmiData, PfmiInput
from gerador_cnab_nox.service import generate_cnab, prepare_workbook
from gerador_cnab_nox.workbook import read_intermediate, write_intermediate

from .test_matching import _credit, _due, _payment
from .test_revision_e2e import edit


@pytest.mark.parametrize(
    "changes,message",
    [
        ({"DOC_CEDENTE": "52998224725", "TIPO_PESSOA_CEDENTE": 1}, "CNPJ válido"),
        ({"NOME_CEDENTE": "CONEXCRED/INTERMEDIACAO"}, "deve ser somente"),
    ],
)
def test_t13_approval_does_not_replace_cnpj_or_exact_principal(
    source_files, tmp_path, changes, message
):
    path = tmp_path / "cessao.xlsx"
    prepare_workbook(
        [PfmiInput(source_files[0], CESSAO, "0")],
        *source_files[1:],
        liquidation_date="03/09/2026",
        first_sequence=1,
        output_path=path,
    )
    edit(
        path,
        {
            "DOC_CEDENTE": "11222333000181",
            "TIPO_PESSOA_CEDENTE": 2,
            "DT_EMISSAO_TITULO": date(2026, 9, 2),
            "APROVADO": "SIM",
            **changes,
        },
    )
    with pytest.raises(ValidationError, match=message):
        generate_cnab(path, output_path=tmp_path / "blocked.txt")


def test_t21_aggregate_above_single_title_limit_is_valid(source_files, tmp_path):
    pfmi, analytic, due = source_files
    book = load_workbook(pfmi)
    book.active.cell(3, 1).value = "119999999940.02"
    book.active.cell(3, 9).value = "119999999940.02"
    book.save(pfmi)
    book.close()
    book = load_workbook(analytic)
    sh = book.active
    sh.cell(2, 4).value = "60000000000.01"
    sh.cell(2, 5).value = "60000000000.01"
    sh.append([c.value for c in sh[2]])
    sh.cell(3, 7).value = "SECOND-CREDIT"
    book.save(analytic)
    book.close()
    path = tmp_path / "large.xlsx"
    result = prepare_workbook(
        pfmi, analytic, due, liquidation_date="03/09/2026", first_sequence=1, output_path=path
    )
    # Security fix: two principal credits with the same document/name and an
    # exact aggregate sum are no longer auto-selected; confirm both manually.
    assert result.selected_count == 0
    edit(path, {"INCLUIR_CNAB": "SIM"}, row=2)
    edit(path, {"INCLUIR_CNAB": "SIM"}, row=3)
    generated = generate_cnab(path, output_path=tmp_path / "large.txt")
    assert generated.detail_count == 2
    assert generated.total_present == Decimal("120000000000.02")


def test_t18_large_original_cents_survive_excel_round_trip(tmp_path):
    amount = Decimal("100000000000000.01")
    batch = prepare_batch(
        PfmiData(
            payments=[
                replace(_payment("B1", "100"), operation_value=amount, beneficiary_value=amount)
            ]
        ),
        [_credit("A", "100")],
        [_due()],
        liquidation_date=date(2026, 9, 3),
        first_sequence=1,
    )
    path = tmp_path / "exact.xlsx"
    write_intermediate(batch, path)
    loaded = read_intermediate(path)
    assert loaded.structural_issues == []
    assert Decimal(loaded.rows[0].originals["TOTAL_PFMI_GRUPO"]) == amount
    assert Decimal(loaded.rows[0].values["TOTAL_PFMI_GRUPO"]) == amount


def test_t13_conex_source_from_other_group_is_visible_and_conflicts_are_local():
    p = replace(
        _payment("B1", "100"),
        modality=CESSAO,
        commission_text="0",
        beneficiary_name="OUTRO",
        beneficiary_document="52998224725",
    )
    other = replace(
        p,
        block_id="B2",
        source_row=4,
        source_sheet="ABA SINTETICA",
        beneficiary_name="CONEXCRED",
        beneficiary_document="11222333000181",
    )
    second_file = replace(
        other, file_id="F0002", file_order=2, beneficiary_document="11444777000161"
    )

    def prepare(payments):
        return prepare_batch(
            PfmiData(payments=payments),
            [_credit("A", "100")],
            [_due()],
            liquidation_date=date(2026, 9, 3),
            first_sequence=1,
        )

    batch = prepare([p, other, second_file])
    r = batch.rows[0].values
    assert r["DOC_CEDENTE"] == "11222333000181"
    assert r["ESTADO_CNPJ_CONEXCRED"] == "UNICO_VALIDO"
    assert json.loads(r["ORIGENS_CNPJ_CONEXCRED"]) == [
        {"aba": "ABA SINTETICA", "linha": 4, "documento": "11222333000181"}
    ]
    assert batch.rows[-1].values["DOC_CEDENTE"] == "11444777000161"
    conflict = replace(other, block_id="B3", beneficiary_document="11444777000161")
    batch = prepare([p, other, conflict, second_file])
    assert all(
        r.values["ESTADO_CNPJ_CONEXCRED"] == "CONFLITANTE" and not r.values["DOC_CEDENTE"]
        for r in batch.rows
        if r.values["ID_ARQUIVO"] == "F0001"
    )
    assert batch.rows[-1].values["DOC_CEDENTE"] == "11444777000161"


@pytest.mark.parametrize(
    "name,approved,allowed",
    [
        ("FULANO DE TAL COMPLETO", "SIM", True),
        ("FULANO DE TAL COMPLETO", "", False),
        ("BELTRANO", "SIM", False),
        ("FULANO DE TAL CREDOR SINTETICO", "SIM", False),
        ("FULANO DE TAL (CREDOR SINTETICO)", "SIM", False),
    ],
)
def test_t14_final_name_correction_keeps_lawyer_identity(tmp_path, name, approved, allowed):
    original = "FULANO DE TAL (CREDOR SINTETICO)"
    p = replace(_payment("B1", "100"), failure="FALENCIA GAMA", cedent_name=original)
    c = replace(_credit("A", "100"), campaign=p.failure, cedent_name=original)
    d = replace(_due(), debtor_name=p.failure)
    rules = {"FALENCIA GAMA": ("FULANO DE TAL", "BELTRANO")}
    batch = prepare_batch(
        PfmiData(payments=[p]),
        [c],
        [d],
        liquidation_date=date(2026, 9, 3),
        first_sequence=1,
        lawyer_rules=rules,
    )
    path = tmp_path / "lawyer.xlsx"
    write_intermediate(batch, path)
    edit(path, {"NOME_CEDENTE": name, "APROVADO": approved, "INCLUIR_CNAB": "SIM"})
    if allowed:
        assert (
            generate_cnab(path, output_path=tmp_path / "ok.txt", lawyer_rules=rules).detail_count
            == 1
        )
    else:
        with pytest.raises(ValidationError):
            generate_cnab(path, output_path=tmp_path / "blocked.txt", lawyer_rules=rules)


def test_t15_normal_issue_and_per_file_summary(source_files, tmp_path):
    path = tmp_path / "normal.xlsx"
    prepare_workbook(
        *source_files, liquidation_date="03/09/2026", first_sequence=1, output_path=path
    )
    loaded = read_intermediate(path)
    row = loaded.rows[0]
    assert row.values["DT_EMISSAO_TITULO"] == row.originals["ASSINATURA_ORIGINAL"]
    assert row.values["DT_EMISSAO_TITULO"].date() == date(2026, 9, 1)
    book = load_workbook(path)
    summary = dict(book["RESUMO"].values)
    assert {
        key: summary["F0001_" + key]
        for key in (
            "PAGAMENTOS",
            "GRUPOS",
            "CANDIDATOS",
            "SELECIONADOS_PREVISTOS",
            "TOTAL_PFMI_VALIDO",
            "TOTAL_NOMINAL_PREVISTO",
            "TOTAL_PRESENTE_PREVISTO",
        )
    } == dict(
        PAGAMENTOS=2,
        GRUPOS=1,
        CANDIDATOS=1,
        SELECIONADOS_PREVISTOS=1,
        TOTAL_PFMI_VALIDO=100,
        TOTAL_NOMINAL_PREVISTO=150,
        TOTAL_PRESENTE_PREVISTO=100,
    )
    book.close()


@pytest.mark.parametrize(
    "sheet,cell",
    [
        ("RESUMO", "B3"),
        ("RESUMO", "B4"),
        ("PENDENCIAS_PFMI", "A2"),
        ("ORIGINAIS_CONTROLE", "B2"),
        ("MANIFESTO", "B3"),
        ("CREDITOS", "P2"),
    ],
)
def test_t18_formulas_block_all_relevant_surfaces(source_files, tmp_path, sheet, cell):
    path = tmp_path / "formula.xlsx"
    prepare_workbook(
        *source_files, liquidation_date="03/09/2026", first_sequence=1, output_path=path
    )
    book = load_workbook(path)
    book[sheet][cell] = "=1+1"
    if sheet == "CREDITOS":
        headers = {c.value: c.column for c in book[sheet][1]}
        book[sheet].cell(2, headers["INCLUIR_CNAB"]).value = "NAO"
    book.save(path)
    book.close()
    with pytest.raises(ValidationError):
        generate_cnab(path, output_path=tmp_path / "blocked.txt")
