from __future__ import annotations

import contextlib
import gc
import os
import threading
import time
from decimal import Decimal
from pathlib import Path
from unittest.mock import Mock

import pytest

import gerador_cnab_nox.gui as gui
from gerador_cnab_nox.errors import ValidationError
from gerador_cnab_nox.gui_messages import explain_issue, explain_pending, money
from gerador_cnab_nox.models import GenerationResult, PreparationResult
from tests.test_gui import (
    _choose_modality,
    _select_pfmi,
    _set_commission,
    _single_normal_sources,
    _wait_for_idle,
)
from tests.test_gui import gui_application as gui_application


def _setup_sources(root, app, monkeypatch, directory):
    pfmis, analytic, due = _single_normal_sources(directory)
    monkeypatch.setattr(gui.filedialog, "askopenfilenames", lambda **_: tuple(map(str, pfmis)))
    app.add_pfmi_button.invoke()
    app.analytic.set(str(analytic))
    app.due.set(str(due))
    app.liquidation.set("05/09/2026")
    app.sequence.set("1001")
    root.update()
    return pfmis, analytic, due


def _wait_for_window(root, width, height):
    """X11 resize/map acknowledgements arrive after update() may have returned."""
    deadline = time.monotonic() + 4
    stable_since = None
    while time.monotonic() < deadline:
        root.update()
        if root.winfo_ismapped() and (root.winfo_width(), root.winfo_height()) == (width, height):
            stable_since = stable_since or time.monotonic()
            if time.monotonic() - stable_since >= 0.12:
                return
        else:
            stable_since = None
        time.sleep(0.01)
    raise AssertionError(f"Janela não estabilizou em {width}x{height}: {root.geometry()}")


def test_navigation_paths_staleness_and_separate_file_results(
    tmp_path, monkeypatch, gui_application
):
    root, app, _ = gui_application
    pfmis, _, _ = _setup_sources(root, app, monkeypatch, tmp_path)
    _select_pfmi(root, app, 0)
    assert app.pfmi_tree.item(app.pfmi_tree.get_children()[0], "values")[1] == pfmis[0].name
    assert app.selected_path.get() == str(pfmis[0])
    assert str(app.selected_path_entry.cget("state")) == "readonly"
    assert str(app.remove_pfmi_button.cget("state")) == "normal"
    assert str(app.move_up_button.cget("state")) == "disabled"

    app.continue_button.invoke()
    assert app.current_step == 2
    excel = tmp_path / "preparado.xlsx"
    monkeypatch.setattr(gui.filedialog, "asksaveasfilename", lambda **_: str(excel))
    app.prepare_button.invoke()
    _wait_for_idle(root, app)
    assert app.review_result.kind == "success"
    assert "validado" not in app.review_result.title.cget("text")
    assert app.prepared_path.get() == app.intermediate.get() == str(excel)

    app.back_config_button.invoke()
    _choose_modality(root, app, gui.CESSAO_LABEL)
    _set_commission(root, app, "0")
    assert "prepare um novo Excel" in app.configuration_notice.get()
    app.direct_excel_button.invoke()
    assert app.current_step == 3
    assert app.intermediate.get() == str(excel)
    app.generation_result.show("Resultado anterior", kind="success")
    app.generated_path.set(str(tmp_path / "anterior.txt"))
    app.intermediate.set(str(tmp_path / "outra cópia.xlsx"))
    assert app.generated_path.get() == ""
    assert app.generation_result.kind == "info"
    assert app.prepared_path.get() == str(excel)
    app.show_step(1)
    assert app.pfmi_commission.get() == "0"


def test_new_batch_button_clears_pfmis_sources_dates_and_results(
    tmp_path, monkeypatch, gui_application
):
    root, app, _ = gui_application
    _setup_sources(root, app, monkeypatch, tmp_path)
    assert app.pfmi_tree.get_children()
    excel = tmp_path / "preparado.xlsx"
    monkeypatch.setattr(gui.filedialog, "asksaveasfilename", lambda **_: str(excel))
    app.prepare_button.invoke()
    _wait_for_idle(root, app)
    assert app.prepared_path.get() == str(excel)

    app.new_batch_button.invoke()

    assert app.current_step == 1
    assert not app.pfmi_tree.get_children()
    assert not app._pfmi_rows
    assert app.pfmi_count.get() == "0 PFMIs"
    assert app.analytic.get() == ""
    assert app.due.get() == ""
    assert app.liquidation.get() == ""
    assert app.sequence.get() == ""
    assert app.prepared_path.get() == ""
    assert app.intermediate.get() == ""
    assert app._prepared_signature is None
    assert app.review_result.title.cget("text") == "Pronto para preparar"
    assert str(app.prepare_button.cget("state")) == "normal"
    assert not app.open_excel_button.winfo_manager()
    assert not app.continue_generate_button.winfo_manager()
    assert not app.prepared_files.winfo_manager()
    assert app.summary_label.winfo_manager() == "pack"

    # A fresh selection right after reset must not resurrect anything cleared.
    novo_lote_dir = tmp_path / "novo_lote"
    novo_lote_dir.mkdir()
    next_pfmis, next_analytic, next_due = _single_normal_sources(novo_lote_dir)
    monkeypatch.setattr(
        gui.filedialog, "askopenfilenames", lambda **_: tuple(map(str, next_pfmis))
    )
    app.add_pfmi_button.invoke()
    assert len(app.pfmi_tree.get_children()) == len(next_pfmis)


def test_new_batch_after_generate_button_resets_and_returns_to_step_1(
    tmp_path, monkeypatch, gui_application
):
    root, app, _ = gui_application
    _setup_sources(root, app, monkeypatch, tmp_path)
    excel = tmp_path / "preparado.xlsx"
    txt = tmp_path / "gerado.txt"
    monkeypatch.setattr(gui.filedialog, "asksaveasfilename", lambda **_: str(excel))
    app.prepare_button.invoke()
    _wait_for_idle(root, app)
    monkeypatch.setattr(gui.filedialog, "asksaveasfilename", lambda **_: str(txt))
    app.intermediate.set(str(excel))
    app.generate_button.invoke()
    _wait_for_idle(root, app)
    assert app.generated_path.get() == str(txt)

    app.new_batch_after_generate_button.invoke()

    assert app.current_step == 1
    assert not app._pfmi_rows
    assert app.generated_path.get() == ""
    assert app.prepared_path.get() == ""
    assert app.intermediate.get() == ""


def test_add_pfmis_replaces_by_default_and_accumulates_explicitly(
    tmp_path, monkeypatch, gui_application
):
    root, app, _ = gui_application
    primeiro_dir, segundo_dir = tmp_path / "primeiro", tmp_path / "segundo"
    primeiro_dir.mkdir()
    segundo_dir.mkdir()
    first_pfmis, _, _ = _single_normal_sources(primeiro_dir)
    second_pfmis, _, _ = _single_normal_sources(segundo_dir)

    monkeypatch.setattr(
        gui.filedialog, "askopenfilenames", lambda **_: tuple(map(str, first_pfmis))
    )
    app.add_pfmi_button.invoke()
    assert len(app.pfmi_tree.get_children()) == len(first_pfmis)

    # Default action (used again) replaces, so files from an unrelated
    # earlier batch (e.g. a different liquidation date) never accumulate
    # silently into the current one.
    monkeypatch.setattr(
        gui.filedialog, "askopenfilenames", lambda **_: tuple(map(str, second_pfmis))
    )
    app.add_pfmi_button.invoke()
    assert len(app.pfmi_tree.get_children()) == len(second_pfmis)
    assert {row.path for row in app._pfmi_rows.values()} == {str(p) for p in second_pfmis}

    # The explicit "Adicionar mais PFMIs" action is the only way to combine
    # files into the same batch (e.g. Normal + Cessão da Cessão).
    monkeypatch.setattr(
        gui.filedialog, "askopenfilenames", lambda **_: tuple(map(str, first_pfmis))
    )
    app.add_more_pfmi_button.invoke()
    assert len(app.pfmi_tree.get_children()) == len(second_pfmis) + len(first_pfmis)
    assert {row.path for row in app._pfmi_rows.values()} == {
        str(p) for p in second_pfmis
    } | {str(p) for p in first_pfmis}


def test_busy_job_keeps_event_loop_and_snapshots_on_main_thread(
    tmp_path, monkeypatch, gui_application
):
    root, app, _ = gui_application
    _setup_sources(root, app, monkeypatch, tmp_path)
    destination = tmp_path / "capturado.xlsx"
    monkeypatch.setattr(gui.filedialog, "asksaveasfilename", lambda **_: str(destination))
    entered, release = threading.Event(), threading.Event()
    calls = []
    main_ident = threading.get_ident()
    real_prepare = gui.prepare_workbook

    def service(pfmis, analytic, due, **kwargs):
        calls.append((threading.get_ident(), pfmis, analytic, due, kwargs))
        entered.set()
        assert release.wait(5)
        return real_prepare(pfmis, analytic, due, **kwargs)

    monkeypatch.setattr(gui, "prepare_workbook", service)
    app.prepare_button.invoke()
    assert entered.wait(2)
    assert app.is_busy
    assert all(str(w.cget("state")) == "disabled" for w, _ in app._controls)
    assert app.pfmi_tree.instate(["disabled"])
    activity = []
    root.after(1, lambda: activity.append("event loop ativo"))
    try:
        for _ in range(5):
            root.update()
            time.sleep(0.005)
        assert activity
        app.prepare_button.invoke()
        app._prepare()
        app._generate()
        app.show_step(1)
        app._request_close()
        assert root.winfo_exists()
        assert app.current_step == 2
        assert "Aguarde a conclusão" in app.activity.get()
        assert len(calls) == 1
        assert calls[0][0] != main_ident
        assert calls[0][-1]["first_sequence"] == "1001"
        assert calls[0][-1]["output_path"] == str(destination)
    finally:
        release.set()
        _wait_for_idle(root, app)
    assert app.review_result.kind == "success"
    assert str(app.prepare_button.cget("state")) == "normal"
    assert str(app.modality_combo.cget("state")) == "readonly"
    assert str(app.commission_entry.cget("state")) == "disabled"
    assert str(app.move_up_button.cget("state")) == "disabled"
    assert app._poll_id is None


def test_preparation_composition_summary_completes_on_tk_thread(
    tmp_path, monkeypatch, gui_application,
):
    from openpyxl import load_workbook

    root, app, _ = gui_application
    root.geometry("1040x700")
    root.deiconify()
    pfmis, _, _ = _setup_sources(root, app, monkeypatch, tmp_path)
    book = load_workbook(pfmis[0])
    sheet = book.active
    sheet.cell(3, 1, 60)
    sheet.cell(3, 9, 60)
    row = [c.value for c in sheet[3]]
    row[0] = row[8] = 40
    row[2] = "REPRESENTANTE (EMPRESA NORMAL)"
    row[3] = "REPRESENTANTE"
    sheet.append(row)
    book.save(pfmis[0])
    book.close()
    target = tmp_path / "composition.xlsx"
    monkeypatch.setattr(gui.filedialog, "asksaveasfilename", lambda **_: str(target))
    main_thread = threading.get_ident()
    read_threads, completion_threads, after_calls, callback_errors, redraws = [], [], [], [], []
    real_read, real_completed = gui.read_intermediate, app._prepared
    real_after, real_redraw = root.after, root.update_idletasks

    def read(path):
        read_threads.append(threading.get_ident())
        return real_read(path)

    def completed(*args):
        completion_threads.append(threading.get_ident())
        return real_completed(*args)

    def after(delay, callback=None, *args):
        after_calls.append((threading.get_ident(), delay, getattr(callback, "__name__", "")))
        return real_after(delay, callback, *args)

    def redraw():
        redraws.append(threading.get_ident())
        return real_redraw()

    monkeypatch.setattr(gui, "read_intermediate", read)
    monkeypatch.setattr(app, "_prepared", completed)
    monkeypatch.setattr(root, "after", after)
    monkeypatch.setattr(root, "update_idletasks", redraw)
    monkeypatch.setattr(
        root, "report_callback_exception", lambda *args: callback_errors.append(args)
    )
    app.prepare_button.invoke()
    _wait_for_idle(root, app)
    assert target.exists()
    assert not callback_errors
    assert read_threads and all(t != main_thread for t in read_threads)
    assert completion_threads == [main_thread]
    assert all(t == main_thread for t, _, _ in after_calls)
    assert (main_thread, 0, "_finish_job") in after_calls
    assert redraws and all(t == main_thread for t in redraws)
    assert app._poll_id is None
    assert app.review_result.title.cget("text") == "Excel preparado"
    assert app.prepared_path.get() == app.intermediate.get() == str(target)
    assert "Preparando o Excel" not in app.status.get()
    assert not app.progress.winfo_ismapped()
    assert not app.activity_label.winfo_ismapped()
    assert len(app.review_result.metrics_frame.winfo_children()) == 6
    assert app.review_result.metric_summary == (
        "operações: 2 · prontas: 0 · pendentes: 2 · composições encontradas: 1 · "
        "títulos envolvidos: 2 · composições aguardando aprovação: 1"
    )
    text = app.review_result.issue_text.get("1.0", "end")
    assert "── Preenchimentos comprovados (automáticos) ──" in text
    assert "── Sugestões aguardando confirmação ──" in text
    assert "── Pendências sem solução automática ──" in text
    assert "PRO_RATA_VALOR_PRESENTE_COMPOSICAO" in text
    for widget in (app.open_excel_button, app.continue_generate_button, app.prepared_files):
        assert widget.winfo_ismapped()
    for button in (app.prepare_button, app.prepare_again_button,
                   app.open_excel_button, app.continue_generate_button):
        assert str(button.cget("state")) == "normal"
    loaded = real_read(target)
    assert all(r.values["COMPOSICAO_APROVADA"] in (None, "") for r in loaded.rows)


def test_prepared_summary_survives_blank_pendencias_after_nominal_suggestion(
    gui_application,
):
    """Regression for the V1-7 pro-rata suggestion: once it clears a
    composition row's only pendency, PENDENCIAS round-trips through openpyxl
    as None (not "") for a still-unapproved row - explain_pending() must not
    raise TypeError, and the full summary panel must render regardless."""
    root, app, _ = gui_application

    class _FakeRow:
        def __init__(self, values):
            self.values = values

    rows = [
        _FakeRow(
            {
                "ID_GRUPO": "F0001-B0001-G0001",
                "NOME_CEDENTE_PFMI": "PRINCIPAL COMPOSICAO",
                "STATUS": "OK",
                "PENDENCIAS": None,  # blank cell round-tripped from Excel
                "INCLUIR_CNAB": "NAO",
                "COMPOSICAO_ID": "COMP1",
                "COMPOSICAO_APROVADA": None,
                "PREENCHIMENTOS_AUTOMATICOS": (
                    '[{"campo": "VL_NOMINAL", "valor_sugerido": 123.45, '
                    '"regra": "PRO_RATA_VALOR_PRESENTE_COMPOSICAO", '
                    '"exige_aprovacao_humana": true}]'
                ),
            }
        ),
        _FakeRow(
            {
                "ID_GRUPO": "F0001-B0001-G0002",
                "NOME_CEDENTE_PFMI": "SATELITE COMPOSICAO",
                "STATUS": "DIVERGENCIA_NOME",
                "PENDENCIAS": "DIVERGENCIA_NOME",
                "INCLUIR_CNAB": "NAO",
                "COMPOSICAO_ID": "COMP1",
                "COMPOSICAO_APROVADA": None,
                "PREENCHIMENTOS_AUTOMATICOS": "",
            }
        ),
    ]
    result = PreparationResult(
        output_path=Path("fake.xlsx"),
        row_count=2,
        selected_count=0,
        warning_count=0,
        warnings=(),
        payment_count=2,
        group_count=2,
        related_suggestions=(),
        compositions=(
            {
                "composicao_id": "COMP1",
                "estado": "PROPOSTA",
                "participantes": "PRINCIPAL COMPOSICAO + SATELITE COMPOSICAO",
                "motivo": "soma exata",
                "total_pfmi": Decimal("123.45"),
                "total_analitico": Decimal("123.45"),
                "diferenca": Decimal("0.00"),
                "nominal_analitico": Decimal("123.45"),
                "component_ids": ("F0001-B0001-G0001", "F0001-B0001-G0002"),
            },
        ),
    )
    review = gui._PreparedReview(result=result, rows=rows)
    app._job_step = 2
    app._job_completed = lambda value: app._prepared(value, 1, "sig")
    # Called synchronously (no root.after in between), so a bug here raises
    # straight into the test instead of only reaching Tk's error hook.
    app._finish_job(True, review)
    assert app.review_result.title.cget("text") == "Excel preparado"
    assert len(app.review_result.metrics_frame.winfo_children()) == 6
    text = app.review_result.issue_text.get("1.0", "end")
    assert "── Preenchimentos comprovados (automáticos) ──" in text
    assert "── Sugestões aguardando confirmação ──" in text
    assert "── Pendências sem solução automática ──" in text
    assert "Este item precisa de conferência manual antes de continuar." in text
    assert not app.is_busy
    assert app._poll_id is None
    assert str(app.prepare_button.cget("state")) == "normal"


def test_explain_pending_never_raises_on_blank_pendencias():
    assert explain_pending(None) == "Este item precisa de conferência manual antes de continuar."
    assert explain_pending("") == "Este item precisa de conferência manual antes de continuar."


@pytest.mark.parametrize("failure_stage", ["read", "callback"])
def test_saved_excel_summary_error_recovers_and_allows_next_preparation(
    tmp_path, monkeypatch, gui_application, failure_stage,
):
    root, app, _ = gui_application
    _setup_sources(root, app, monkeypatch, tmp_path)
    target = tmp_path / "saved.xlsx"
    monkeypatch.setattr(gui.filedialog, "asksaveasfilename", lambda **_: str(target))
    callback_errors = []
    monkeypatch.setattr(
        root, "report_callback_exception", lambda *args: callback_errors.append(args)
    )
    real_read, real_completed = gui.read_intermediate, app._prepared

    def fail(*_args):
        raise KeyError("fonte")

    if failure_stage == "read":
        monkeypatch.setattr(gui, "read_intermediate", fail)
    else:
        monkeypatch.setattr(app, "_prepared", fail)
    app.prepare_button.invoke()
    _wait_for_idle(root, app)
    assert target.exists()
    assert not callback_errors
    assert app.review_result.title.cget("text") == "Excel salvo; resumo indisponível"
    assert app.prepared_path.get() == str(target)
    assert not app.is_busy
    assert app._poll_id is None
    assert not app.progress.winfo_manager()
    assert str(app.prepare_again_button.cget("state")) == "normal"
    assert str(app.open_excel_button.cget("state")) == "normal"
    monkeypatch.setattr(gui, "read_intermediate", real_read)
    monkeypatch.setattr(app, "_prepared", real_completed)
    next_target = tmp_path / "next.xlsx"
    monkeypatch.setattr(gui.filedialog, "asksaveasfilename", lambda **_: str(next_target))
    app.prepare_again_button.invoke()
    _wait_for_idle(root, app)
    assert target.exists() and next_target.exists()
    assert not callback_errors
    assert app.review_result.title.cget("text") == "Excel preparado"
    assert len(app.review_result.metrics_frame.winfo_children()) == 6


def test_all_issues_remain_readable_and_unknown_guidance_is_neutral(
    tmp_path, monkeypatch, gui_application
):
    root, app, _ = gui_application
    root.geometry("900x600")
    root.deiconify()
    _wait_for_window(root, 900, 600)
    app.direct_excel_button.invoke()
    app.intermediate.set(str(tmp_path / "revisado.xlsx"))
    issues = [
        f"CREDITOS linha {i} (REF-{i}): DT_EMISSAO_TITULO inválida (data ausente)"
        for i in range(2, 33)
    ] + ["APONTAMENTO_NOVO"]
    service = Mock(side_effect=ValidationError(issues))
    monkeypatch.setattr(gui, "generate_cnab", service)
    target = tmp_path / "bloqueado.txt"
    monkeypatch.setattr(gui.filedialog, "asksaveasfilename", lambda **_: str(target))
    app.generate_button.invoke()
    _wait_for_idle(root, app)
    panel = app.generation_result
    title_top = panel.title.winfo_rooty() - app.body.canvas.winfo_rooty()
    assert 0 <= title_top <= 16
    assert title_top + panel.title.winfo_height() <= app.body.canvas.winfo_height()
    assert not target.exists()
    assert panel.kind == "error"
    assert len(panel.issues) == len(issues)
    friendly = panel.issue_text.get("1.0", "end")
    assert "CREDITOS, linha 32" in friendly
    assert "REF-" not in friendly
    assert "detalhes técnicos" in panel.issues[-1]
    assert panel.detail_text.get("1.0", "end-1c").splitlines() == issues
    panel.details_button.invoke()
    assert panel.details_open
    panel.details_button.invoke()
    assert not panel.details_open
    assert service.call_args.args == (str(tmp_path / "revisado.xlsx"),)
    assert not app.is_busy


@pytest.mark.parametrize(
    "failure", [PermissionError("private path"), RuntimeError("secret detail")]
)
def test_errors_recover_controls_without_raw_exception(
    tmp_path, monkeypatch, gui_application, failure
):
    root, app, _ = gui_application
    _setup_sources(root, app, monkeypatch, tmp_path)
    monkeypatch.setattr(gui, "prepare_workbook", Mock(side_effect=failure))
    target = tmp_path / "falhou.xlsx"
    monkeypatch.setattr(gui.filedialog, "asksaveasfilename", lambda **_: str(target))
    app.prepare_button.invoke()
    _wait_for_idle(root, app)
    assert app.review_result.kind == "error"
    assert str(failure) not in app.status.get()
    assert not target.exists()
    assert not app.is_busy
    assert str(app.analytic_entry.cget("state")) == "normal"


def test_direct_excel_bytes_match_service_without_sources(tmp_path, monkeypatch, gui_application):
    root, app, _ = gui_application
    pfmis, analytic, due = _single_normal_sources(tmp_path)
    excel = tmp_path / "somente_excel.xlsx"
    gui.prepare_workbook(
        pfmis[0],
        analytic,
        due,
        liquidation_date="05/09/2026",
        first_sequence="1001",
        output_path=excel,
    )
    expected = tmp_path / "expected.txt"
    gui.generate_cnab(excel, output_path=expected)
    for source in (*pfmis, analytic, due):
        source.unlink()
    assert not app.pfmi_tree.get_children()
    app.direct_excel_button.invoke()
    app.intermediate.set(str(excel))
    target = tmp_path / "pela_gui.txt"
    monkeypatch.setattr(gui.filedialog, "asksaveasfilename", lambda **_: str(target))
    app.generate_button.invoke()
    _wait_for_idle(root, app)
    assert target.read_bytes() == expected.read_bytes()
    assert app.generation_result.kind == "success"
    assert "R$ 150,00" in app.status.get()
    assert "R$ 100,00" in app.status.get()
    assert "Sequência: 1001–1001" in app.status.get()
    assert app.generated_path.get() == str(target)


def test_saved_output_audit_warning_stays_success(tmp_path, monkeypatch, gui_application):
    root, app, _ = gui_application
    root.geometry("900x600")
    root.deiconify()
    _wait_for_window(root, 900, 600)
    app.intermediate.set(str(tmp_path / "revisado.xlsx"))
    target = tmp_path / "salvo.txt"
    result = GenerationResult(
        target,
        1,
        Decimal("1234.56"),
        Decimal("1000"),
        1,
        1,
        1,
        ("LOG_AUDITORIA_NAO_GRAVADO; o TXT foi gerado normalmente",),
    )
    monkeypatch.setattr(gui, "generate_cnab", Mock(return_value=result))
    monkeypatch.setattr(gui.filedialog, "asksaveasfilename", lambda **_: str(target))
    app.generate_button.invoke()
    _wait_for_idle(root, app)
    assert app.generation_result.kind == "success"
    assert app.generated_path.get() == str(target)
    assert "O arquivo foi salvo" in app.generation_result.issues[0]
    assert "R$ 1.234,56" in app.status.get()
    title = app.generation_result.title
    title_top = title.winfo_rooty() - app.body.canvas.winfo_rooty()
    assert 0 <= title_top <= 16
    assert title_top + title.winfo_height() <= app.body.canvas.winfo_height()


def test_open_file_and_folder_use_arguments_and_report_association_failure(
    tmp_path, monkeypatch, gui_application
):
    root, app, _ = gui_application
    path = tmp_path / "Excel com espaços e acentuação.xlsx"
    path.touch()
    process = Mock()
    process.poll.side_effect = [None, 3]
    launch = Mock(return_value=process)
    monkeypatch.setattr(gui.sys, "platform", "linux")
    monkeypatch.setattr(gui.subprocess, "Popen", launch)
    app.prepared_path.set(str(path))
    app.open_excel_button.invoke()
    assert launch.call_args.args[0] == ["xdg-open", str(path)]
    assert "shell" not in launch.call_args.kwargs
    root.after_cancel(app._opener_poll_id)
    app._poll_openers()
    root.after_cancel(app._opener_poll_id)
    app._poll_openers()
    assert "abra manualmente" in app.opening_review_notice.cget("text")
    assert path.exists()
    assert app._opener_jobs == []
    launch.reset_mock()
    process.poll.side_effect = None
    process.poll.return_value = 0
    app.open_prepared_folder_button.invoke()
    assert launch.call_args.args[0] == ["xdg-open", str(tmp_path)]
    root.after_cancel(app._opener_poll_id)
    app._poll_openers()
    launch.side_effect = OSError("missing xdg-open")
    app.open_excel_button.invoke()
    assert "abra manualmente" in app.opening_review_notice.cget("text")
    path.unlink()
    app.open_excel_button.invoke()
    assert "não está mais" in app.opening_review_notice.cget("text")


def test_windows_opener_uses_startfile_without_shell(tmp_path, monkeypatch, gui_application):
    _, app, _ = gui_application
    path = tmp_path / "arquivo.xlsx"
    path.touch()
    launch = Mock()
    monkeypatch.setattr(gui.os, "startfile", launch, raising=False)
    monkeypatch.setattr(gui.sys, "platform", "win32")
    app._open_local(str(path), 2)
    launch.assert_called_once_with(str(path))


def test_no_operation_when_save_dialog_cancelled_preserves_previous_result(
    tmp_path, monkeypatch, gui_application
):
    root, app, _ = gui_application
    _setup_sources(root, app, monkeypatch, tmp_path)
    prepare = Mock()
    generate = Mock()
    monkeypatch.setattr(gui, "prepare_workbook", prepare)
    monkeypatch.setattr(gui, "generate_cnab", generate)
    monkeypatch.setattr(gui.filedialog, "asksaveasfilename", lambda **_: "")
    app.prepare_button.invoke()
    app.intermediate.set(str(tmp_path / "revisado.xlsx"))
    app.generation_result.show("TXT salvo com sucesso", kind="success")
    app.generate_button.invoke()
    assert app.generation_result.title.cget("text") == "TXT salvo com sucesso"
    assert not app.is_busy
    prepare.assert_not_called()
    generate.assert_not_called()


@pytest.mark.parametrize(
    ("message", "expected"),
    [
        ("COMISSAO_PARAMETRO_INVALIDO", "Volte à configuração, corrija a comissão"),
        ("linha sem crédito original do Analítico não pode ser selecionada", "Analítico"),
        ("DIVERGENCIA_VALOR", "centavos"),
        ("DATA_LIQUIDACAO: data inválida", "RESUMO"),
        ("PRIMEIRA_SEQUENCIA: número inválido", "RESUMO"),
        ("DIVERGENCIA_NOME", "somente divergência de nome"),
        ("CREDITOS linha 9 (ID): DOC_SACADO inválido", "documento do sacado"),
        ("CREDITOS linha 9 (ID): DOC_CEDENTE inválido", "documento do cedente"),
        ("controle de proveniência foi alterado", "novo Excel"),
        ("LOG_AUDITORIA_NAO_GRAVADO", "foi salvo"),
        ("NOME_TRUNCADO_CNAB", "40 caracteres"),
        ("Grupo G1: nenhum crédito selecionado", "Grupo G1"),
        ("INCLUIR_CNAB deve ser SIM ou NAO", "SIM ou NAO"),
        ("mesmo crédito do Analítico foi selecionado mais de uma vez", "mais de um grupo"),
        ("DT_EMISSAO_TITULO é posterior à DATA_LIQUIDACAO", "posterior"),
        ("pendência de origem PFMI", "PFMI de origem"),
        ("DOC_CEDENTE foi alterado no Excel", "alteração válida"),
        ("DESCONHECIDO", "detalhes técnicos"),
    ],
)
def test_guidance_preserves_meaning(message, expected):
    assert expected in explain_issue(message)


def test_local_currency_format():
    assert money(Decimal("0")) == "0,00"
    assert money(Decimal("99999999999.99")) == "99.999.999.999,99"


@pytest.mark.parametrize("desktop", [(1366, 768), (1920, 1080)])
@pytest.mark.parametrize("scale", [1.0, 1.5])
def test_desktop_layout_keyboard_and_scaling(
    desktop, scale, tmp_path, monkeypatch
):
    # The real Tk window uses the usable area of each simulated desktop; no service is run.
    from tkinter import TclError, Tk, font

    # Each scenario owns one interpreter; two simultaneous Tk roots interfere with focus.
    try:
        root = Tk()
    except TclError as exc:
        if os.environ.get("CNAB_NOX_REQUIRE_GUI") == "1":
            pytest.fail(f"Tk obrigatório indisponível: {exc}")
        pytest.skip(f"Tk indisponível nesta sessão: {exc}")
    root.withdraw()
    # Test client-area layout and real Tk focus without desktop activation policies.
    # Normal managed windows are covered separately by the visible synthetic smoke.
    root.overrideredirect(True)
    root.winfo_screenwidth = lambda: desktop[0]
    root.winfo_screenheight = lambda: desktop[1]
    root.tk.call("tk", "scaling", (96 / 72) * scale)
    app = gui.Application(root)
    root.deiconify()
    _wait_for_window(root, 1040, min(700, desktop[1] - 100))
    for field in (app.sequence_entry, app.modality_combo, app.selected_path_entry):
        assert font.Font(root=root, font=field.cget("font")).actual("size") == 11
    paths = tuple(
        str(tmp_path / ("pasta de exemplo com nome longo " * 3) / str(i) / "PFMI sintética.xlsx")
        for i in range(30)
    )
    monkeypatch.setattr(gui.filedialog, "askopenfilenames", lambda **_: paths)
    app.add_pfmi_button.invoke()
    root.update()
    assert app.pfmi_count.get() == "30 PFMIs"
    assert app.selected_path.get() == paths[-1]
    assert all(
        app.pfmi_tree.item(item, "values")[1] == "PFMI sintética.xlsx"
        for item in app.pfmi_tree.get_children()
    )

    def visible_inside(widget):
        assert widget.winfo_ismapped()
        x = widget.winfo_rootx() - root.winfo_rootx()
        y = widget.winfo_rooty() - root.winfo_rooty()
        assert 0 <= x < root.winfo_width()
        assert x + widget.winfo_width() <= root.winfo_width()
        assert 0 <= y < root.winfo_height()
        assert y + widget.winfo_height() <= root.winfo_height()

    try:
        app._prepared(
            PreparationResult(tmp_path / "sintetico.xlsx", 30, 30, 0, (), 30, 30),
            30,
            app._signature(),
        )
        app._generated(
            GenerationResult(
                tmp_path / "sintetico.txt", 30, Decimal("1234.56"), Decimal("1000"), 1, 30, 0, ()
            )
        )
        for geometry in (f"1040x{min(700, desktop[1] - 100)}", "900x600"):
            root.geometry(geometry)
            width, height = map(int, geometry.split("x"))
            _wait_for_window(root, width, height)
            for step in (1, 2, 3):
                app.show_step(step)
                root.update()
                visible_inside(app.title_label)
                visible_inside(app.direct_excel_button)
                for button in app.step_buttons:
                    visible_inside(button)
                    assert button.winfo_width() >= button.winfo_reqwidth()
                for button in app.footers[step].winfo_children():
                    if button.winfo_ismapped():
                        visible_inside(button)
                        assert button.winfo_width() >= button.winfo_reqwidth()
                assert [p.winfo_ismapped() for p in app.pages.values()] == [
                    step == 1,
                    step == 2,
                    step == 3,
                ]
                if step == 2:
                    visible_inside(app.open_excel_button)
                    assert app.open_excel_button.cget("style") == "Primary.TButton"
                if step == 3:
                    visible_inside(app.open_generated_folder_button)
                    assert app.open_generated_folder_button.cget("style") == "Primary.TButton"
            app.show_step(1)
            root.update()
            # Follow the actual Tk keyboard traversal chain, not a manually ordered list.
            current = app.add_pfmi_button
            seen = set()
            root.focus_force()
            for _ in range(80):
                current = current.tk_focusNext()
                if str(current) in seen:
                    break
                seen.add(str(current))
                # The desktop WM may focus another window between visual scenarios.
                # Acquire keyboard focus before checking the application's reveal behavior.
                current.focus_force()
                deadline = time.monotonic() + 2
                while time.monotonic() < deadline:
                    root.update()
                    if root.focus_get() == current:
                        break
                    time.sleep(0.005)
                assert root.focus_get() == current
                if current == app.sequence_entry:
                    top = current.winfo_rooty() - app.body.canvas.winfo_rooty()
                    assert top >= 0
                    assert top + current.winfo_height() <= app.body.canvas.winfo_height()
            assert str(app.sequence_entry) in seen
            assert str(app.intermediate_entry) not in seen
            app.body.canvas.yview_moveto(0)
            root.update()
            assert app.body.canvas.yview()[0] == 0
            app.instructions_label.event_generate("<Button-5>")
            root.update()
            assert app.body.canvas.yview()[0] > 0
            app.instructions_label.event_generate("<MouseWheel>", delta=120)
            root.update()
            assert app.body.canvas.yview()[0] == 0
    finally:
        # Same explicit cleanup as the shared gui_application fixture: cancel
        # pending after() polling jobs before destroy, then force GC so
        # PhotoImage/Tcl resources do not pile up across the 4 parametrized
        # interpreters this test creates.
        for attr in ("_poll_id", "_opener_poll_id", "_focus_reveal_id"):
            job_id = getattr(app, attr, None)
            if job_id:
                with contextlib.suppress(TclError):
                    root.after_cancel(job_id)
        for img in getattr(app, "step_icons", {}).values():
            with contextlib.suppress(TclError):
                img.tk.call("image", "delete", img.name)
        root.destroy()
        gc.collect()
