from __future__ import annotations

import contextlib
import gc
import time
from datetime import date
from pathlib import Path
from tkinter import TclError, Tk

import pytest
from openpyxl import Workbook, load_workbook

import gerador_cnab_nox.gui as gui_module
from gerador_cnab_nox.gui import CESSAO_LABEL, NORMAL_LABEL, Application
from gerador_cnab_nox.models import CESSAO, NORMAL

VALID_CPF = "52998224725"
VALID_CNPJ = "11222333000181"


@pytest.fixture
def gui_application(monkeypatch, _shared_tk_root):
    root = _shared_tk_root
    root.withdraw()
    dialogs: dict[str, list[tuple[str, str]]] = {
        "info": [],
        "warning": [],
        "error": [],
    }

    app = Application(root)
    # Record actual inline result rendering, not modal dialogs.
    for panel in app.results.values():
        original = panel.show

        def record(title, body="", *, _show=original, **kwargs):
            _show(title, body, **kwargs)
            kind = kwargs.get("kind", "info")
            key = "info" if kind == "success" else kind
            dialogs[key].append((title, body))

        monkeypatch.setattr(panel, "show", record)
    root.update_idletasks()
    try:
        yield root, app, dialogs
    finally:
        _wait_for_idle(root, app)
        # Pending after() polling jobs are never cancelled by destroy() itself
        # (WM_DELETE_WINDOW is not triggered by a direct destroy() call).
        # Cancel them explicitly, then force GC so PhotoImage/Tcl image
        # resources tied to reference cycles are freed before the next test,
        # instead of piling up across the whole GUI suite in one process.
        for attr in ("_poll_id", "_opener_poll_id", "_focus_reveal_id"):
            job_id = getattr(app, attr, None)
            if job_id:
                with contextlib.suppress(TclError):
                    root.after_cancel(job_id)
        # Photo images tied to Tk's Windows photo-image cleanup on interpreter
        # teardown are not reliably freed across hundreds of interpreters in
        # one process; deleting them explicitly while the interpreter is
        # still alive avoids the cumulative resource exhaustion.
        for img in getattr(app, "step_icons", {}).values():
            with contextlib.suppress(TclError):
                img.tk.call("image", "delete", img.name)
        for child in list(root.winfo_children()):
            child.destroy()
        gc.collect()


def _wait_for_idle(root, app, timeout=15):
    deadline = time.monotonic() + timeout
    while app.is_busy:
        root.update()
        if time.monotonic() > deadline:
            raise AssertionError("A operação gráfica não concluiu no prazo do teste")
        time.sleep(0.005)
    root.update()


def _write_sources(
    directory: Path,
    groups: list[dict[str, object]],
) -> tuple[list[Path], Path, Path]:
    pfmis: list[Path] = []
    for index, group in enumerate(groups, start=1):
        pfmi = directory / f"pfmi_{index}.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["FUNDO SINTETICO NOX"])
        sheet.append(
            [
                "VALOR DA OPERAÇÃO",
                "FALÊNCIA",
                "CREDOR / CEDENTE",
                "BENEFICIÁRIOS / FAVORECIDOS",
                "BCO",
                "AG",
                "CTA",
                "CPF/ CNPJ",
                "VALOR",
                "PIX",
            ]
        )
        beneficiary = "CONEXCRED" if group["cession"] else "BENEFICIARIO SINTETICO"
        beneficiary_document = VALID_CNPJ if group["cession"] else VALID_CPF
        sheet.append(
            [
                group["operation"],
                group["failure"],
                group["cedent"],
                beneficiary,
                "",
                "",
                "",
                beneficiary_document,
                group["operation"],
                "",
            ]
        )
        workbook.save(pfmi)
        workbook.close()
        pfmis.append(pfmi)

    analytic = directory / "analitico_sintetico.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "Cpf",
            "Contato",
            "Campanha",
            "Valor Receber",
            "Valor Aquisicao",
            "Data Assinatura",
            "ID",
        ]
    )
    for index, group in enumerate(groups, start=1):
        sheet.append(
            [
                VALID_CPF,
                group["cedent"],
                group["failure"],
                group["nominal"],
                group["acquisition"],
                date(2026, 9, 1),
                f"CREDITO-{index}",
            ]
        )
    workbook.save(analytic)
    workbook.close()

    due = directory / "vencimentos_sintetico.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["NOME_SACADO", "DOC_SACADO", "DATA_VENCIMENTO"])
    for group in groups:
        sheet.append([group["failure"], VALID_CNPJ, date(2030, 1, 31)])
    workbook.save(due)
    workbook.close()
    return pfmis, analytic, due


def _single_normal_sources(directory: Path) -> tuple[list[Path], Path, Path]:
    return _write_sources(
        directory,
        [
            {
                "cedent": "EMPRESA NORMAL",
                "failure": "FALENCIA NORMAL",
                "operation": 100,
                "nominal": 150,
                "acquisition": 100,
                "cession": False,
            }
        ],
    )


def _select_pfmi(root: Tk, app: Application, index: int) -> str:
    item_id = app.pfmi_tree.get_children()[index]
    app.pfmi_tree.selection_set(item_id)
    app.pfmi_tree.focus(item_id)
    app.pfmi_tree.event_generate("<<TreeviewSelect>>")
    root.update()
    return item_id


def _choose_modality(root: Tk, app: Application, label: str) -> None:
    app.modality_combo.set(label)
    app.modality_combo.event_generate("<<ComboboxSelected>>")
    root.update()


def _set_commission(root: Tk, app: Application, value: str) -> None:
    app.commission_entry.delete(0, "end")
    app.commission_entry.insert(0, value)
    root.update()


def test_t01_real_widgets_keep_order_modalities_and_individual_rates(
    tmp_path: Path,
    monkeypatch,
    gui_application,
) -> None:
    root, app, dialogs = gui_application
    groups = [
        {
            "cedent": "EMPRESA ZERO",
            "failure": "FALENCIA ZERO",
            "operation": 80,
            "nominal": 200,
            "acquisition": 80,
            "cession": True,
        },
        {
            "cedent": "EMPRESA FRACAO",
            "failure": "FALENCIA FRACAO",
            "operation": 100.50,
            "nominal": 200,
            "acquisition": 100,
            "cession": True,
        },
        {
            "cedent": "EMPRESA NORMAL",
            "failure": "FALENCIA NORMAL",
            "operation": 90,
            "nominal": 130,
            "acquisition": 90,
            "cession": False,
        },
    ]
    pfmis, analytic, due = _write_sources(tmp_path, groups)
    output = tmp_path / "intermediario_misto.xlsx"
    monkeypatch.setattr(
        gui_module.filedialog,
        "askopenfilenames",
        lambda **_kwargs: tuple(str(path) for path in pfmis),
    )
    monkeypatch.setattr(
        gui_module.filedialog,
        "asksaveasfilename",
        lambda **_kwargs: str(output),
    )

    app.add_pfmi_button.invoke()
    root.update()
    assert app.pfmi_count.get() == "3 PFMIs"
    assert len(app.pfmi_tree.get_children()) == 3

    _select_pfmi(root, app, 0)
    assert app.pfmi_modality.get() == NORMAL_LABEL
    _choose_modality(root, app, CESSAO_LABEL)
    assert app.pfmi_commission.get() == "15"
    _set_commission(root, app, "0")
    _choose_modality(root, app, NORMAL_LABEL)
    assert str(app.commission_entry.cget("state")) == "disabled"
    _choose_modality(root, app, CESSAO_LABEL)
    assert app.pfmi_commission.get() == "0"

    _select_pfmi(root, app, 1)
    _choose_modality(root, app, CESSAO_LABEL)
    assert app.pfmi_commission.get() == "15"
    _set_commission(root, app, "")
    _choose_modality(root, app, NORMAL_LABEL)
    _choose_modality(root, app, CESSAO_LABEL)
    assert app.pfmi_commission.get() == ""
    _set_commission(root, app, "0,5")

    normal_id = _select_pfmi(root, app, 2)
    assert app.pfmi_modality.get() == NORMAL_LABEL
    app.move_up_button.invoke()
    app.move_up_button.invoke()
    app.move_down_button.invoke()
    app.move_up_button.invoke()
    root.update()
    assert app.pfmi_tree.get_children()[0] == normal_id

    app.analytic.set(str(analytic))
    app.due.set(str(due))
    app.liquidation.set("03/09/2026")
    app.sequence.set("500")
    app.prepare_button.invoke()
    _wait_for_idle(root, app)
    root.update()

    assert output.exists()
    assert app.intermediate.get() == str(output)
    assert "operações: 3" in app.status.get()
    assert "prontas: 1" in app.status.get()
    assert "pendentes: 2" in app.status.get()
    assert not dialogs["error"]

    workbook = load_workbook(output, data_only=True)
    sheet = workbook["CREDITOS"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    rows = [
        {
            "arquivo": sheet.cell(row, headers["ARQUIVO_PFMI"]).value,
            "ordem": sheet.cell(row, headers["ORDEM_ARQUIVO"]).value,
            "modalidade": sheet.cell(row, headers["MODALIDADE"]).value,
            "comissao": sheet.cell(row, headers["COMISSAO_TEXTO"]).value,
        }
        for row in range(2, sheet.max_row + 1)
    ]
    workbook.close()
    assert rows == [
        {"arquivo": pfmis[2].name, "ordem": 1, "modalidade": NORMAL, "comissao": None},
        {"arquivo": pfmis[0].name, "ordem": 2, "modalidade": CESSAO, "comissao": "0"},
        {"arquivo": pfmis[1].name, "ordem": 3, "modalidade": CESSAO, "comissao": "0,5"},
    ]

    app.remove_pfmi_button.invoke()
    root.update()
    assert app.pfmi_count.get() == "2 PFMIs"
    assert len(app.pfmi_tree.get_children()) == 2


def test_missing_inputs_cancelled_dialogs_and_bad_destination_are_controlled(
    tmp_path: Path,
    monkeypatch,
    gui_application,
) -> None:
    root, app, dialogs = gui_application
    pfmis, analytic, due = _single_normal_sources(tmp_path)

    monkeypatch.setattr(
        gui_module.filedialog,
        "askopenfilenames",
        lambda **_kwargs: (),
    )
    app.add_pfmi_button.invoke()
    root.update()
    assert not app.pfmi_tree.get_children()

    app.prepare_button.invoke()
    _wait_for_idle(root, app)
    root.update()
    assert dialogs["warning"][-1][0] == "Arquivos obrigatórios"

    monkeypatch.setattr(
        gui_module.filedialog,
        "askopenfilenames",
        lambda **_kwargs: (str(pfmis[0]),),
    )
    app.add_pfmi_button.invoke()
    app.analytic.set(str(analytic))
    app.due.set(str(due))
    monkeypatch.setattr(
        gui_module.filedialog,
        "asksaveasfilename",
        lambda **_kwargs: "",
    )
    app.prepare_button.invoke()
    _wait_for_idle(root, app)
    root.update()
    assert not list(tmp_path.glob("excel_cnab_nox*.xlsx"))

    app.intermediate.set("")
    app.generate_button.invoke()
    _wait_for_idle(root, app)
    root.update()
    assert dialogs["warning"][-1][0] == "Excel obrigatório"

    app.intermediate.set(str(pfmis[0]))
    app.generate_button.invoke()
    _wait_for_idle(root, app)
    root.update()
    assert not list(tmp_path.glob("cnab_nox*.txt"))

    blocker = tmp_path / "destino_indisponivel"
    blocker.write_text("arquivo no lugar de diretório", encoding="utf-8")
    bad_output = blocker / "intermediario.xlsx"
    monkeypatch.setattr(
        gui_module.filedialog,
        "asksaveasfilename",
        lambda **_kwargs: str(bad_output),
    )
    app.prepare_button.invoke()
    _wait_for_idle(root, app)
    root.update()
    assert dialogs["error"]
    assert "Traceback" not in dialogs["error"][-1][1]
    assert not bad_output.exists()
    assert pfmis[0].exists() and analytic.exists() and due.exists()


def test_t19_t22_t23_prepare_pending_correct_excel_and_generate_only_from_excel(
    tmp_path: Path,
    monkeypatch,
    gui_application,
) -> None:
    root, app, dialogs = gui_application
    work = tmp_path / "fluxo"
    work.mkdir()
    pfmis, analytic, due = _single_normal_sources(work)
    intermediate = work / "intermediario.xlsx"
    txt = work / "resultado.txt"
    (work / "logs").write_text("log indisponível de propósito", encoding="utf-8")

    monkeypatch.setattr(
        gui_module.filedialog,
        "askopenfilenames",
        lambda **_kwargs: (str(pfmis[0]),),
    )
    destinations = iter((str(intermediate), str(txt)))
    monkeypatch.setattr(
        gui_module.filedialog,
        "asksaveasfilename",
        lambda **_kwargs: next(destinations),
    )

    app.add_pfmi_button.invoke()
    app.analytic.set(str(analytic))
    app.due.set(str(due))
    app.prepare_button.invoke()
    _wait_for_idle(root, app)
    root.update()

    assert intermediate.exists()
    assert "operações: 1" in app.status.get()
    assert "prontas: 1" in app.status.get()
    assert "pendentes: 0" in app.status.get()
    assert "LOG_AUDITORIA_NAO_GRAVADO" in app.results[app.current_step].details

    workbook = load_workbook(intermediate)
    summary = workbook["RESUMO"]
    assert "DATA_LIQUIDACAO_AUSENTE" in summary["B5"].value
    assert "PRIMEIRA_SEQUENCIA_AUSENTE" in summary["B5"].value
    summary["B3"] = "03/09/2026"
    summary["B4"] = 730
    workbook.save(intermediate)
    workbook.close()

    for source in (*pfmis, analytic, due):
        source.unlink()
    _select_pfmi(root, app, 0)
    _choose_modality(root, app, CESSAO_LABEL)
    _set_commission(root, app, "99")
    app.analytic.set(str(work / "analitico_inexistente.xlsx"))
    app.due.set(str(work / "base_inexistente.xlsx"))
    app.liquidation.set("31/12/2099")
    app.sequence.set("9999")

    app.generate_button.invoke()
    _wait_for_idle(root, app)
    root.update()

    assert txt.exists()
    records = txt.read_bytes().split(b"\r\n")[:-1]
    assert len(records) == 3
    assert all(len(record) == 444 for record in records)
    assert records[1][37:62].strip() == b"730"
    assert "Títulos: 1" in app.status.get()
    assert "Sequência: 730–730" in app.status.get()
    assert "Avisos: 1" in app.status.get()
    assert "LOG_AUDITORIA_NAO_GRAVADO" in app.results[app.current_step].details
    assert not dialogs["error"]
