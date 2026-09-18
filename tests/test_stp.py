from __future__ import annotations

import json
from dataclasses import replace
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from gerador_cnab_nox.debtor_fallbacks import BEGIN, END, load_fallbacks
from gerador_cnab_nox.errors import InputFileError, ValidationError
from gerador_cnab_nox.matching import prepare_batch
from gerador_cnab_nox.models import PfmiData
from gerador_cnab_nox.service import generate_cnab, prepare_workbook
from gerador_cnab_nox.stp import SYSTEM_APPROVAL, strong_name_match
from gerador_cnab_nox.validation import validate_for_generation
from gerador_cnab_nox.workbook import read_intermediate, write_intermediate

from .conftest import VALID_CNPJ, VALID_CPF
from .test_pfmi_composicoes import _credit, _due, _pay


def _batch(credits=None, *, principal="MARIA ANGELICA NALIN 1", satellite="JOSE VALDIR"):
    # Different documents ensure the principal is resolved by name evidence, not accident.
    payments = [
        replace(_pay(principal, "16000", doc=VALID_CNPJ), beneficiary_name=principal),
        replace(_pay(satellite, "8000", doc=VALID_CPF), beneficiary_name="JOSE VALDIR"),
    ]
    credits = credits if credits is not None else [
        _credit("MARIA-1", "MARIA ANGELICA", acquisition="24000", nominal="48000",
                source_sheet="Analitico", source_row=2918),
    ]
    return prepare_batch(
        PfmiData(payments=payments), credits, [_due()],
        liquidation_date=date(2026, 9, 17), first_sequence=1,
    )


def _save_and_include(batch, tmp_path):
    path = tmp_path / "review.xlsx"
    write_intermediate(batch, path)
    book = load_workbook(path)
    sheet = book["CREDITOS"]
    headers = {c.value: c.column for c in sheet[1]}
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, headers["INCLUIR_CNAB"], "SIM")
    book.save(path)
    book.close()
    return path


def test_maria_and_jose_generated_excel_needs_only_binary_decision(tmp_path):
    batch = _batch()
    assert len(batch.rows) == 2
    for row in batch.rows:
        assert row.values["COMPOSICAO_SELECAO_MANUAL_ABA"] == "Analitico"
        assert row.values["COMPOSICAO_SELECAO_MANUAL_LINHA"] == 2918
        assert row.values["COMPOSICAO_APROVADA"] == "SIM"
        assert row.values["APROVADO"] == "SIM"
        assert row.values["INCLUIR_CNAB"] == "SIM"
        assert "AUTOMATED_MATCH_APPLIED" in row.values["ALERTAS"]
    assert [r.values["VL_NOMINAL"] for r in batch.rows] == [Decimal(32000), Decimal(16000)]
    path = _save_and_include(batch, tmp_path)
    # No Analitico or PFMI source is reopened during generation.
    result = generate_cnab(path, output_path=tmp_path / "payments.txt")
    assert result.detail_count == 2
    assert result.total_present == Decimal(24000)
    assert result.total_nominal == Decimal(48000)


@pytest.mark.parametrize("name", ["ESPOLIO DE FRANCISCO MOURA", "MARIA ANGELICA NALIN 1"])
def test_strong_document_autofills_name_approval(name, tmp_path):
    batch = prepare_batch(
        PfmiData(payments=[_pay(name, "100", doc=VALID_CPF)]),
        [_credit("A", "NOME CADASTRAL DIFERENTE", source_sheet="Analitico")], [_due()],
        liquidation_date=date(2026, 9, 17), first_sequence=1,
    )
    row = batch.rows[0]
    assert row.values["APROVADO"] == "SIM"
    assert row.values["STATUS"] == "OK_COM_ALERTA_NOME"
    assert not row.values["PENDENCIAS"]
    result = generate_cnab(_save_and_include(batch, tmp_path), output_path=tmp_path / "out.txt")
    assert result.detail_count == 1


@pytest.mark.parametrize("offset", ["0", "50.00"])
def test_ambiguous_value_aborts_before_document_or_name_disambiguation(offset):
    candidates = [
        _credit("A", "MARIA ANGELICA", acquisition="24000", source_sheet="Analitico"),
        _credit("B", "OUTRO TITULAR", acquisition=str(Decimal(24000) + Decimal(offset)),
                source_sheet="Analitico", source_row=3),
    ]
    batch = _batch(candidates)
    assert all(r.values["COMPOSICAO_APROVADA"] not in {"SIM", SYSTEM_APPROVAL} for r in batch.rows)
    assert all(r.values["STATUS"] == "COMPOSICAO_BLOQUEADA_SELECAO_MANUAL" for r in batch.rows)


def test_value_alone_never_proves_identity():
    batch = _batch([
        _credit("A", "PESSOA TOTALMENTE DIFERENTE", acquisition="24000", doc=VALID_CPF,
                source_sheet="Analitico"),
    ])
    assert all(r.values["COMPOSICAO_APROVADA"] not in {"SIM", SYSTEM_APPROVAL} for r in batch.rows)
    assert not strong_name_match("MARIA", "MARIA ANGELICA")


def test_exact_document_composition_ignores_unrelated_equal_values(tmp_path):
    batch = _batch([
        _credit("CONFIRMED", "NOME CADASTRAL DIFERENTE", acquisition="24000", nominal="48000",
                doc=VALID_CNPJ, source_sheet="Analitico", source_row=10),
        _credit("OTHER", "OUTRO TITULAR", acquisition="24000", doc=VALID_CPF,
                source_sheet="Analitico", source_row=11),
    ])
    assert batch.compositions[0]["auto_approved"]
    for row in batch.rows:
        assert row.values["COMPOSICAO_APROVADA"] == "SIM"
        assert row.values["COMPOSICAO_SELECAO_MANUAL_LINHA"] == 10
        assert not row.values["PENDENCIAS"]
        assert not any(e.get("exige_aprovacao_humana")
                       for e in json.loads(row.values["PREENCHIMENTOS_AUTOMATICOS"])
                       if e["campo"] == "VL_NOMINAL")
    result = generate_cnab(_save_and_include(batch, tmp_path), output_path=tmp_path / "out.txt")
    assert result.detail_count == 2
    assert result.total_present == Decimal(24000)


@pytest.mark.parametrize("offset", ["0", "50"])
def test_same_document_two_compatible_credits_still_require_selection(offset):
    batch = _batch([
        _credit("A", "MARIA ANGELICA", acquisition="24000", doc=VALID_CNPJ,
                source_sheet="Analitico", source_row=10),
        _credit("B", "MARIA ANGELICA", acquisition=str(Decimal(24000) + Decimal(offset)),
                doc=VALID_CNPJ, source_sheet="Analitico", source_row=11),
    ])
    assert not batch.compositions[0]["auto_approved"]
    assert all(r.values["COMPOSICAO_APROVADA"] not in {"SIM", SYSTEM_APPROVAL} for r in batch.rows)


def test_exact_document_composition_selects_only_same_failure():
    batch = _batch([
        _credit("A", "MARIA ANGELICA", acquisition="24000", doc=VALID_CNPJ,
                source_sheet="Analitico", source_row=10),
        _credit("B", "MARIA ANGELICA", acquisition="24000", doc=VALID_CNPJ,
                failure="OUTRA FALENCIA", source_sheet="Analitico", source_row=11),
    ])
    assert batch.compositions[0]["auto_approved"]
    assert all(r.values["COMPOSICAO_SELECAO_MANUAL_LINHA"] == 10 for r in batch.rows)


@pytest.mark.parametrize("field", ["block_id", "source_sheet", "file_id"])
def test_bare_satellite_never_crosses_physical_boundaries(field):
    principal = _pay("MARIA ANGELICA NALIN 1", "16000", doc=VALID_CNPJ)
    satellite = replace(_pay("JOSE VALDIR", "8000"), **{field: "OUTRO"})
    batch = prepare_batch(
        PfmiData(payments=[principal, satellite]),
        [_credit("A", "MARIA ANGELICA", acquisition="24000", source_sheet="Analitico")],
        [_due()], liquidation_date=date(2026, 9, 17), first_sequence=1,
    )
    assert not batch.compositions


@pytest.mark.parametrize("field", ["APROVADO", "COMPOSICAO_APROVADA"])
def test_typed_system_approval_without_provenance_is_rejected(tmp_path, field):
    batch = prepare_batch(
        PfmiData(payments=[_pay("TITULAR", "100")]), [_credit("A", "TITULAR")], [_due()],
        liquidation_date=date(2026, 9, 17), first_sequence=1,
    )
    path = _save_and_include(batch, tmp_path)
    loaded = read_intermediate(path)
    loaded.rows[0].values[field] = SYSTEM_APPROVAL
    with pytest.raises(ValidationError, match="sem origem comprovada"):
        validate_for_generation(loaded)


@pytest.mark.parametrize("field,value", [
    ("COMPOSICAO_SELECAO_MANUAL_LINHA", 99), ("NOME_CEDENTE", "OUTRO TITULAR"),
    ("VL_NOMINAL", Decimal("1.00")),
])
def test_edit_cannot_inherit_system_approval(tmp_path, field, value):
    loaded = read_intermediate(_save_and_include(_batch(), tmp_path))
    # Compatibility with protected snapshots produced before SIM auto-assignment.
    for row in loaded.rows:
        for approval in ("APROVADO", "COMPOSICAO_APROVADA", "SELECAO_MANUAL_APROVADA"):
            row.values[approval] = row.originals[approval] = SYSTEM_APPROVAL
    loaded.rows[0].values[field] = value
    with pytest.raises(ValidationError, match="campos aprovados alterados"):
        validate_for_generation(loaded)


def _fallback_file(tmp_path, document=VALID_CNPJ):
    path = tmp_path / "fallback.md"
    data = {"COSTEIRA": {"cnpj": document, "nome": "MASSA FALIDA COSTEIRA"}}
    path.write_text(f"{BEGIN}\n```json\n{json.dumps(data)}\n```\n{END}\n", encoding="utf-8")
    return path


@pytest.mark.parametrize("date_value", [date(2026, 9, 17), date(2027, 3, 2)])
def test_fallback_reads_markdown_and_uses_batch_liquidation(tmp_path, date_value):
    fallbacks = load_fallbacks(_fallback_file(tmp_path, "11.222.333/0001-81"))
    batch = prepare_batch(
        PfmiData(payments=[_pay("TITULAR", "100", failure="COSTEIRA")]),
        [_credit("A", "TITULAR", failure="COSTEIRA")], [],
        liquidation_date=date_value, first_sequence=1, debtor_fallbacks=fallbacks,
    )
    row = batch.rows[0]
    assert row.values["DOC_SACADO"] == VALID_CNPJ
    assert row.values["NOME_SACADO"] == "MASSA FALIDA COSTEIRA"
    assert row.values["DT_VENCIMENTO"] == date_value
    assert "FALLBACK_SACADO" in row.values["ALERTAS"]
    path = _save_and_include(batch, tmp_path)
    book = load_workbook(path)
    headers = {c.value: c.column for c in book["CREDITOS"][1]}
    cell = book["CREDITOS"].cell(2, headers["DT_VENCIMENTO"])
    assert cell.number_format == "DD/MM/YYYY"
    book.close()
    assert generate_cnab(path, output_path=tmp_path / "out.txt").detail_count == 1


@pytest.mark.parametrize("document", ["123", "11.222.333/0001-81a", "", 11222333000181])
def test_invalid_fallback_config_blocks(tmp_path, document):
    with pytest.raises(InputFileError):
        load_fallbacks(_fallback_file(tmp_path, document))


def test_development_cnpjs_are_filled_but_do_not_bypass_checksum(tmp_path):
    config = load_fallbacks(_fallback_file(tmp_path, "00000000000100"))
    assert config["COSTEIRA"]["document"] == "00000000000100"
    batch = prepare_batch(
        PfmiData(payments=[_pay("TITULAR", "100", failure="COSTEIRA")]),
        [_credit("A", "TITULAR", failure="COSTEIRA")], [],
        liquidation_date=date(2026, 9, 17), first_sequence=1, debtor_fallbacks=config,
    )
    path = _save_and_include(batch, tmp_path)
    with pytest.raises(ValidationError, match="DOC_SACADO"):
        generate_cnab(path, output_path=tmp_path / "blocked.txt")
    assert not (tmp_path / "blocked.txt").exists()


def test_binary_no_excludes_whole_operation_without_blocking_other_credits(tmp_path):
    batch = _batch()
    for row in batch.rows:
        row.values["INCLUIR_CNAB"] = "NAO"
    other = prepare_batch(
        PfmiData(payments=[_pay("OUTRA OPERACAO", "100", block="B0002")]),
        [_credit("OTHER", "OUTRA OPERACAO")], [_due()],
        liquidation_date=date(2026, 9, 17), first_sequence=100,
    )
    # The auto-approved composition remains NAO; only the other operation is selected.
    batch.rows.extend(other.rows)
    path = tmp_path / "review.xlsx"
    write_intermediate(batch, path)
    result = generate_cnab(path, output_path=tmp_path / "out.txt")
    assert result.detail_count == 1
    assert result.total_present == Decimal(100)


def test_partial_system_composition_stays_blocked(tmp_path):
    loaded = read_intermediate(_save_and_include(_batch(), tmp_path))
    loaded.rows[-1].values["INCLUIR_CNAB"] = "NAO"
    with pytest.raises(ValidationError, match="tudo ou nada"):
        validate_for_generation(loaded)


def test_fallback_does_not_replace_valid_base_and_resolves_aliases(tmp_path):
    fallback = load_fallbacks(_fallback_file(tmp_path))
    payments = PfmiData(payments=[_pay("TITULAR", "100", failure="COSTEIRA ALIAS")])
    kwargs = dict(
        liquidation_date=date(2026, 9, 17), first_sequence=1, debtor_fallbacks=fallback,
        failure_aliases={"COSTEIRA ALIAS": "COSTEIRA"},
    )
    credits = [_credit("A", "TITULAR", failure="COSTEIRA")]
    batch = prepare_batch(payments, credits, [_due("COSTEIRA")], **kwargs)
    assert batch.rows[0].values["DT_VENCIMENTO"] == date(2030, 1, 1)
    assert "FALLBACK_SACADO" not in batch.rows[0].values["ALERTAS"]
    batch = prepare_batch(payments, credits, [], **kwargs)
    assert batch.rows[0].values["DT_VENCIMENTO"] == date(2026, 9, 17)
    assert "FALLBACK_SACADO" in batch.rows[0].values["ALERTAS"]


def test_human_location_override_supersedes_robot_and_refreshes_only_automatic_fields(tmp_path):
    loaded = read_intermediate(_save_and_include(_batch(), tmp_path))
    corrected = _credit(
        "CORRECT", "MARIA ANGELICA CORRIGIDA", acquisition="24000", nominal="72000",
        doc=VALID_CNPJ, source_sheet="Escolha humana", source_row=42,
    )
    for row in loaded.rows:
        row.values["COMPOSICAO_SELECAO_MANUAL_ABA"] = corrected.source_sheet
        row.values["COMPOSICAO_SELECAO_MANUAL_LINHA"] = corrected.source_row
        row.values["COMPOSICAO_APROVADA"] = "SIM"
    result = validate_for_generation(loaded, analytic=[corrected])
    assert all(r["ID_CREDITO"] == corrected.credit_id for r in result.rows)
    assert result.rows[0]["NOME_CEDENTE"] == corrected.cedent_name
    assert result.rows[0]["DOC_CEDENTE"] == VALID_CNPJ
    assert [r["VL_NOMINAL"] for r in result.rows] == [Decimal(48000), Decimal(24000)]
    assert any("MANUAL_OVERRIDE=SIM" in warning for warning in result.warnings)
    assert loaded.rows[0].originals["ID_CREDITO"] != corrected.credit_id


def test_explicit_human_values_are_never_replaced_by_new_robot_defaults(tmp_path):
    loaded = read_intermediate(_save_and_include(_batch(), tmp_path))
    corrected = _credit(
        "CORRECT", "MARIA CADASTRAL", acquisition="24000", nominal="72000",
        source_sheet="Analitico", source_row=42,
    )
    for row, nominal in zip(loaded.rows, (Decimal(30000), Decimal(42000)), strict=True):
        row.values.update({
            "COMPOSICAO_SELECAO_MANUAL_LINHA": 42, "COMPOSICAO_APROVADA": "SIM",
            "APROVADO": "SIM", "VL_NOMINAL": nominal,
        })
    loaded.rows[0].values["NOME_CEDENTE"] = "NOME CONFERIDO PELO OPERADOR"
    result = validate_for_generation(loaded, analytic=[corrected])
    assert result.rows[0]["NOME_CEDENTE"] == "NOME CONFERIDO PELO OPERADOR"
    assert [r["VL_NOMINAL"] for r in result.rows] == [Decimal(30000), Decimal(42000)]


def test_human_can_change_a_name_without_reopening_unchanged_system_location(tmp_path):
    loaded = read_intermediate(_save_and_include(_batch(), tmp_path))
    loaded.rows[0].values["NOME_CEDENTE"] = "MARIA ANGELICA CONFERIDA"
    loaded.rows[0].values["APROVADO"] = "SIM"
    result = validate_for_generation(loaded)
    assert result.rows[0]["NOME_CEDENTE"] == "MARIA ANGELICA CONFERIDA"


def test_human_can_replace_an_ordinary_automatic_credit_by_location(tmp_path):
    batch = prepare_batch(
        PfmiData(payments=[_pay("TITULAR", "100")]), [_credit("OLD", "TITULAR")], [_due()],
        liquidation_date=date(2026, 9, 17), first_sequence=1,
    )
    loaded = read_intermediate(_save_and_include(batch, tmp_path))
    row = loaded.rows[0]
    row.values.update({
        "COMPOSICAO_SELECAO_MANUAL_ABA": "Analitico",
        "COMPOSICAO_SELECAO_MANUAL_LINHA": 42, "APROVADO": "SIM",
    })
    corrected = _credit("NEW", "TITULAR CORRETO", nominal="700", source_row=42,
                        source_sheet="Analitico", doc=VALID_CNPJ)
    result = validate_for_generation(loaded, analytic=[corrected])
    assert result.rows[0]["ID_CREDITO"] == corrected.credit_id
    assert result.rows[0]["VL_NOMINAL"] == Decimal(700)
    assert result.rows[0]["DOC_CEDENTE"] == VALID_CNPJ


def test_partial_human_location_override_is_rejected(tmp_path):
    loaded = read_intermediate(_save_and_include(_batch(), tmp_path))
    loaded.rows[0].values["COMPOSICAO_SELECAO_MANUAL_LINHA"] = 42
    loaded.rows[0].values["COMPOSICAO_APROVADA"] = "SIM"
    corrected = _credit("NEW", "MARIA ANGELICA", acquisition="24000", source_row=42,
                        source_sheet="Analitico")
    with pytest.raises(ValidationError, match="iguais em todas as linhas"):
        validate_for_generation(loaded, analytic=[corrected])


def test_human_override_excel_to_real_txt_with_a_different_credit(tmp_path):
    path = _save_and_include(_batch(), tmp_path)
    book = load_workbook(path)
    sheet = book["CREDITOS"]
    headers = {c.value: c.column for c in sheet[1]}
    for index in range(2, sheet.max_row + 1):
        sheet.cell(index, headers["COMPOSICAO_SELECAO_MANUAL_LINHA"], 3)
        sheet.cell(index, headers["COMPOSICAO_APROVADA"], "SIM")
    book.save(path)
    book.close()
    source = tmp_path / "analitico.xlsx"
    book = Workbook()
    book.active.title = "Analitico"
    book.active.append([
        "Cpf", "Contato", "Campanha", "Valor Receber", "Valor Aquisicao", "Data Assinatura", "ID",
    ])
    book.active.append([
        VALID_CPF, "MARIA ANGELICA", "PROSPECT", 48000, 24000, date(2026, 9, 1), "OLD",
    ])
    book.active.append([
        VALID_CNPJ, "MARIA CORRIGIDA", "PROSPECT", 72000, 24000, date(2026, 9, 2), "NEW",
    ])
    book.save(source)
    book.close()
    result = generate_cnab(path, analytic_path=source, output_path=tmp_path / "override.txt")
    assert result.total_nominal == Decimal(72000)
    assert result.total_present == Decimal(24000)
    assert result.reconciliation.clearedRecords[0].documentNumber == VALID_CNPJ
    assert b"MARIA CORRIGIDA" in result.output_path.read_bytes()
    assert any("MANUAL_OVERRIDE=SIM" in warning for warning in result.warnings)


def test_service_captures_fallback_without_generation_dependency(source_files, tmp_path):
    pfmi, analytic, due = source_files
    book = load_workbook(pfmi)
    book.active["B3"] = book.active["B4"] = "COSTEIRA"
    book.save(pfmi)
    book.close()
    book = load_workbook(analytic)
    book.active["C2"] = "COSTEIRA"
    book.save(analytic)
    book.close()
    config = _fallback_file(tmp_path)
    path = tmp_path / "review.xlsx"
    prepare_workbook(
        pfmi, analytic, due, liquidation_date="17/09/2026", first_sequence=1,
        output_path=path, fallback_document_path=config,
    )
    for source in (pfmi, analytic, due, config):
        source.unlink()
    result = generate_cnab(path, output_path=tmp_path / "out.txt")
    assert result.detail_count == 1
    assert result.total_present == Decimal(100)
