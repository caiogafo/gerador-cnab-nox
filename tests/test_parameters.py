from decimal import Decimal

import pytest

from gerador_cnab_nox.models import CESSAO, NORMAL, PfmiInput, PreparationRequest
from gerador_cnab_nox.normalize import commission_rate, money


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("0", "0"),
        ("15", "15"),
        ("15,0", "15"),
        ("15.0", "15"),
        ("0,5", "0.5"),
        ("0.5", "0.5"),
        ("101", "101"),
        (" 2.25 ", "2.25"),
        ("0.0000000000000000000000000001", "1e-28"),
    ],
)
def test_t02_valid_rates(text, expected):
    assert commission_rate(text, CESSAO) == Decimal(expected)


@pytest.mark.parametrize(
    "text",
    [
        "",
        None,
        "-1",
        "NaN",
        "inf",
        "Infinity",
        "x",
        "1e2",
        "1,2.3",
        "9" * 29,
        "0." + "0" * 28 + "1",
        "1.1234567890123456789012345678",
    ],
)
def test_t03_invalid_rate_is_explicit(text):
    with pytest.raises(ValueError, match="COMISSAO_PARAMETRO_INVALIDO"):
        commission_rate(text, CESSAO)
    assert commission_rate(text, NORMAL) == 0


def test_t01_request_keeps_order_and_individual_zero_blank():
    request = PreparationRequest(
        [PfmiInput("b.xlsx", CESSAO, "0"), PfmiInput("a.xlsx", CESSAO, "")],
        "analytic.xlsx",
        "base.xlsx",
        "bad",
        "",
    )
    assert [(p.path, p.comissao_texto) for p in request.pfmis] == [("b.xlsx", "0"), ("a.xlsx", "")]


@pytest.mark.parametrize("value", ["NaN", "Infinity", "1e999999", "100000000000.00"])
def test_extreme_money_has_controlled_failure(value):
    with pytest.raises(ValueError):
        money(value)


@pytest.mark.parametrize(
    "rate", ["", "-1", "NaN", "Infinity", "abc", "9" * 29, "0." + "0" * 28 + "1"]
)
def test_t03_invalid_parameter_prepares_but_cannot_be_repaired_in_finals(
    tmp_path, source_files, rate
):
    from datetime import date

    from openpyxl import load_workbook

    from gerador_cnab_nox.errors import ValidationError
    from gerador_cnab_nox.service import generate_cnab, prepare_workbook
    from gerador_cnab_nox.workbook import read_intermediate

    pfmi, analytic, due = source_files
    output = tmp_path / "parameter.xlsx"
    prepare_workbook(
        [PfmiInput(pfmi, CESSAO, rate)],
        analytic,
        due,
        liquidation_date="bad",
        first_sequence="",
        output_path=output,
    )
    loaded = read_intermediate(output)
    assert not loaded.structural_issues
    assert loaded.rows[0].originals["COMISSAO_TEXTO"] == ("'-1" if rate == "-1" else rate or None)
    assert loaded.rows[0].originals["PENDENCIAS_PARAMETROS"] == "COMISSAO_PARAMETRO_INVALIDO"
    book = load_workbook(output)
    sh = book["CREDITOS"]
    h = {c.value: c.column for c in sh[1]}
    for key, value in {
        "INCLUIR_CNAB": "SIM",
        "APROVADO": "SIM",
        "VL_PRESENTE": 100,
        "DOC_CEDENTE": "11222333000181",
        "TIPO_PESSOA_CEDENTE": 2,
        "DT_EMISSAO_TITULO": date(2026, 9, 1),
    }.items():
        sh.cell(2, h[key]).value = value
    book["RESUMO"]["B3"] = date(2026, 9, 3)
    book["RESUMO"]["B4"] = 1
    book.save(output)
    book.close()
    with pytest.raises(ValidationError, match="COMISSAO_PARAMETRO_INVALIDO"):
        generate_cnab(output, output_path=tmp_path / "blocked.txt")
    assert not (tmp_path / "blocked.txt").exists()
