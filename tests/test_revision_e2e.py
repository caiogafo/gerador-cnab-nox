from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from gerador_cnab_nox.errors import ValidationError
from gerador_cnab_nox.models import CESSAO, IMMUTABLE_FIELDS, PfmiInput
from gerador_cnab_nox.service import generate_cnab, prepare_workbook
from gerador_cnab_nox.workbook import read_intermediate


def edit(path, values, row=2, summary=None):
    book = load_workbook(path)
    sh = book["CREDITOS"]
    headers = {c.value: c.column for c in sh[1]}
    for key, value in values.items():
        sh.cell(row, headers[key]).value = value
    for key, value in (summary or {}).items():
        book["RESUMO"][key] = value
    book.save(path)
    book.close()


@pytest.fixture
def prepared(tmp_path, source_files):
    path = tmp_path / "revisao.xlsx"
    prepare_workbook(
        *source_files, liquidation_date="03/09/2026", first_sequence=10, output_path=path
    )
    return path


@pytest.mark.parametrize("field", IMMUTABLE_FIELDS)
def test_t18_every_visible_original_control_is_immutable(prepared, tmp_path, field):
    edit(prepared, {field: "ALTERACAO_INDEVIDA"})
    with pytest.raises(ValidationError):
        generate_cnab(prepared, output_path=tmp_path / "blocked.txt")
    assert not (tmp_path / "blocked.txt").exists()


@pytest.mark.parametrize(
    "mutation",
    ["summary", "manifest", "original", "pending", "schema", "duplicate_column", "insert_blank"],
)
def test_t18_structural_and_snapshot_mutations(prepared, tmp_path, mutation):
    book = load_workbook(prepared)
    if mutation == "summary":
        book["RESUMO"]["B6"] = 999
    elif mutation == "manifest":
        book["MANIFESTO"]["B3"] = 999
    elif mutation == "original":
        book["ORIGINAIS_CONTROLE"]["B2"] = "MUTADO"
    elif mutation == "pending":
        book["PENDENCIAS_PFMI"].append(["FALSA_LINHA"])
    elif mutation == "schema":
        book["RESUMO"]["B2"] = "CNAB-NOX-V1-1"
    elif mutation == "duplicate_column":
        book["CREDITOS"].cell(1, book["CREDITOS"].max_column + 1).value = "ID_LINHA"
    else:
        book["CREDITOS"].insert_rows(2)
    book.save(prepared)
    book.close()
    with pytest.raises(ValidationError):
        generate_cnab(prepared, output_path=tmp_path / "blocked.txt")


def test_t16_t19_final_corrections_prevail_and_sources_are_not_read(
    prepared, source_files, tmp_path
):
    before = read_intermediate(prepared).rows[0].originals.copy()
    # A new name is a human correction, acknowledged only as a name divergence.
    edit(
        prepared,
        {
            "NOME_CEDENTE": "EMPRESA CORRIGIDA",
            "DOC_CEDENTE": "11222333000181",
            "TIPO_PESSOA_CEDENTE": 2,
            "VL_NOMINAL": 160,
            "DT_EMISSAO_TITULO": date(2026, 9, 2),
            "APROVADO": "SIM",
        },
    )
    after = read_intermediate(prepared).rows[0].originals
    assert before == after
    for source in source_files:
        source.unlink()
    result = generate_cnab(prepared, output_path=tmp_path / "final.txt")
    assert result.total_nominal == Decimal("160") and result.total_present == Decimal("100")
    assert result.warning_count >= 4
    detail = result.output_path.read_bytes().split(b"\r\n")[1]
    assert detail[334:374].rstrip() == b"EMPRESA CORRIGIDA"
    assert detail[380:394] == b"11222333000181" and detail[150:156] == b"020926"


def test_t16_final_present_can_correct_bad_suggestion(source_files, tmp_path):
    pfmi, analytic, due = source_files
    book = load_workbook(analytic)
    book.active.cell(2, 5).value = 99
    book.save(analytic)
    book.close()
    path = tmp_path / "correcao.xlsx"
    prepare_workbook(
        pfmi, analytic, due, liquidation_date="03/09/2026", first_sequence=1, output_path=path
    )
    before = read_intermediate(path).rows[0].originals
    assert before["AQUISICAO_ORIGINAL"] == 99 and before["AQUISICAO_ESPERADA"] == 99
    edit(path, {"INCLUIR_CNAB": "SIM", "VL_PRESENTE": 100})
    result = generate_cnab(path, output_path=tmp_path / "corrigido.txt")
    assert result.total_present == Decimal("100")
    assert read_intermediate(path).rows[0].originals == before
    assert any("VL_PRESENTE foi alterado" in w for w in result.warnings)


@pytest.mark.parametrize(
    "field,value,message",
    [
        ("VL_PRESENTE", 100.01, "DIVERGENCIA_VALOR"),
        ("DOC_CEDENTE", "123", "DOC_CEDENTE"),
        ("DT_EMISSAO_TITULO", date(2026, 9, 4), "posterior"),
        ("VL_PRESENTE", "=50+50", "fórmulas"),
        ("INCLUIR_CNAB", "", "INCLUIR_CNAB"),
        ("MOVIMENTO", 2, "MOVIMENTO"),
        ("VALOR_PAGO_TITULO", 1, "VALOR_PAGO_TITULO"),
        ("TAXA_INDEXADOR", 15, "TAXA_INDEXADOR"),
        ("VL_NOMINAL", "100000000000.00", "VL_NOMINAL"),
        ("NOME_CEDENTE", "NOME DIFERENTE", "DIVERGENCIA_NOME"),
    ],
)
def test_t15_t17_approval_never_releases_other_pending(prepared, tmp_path, field, value, message):
    edit(prepared, {field: value, "APROVADO": "NAO" if field == "NOME_CEDENTE" else "SIM"})
    with pytest.raises(ValidationError, match=message):
        generate_cnab(prepared, output_path=tmp_path / "blocked.txt")


def test_t13_t15_cession_requires_new_issue_and_principal_name(source_files, tmp_path):
    pfmi, analytic, due = source_files
    path = tmp_path / "cessao.xlsx"
    prepare_workbook(
        [PfmiInput(pfmi, CESSAO, "0")],
        analytic,
        due,
        liquidation_date="03/09/2026",
        first_sequence=1,
        output_path=path,
    )
    row = read_intermediate(path).rows[0]
    assert row.originals["ASSINATURA_ORIGINAL"].date() == date(2026, 9, 1)
    assert row.values["DT_EMISSAO_TITULO"] is None
    edit(path, {"DOC_CEDENTE": "11222333000181", "TIPO_PESSOA_CEDENTE": 2})
    with pytest.raises(ValidationError, match="DT_EMISSAO_TITULO"):
        generate_cnab(path, output_path=tmp_path / "blocked.txt")
    edit(
        path,
        {
            "DT_EMISSAO_TITULO": date(2026, 9, 2),
            "NOME_CEDENTE": "CONEXCRED INTERMEDIACAO (ORIGINAL)",
            "APROVADO": "SIM",
        },
    )
    with pytest.raises(ValidationError, match="deve ser somente"):
        generate_cnab(path, output_path=tmp_path / "blocked.txt")
    edit(path, {"NOME_CEDENTE": "CONEXCRED INTERMEDIACAO"})
    assert generate_cnab(path, output_path=tmp_path / "final.txt").detail_count == 1


def test_t11_base_conflict_can_be_corrected_in_excel(source_files, tmp_path):
    pfmi, analytic, due = source_files
    book = load_workbook(due)
    book.active.append(["FALENCIA ALFA SA", "52998224725", date(2031, 1, 1)])
    book.save(due)
    book.close()
    path = tmp_path / "conflito.xlsx"
    prepare_workbook(
        pfmi, analytic, due, liquidation_date="03/09/2026", first_sequence=1, output_path=path
    )
    assert read_intermediate(path).rows[0].values["STATUS"] == "CONFLITO_BASE_VENCIMENTOS"
    edit(
        path,
        {
            "DOC_SACADO": "11222333000181",
            "NOME_SACADO": "FALENCIA ALFA SA",
            "TIPO_PESSOA_SACADO": 2,
            "DT_VENCIMENTO": date(2030, 1, 31),
        },
    )
    assert generate_cnab(path, output_path=tmp_path / "final.txt").detail_count == 1


def test_t17_opposite_cents_do_not_compensate_between_groups(tmp_path):
    from .test_gui import _write_sources

    groups = [
        dict(
            cedent="A",
            failure="FALENCIA A",
            operation=100,
            nominal=150,
            acquisition=100,
            cession=False,
        ),
        dict(
            cedent="B",
            failure="FALENCIA B",
            operation=100,
            nominal=150,
            acquisition=100,
            cession=False,
        ),
    ]
    pfmis, analytic, due = _write_sources(tmp_path, groups)
    output = tmp_path / "opostos.xlsx"
    prepare_workbook(
        [PfmiInput(p) for p in pfmis],
        analytic,
        due,
        liquidation_date="03/09/2026",
        first_sequence=1,
        output_path=output,
    )
    edit(output, {"VL_PRESENTE": 100.01, "APROVADO": "SIM"}, row=2)
    edit(output, {"VL_PRESENTE": 99.99, "APROVADO": "SIM"}, row=3)
    with pytest.raises(ValidationError) as exc:
        generate_cnab(output, output_path=tmp_path / "blocked.txt")
    assert sum("Grupo" in issue and "DIVERGENCIA_VALOR" in issue for issue in exc.value.issues) == 2


def test_t09_manual_selection_cannot_reuse_credit_across_files(source_files, tmp_path):
    pfmi, analytic, due = source_files
    output = tmp_path / "duplicados.xlsx"
    prepare_workbook(
        [PfmiInput(pfmi), PfmiInput(pfmi)],
        analytic,
        due,
        liquidation_date="03/09/2026",
        first_sequence=1,
        output_path=output,
    )
    loaded = read_intermediate(output)
    assert len(loaded.rows) == 2 and all(r.values["INCLUIR_CNAB"] == "NAO" for r in loaded.rows)
    edit(output, {"INCLUIR_CNAB": "SIM", "APROVADO": "SIM"}, row=2)
    edit(output, {"INCLUIR_CNAB": "SIM", "APROVADO": "SIM"}, row=3)
    with pytest.raises(ValidationError, match="mesmo crédito"):
        generate_cnab(output, output_path=tmp_path / "blocked.txt")


def test_t12_placeholder_cannot_be_filled_into_credit(source_files, tmp_path):
    pfmi, analytic, due = source_files
    book = load_workbook(analytic)
    book.active.cell(2, 3).value = "OUTRA FALENCIA"
    book.save(analytic)
    book.close()
    output = tmp_path / "placeholder.xlsx"
    prepare_workbook(
        pfmi, analytic, due, liquidation_date="03/09/2026", first_sequence=1, output_path=output
    )
    edit(
        output,
        {
            "INCLUIR_CNAB": "SIM",
            "APROVADO": "SIM",
            "DOC_CEDENTE": "52998224725",
            "NOME_CEDENTE": "ACME LTDA",
            "VL_NOMINAL": 150,
            "VL_PRESENTE": 100,
            "DT_EMISSAO_TITULO": date(2026, 9, 1),
            "TIPO_PESSOA_CEDENTE": 1,
        },
    )
    with pytest.raises(ValidationError, match="sem crédito original"):
        generate_cnab(output, output_path=tmp_path / "blocked.txt")


def test_t18_removing_pending_rows_or_status_does_not_release(source_files, tmp_path):
    pfmi, analytic, due = source_files
    book = load_workbook(pfmi)
    book.active.append(["DESCONHECIDA"])
    book.save(pfmi)
    book.close()
    output = tmp_path / "pending.xlsx"
    prepare_workbook(
        pfmi, analytic, due, liquidation_date="03/09/2026", first_sequence=1, output_path=output
    )
    edit(output, {"STATUS": "OK", "PENDENCIAS": "", "APROVADO": "SIM"})
    with pytest.raises(ValidationError, match="PFMI"):
        generate_cnab(output, output_path=tmp_path / "blocked.txt")
    book = load_workbook(output)
    book["PENDENCIAS_PFMI"].delete_rows(2)
    book.save(output)
    book.close()
    with pytest.raises(ValidationError, match="PENDENCIAS_PFMI"):
        generate_cnab(output, output_path=tmp_path / "blocked.txt")


@pytest.mark.parametrize("length", [40, 41])
def test_t20_long_names_keep_full_excel_and_truncate_with_warning(prepared, tmp_path, length):
    name = "A" * length
    edit(prepared, {"NOME_CEDENTE": name, "APROVADO": "SIM"})
    result = generate_cnab(prepared, output_path=tmp_path / "long.txt")
    assert read_intermediate(prepared).rows[0].values["NOME_CEDENTE"] == name
    assert result.output_path.read_bytes().split(b"\r\n")[1][334:374] == b"A" * 40
    assert any("NOME_TRUNCADO_CNAB" in w for w in result.warnings) == (length == 41)


@pytest.mark.parametrize("name", ["\x7fNome", "Nome\u200b", "Nome😀", "---"])
def test_t20_incompatible_names_are_not_silently_replaced(prepared, tmp_path, name):
    edit(prepared, {"NOME_CEDENTE": name, "APROVADO": "SIM"})
    with pytest.raises(ValidationError, match="NOME_CEDENTE"):
        generate_cnab(prepared, output_path=tmp_path / "blocked.txt")


def test_t21_sequence_capacity_and_valid_upper_limit(prepared, tmp_path):
    edit(prepared, {}, summary={"B4": 9999999999})
    assert generate_cnab(prepared, output_path=tmp_path / "max.txt").first_sequence == 9999999999
    edit(prepared, {}, summary={"B4": 10000000000})
    with pytest.raises(ValidationError, match="Faixa de sequência"):
        generate_cnab(prepared, output_path=tmp_path / "blocked.txt")


def test_t22_excel_atomic_save_failure_cleans_temporary(source_files, tmp_path, monkeypatch):
    import gerador_cnab_nox.workbook as module

    path = tmp_path / "intermediario.xlsx"

    def failure(*args):
        raise PermissionError("synthetic denied")

    with monkeypatch.context() as m:
        m.setattr(module.os, "link", failure)
        with pytest.raises(PermissionError):
            prepare_workbook(*source_files, output_path=path)
    assert not path.exists() and not list(tmp_path.glob(".*.tmp"))


def test_t22_posix_directory_permission(source_files, tmp_path):
    import os

    if os.name != "posix":
        pytest.skip("Permissao chmod POSIX; validar ACL real do Windows no roteiro M7")
    folder = tmp_path / "readonly"
    folder.mkdir()
    folder.chmod(0o500)
    try:
        with pytest.raises(PermissionError):
            prepare_workbook(*source_files, output_path=folder / "blocked.xlsx")
    finally:
        folder.chmod(0o700)
    assert not list(folder.iterdir())


def test_t22_successive_names_are_exclusive(source_files, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    first = prepare_workbook(*source_files)
    second = prepare_workbook(*source_files)
    assert first.output_path != second.output_path
    assert first.output_path.exists() and second.output_path.exists()


def test_t23_approving_alternative_later_keeps_reserved_sequence_and_later_groups(tmp_path):
    from dataclasses import replace

    from gerador_cnab_nox.matching import prepare_batch
    from gerador_cnab_nox.models import PfmiData
    from gerador_cnab_nox.workbook import write_intermediate

    from .test_matching import _credit, _due, _payment

    px = replace(_payment("BX", "100"), failure="FALX", cedent_name="PESSOA X")
    cx = replace(_credit("X", "100"), campaign="FALX", cedent_name="PESSOA X")
    py = replace(_payment("BY", "100"), failure="FALY", cedent_name="PESSOA Y")
    cy = replace(_credit("Y", "100"), campaign="FALY", cedent_name="PESSOA Y DIFERENTE")
    due = [replace(_due(), debtor_name="FALX"), replace(_due(), debtor_name="FALY")]

    batch = prepare_batch(
        PfmiData(payments=[px, py]),
        [cx, cy],
        due,
        liquidation_date=date(2026, 9, 3),
        first_sequence=500,
    )
    path = tmp_path / "alternativas.xlsx"
    write_intermediate(batch, path)

    loaded = read_intermediate(path)
    x_row = next(r for r in loaded.rows if r.values["NOME_CEDENTE_PFMI"] == "PESSOA X")
    y_row = next(r for r in loaded.rows if r.values["NOME_CEDENTE_PFMI"] == "PESSOA Y")
    assert x_row.values["SEU_NUMERO"] == 500 and x_row.values["INCLUIR_CNAB"] == "SIM"
    assert y_row.values["SEU_NUMERO"] == 501 and y_row.values["INCLUIR_CNAB"] == "NAO"

    book = load_workbook(path)
    sh = book["CREDITOS"]
    headers = {c.value: c.column for c in sh[1]}
    y_line = next(
        r for r in range(2, sh.max_row + 1)
        if sh.cell(r, headers["ID_LINHA"]).value == y_row.values["ID_LINHA"]
    )
    # Approve Y without touching SEU_NUMERO/NU_DOCUMENTO at all.
    sh.cell(y_line, headers["INCLUIR_CNAB"]).value = "SIM"
    sh.cell(y_line, headers["APROVADO"]).value = "SIM"
    book.save(path)
    book.close()

    generated = generate_cnab(path, output_path=tmp_path / "alternativas.txt")
    assert generated.first_sequence == 500
    assert generated.last_sequence == 501
