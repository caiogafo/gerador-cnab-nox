from __future__ import annotations

import json
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from gerador_cnab_nox.errors import ValidationError
from gerador_cnab_nox.models import GenerationResult
from gerador_cnab_nox.reconciliation import reconcile_credit, summarize, tolerance_cents
from gerador_cnab_nox.service import generate_cnab, prepare_workbook
from gerador_cnab_nox.validation import validate_for_generation

from .test_gui import gui_application as gui_application
from .test_pfmi_composicoes import (
    PRINCIPAL,
    REPRESENTANTE,
    _composition_divergence_edits,
    _credit,
    _pay,
    _prepare,
)
from .test_validation import _loaded, _prepared


@pytest.mark.parametrize(("difference", "status"), [
    ("0", "PERFECT_MATCH"), ("0.01", "TOLERATED_WARNING"),
    ("-0.01", "TOLERATED_WARNING"), ("100", "TOLERATED_WARNING"),
    ("-100", "TOLERATED_WARNING"), ("100.01", "CRITICAL_ERROR"),
    ("-100.01", "CRITICAL_ERROR"),
])
def test_default_inclusive_absolute_limit(monkeypatch, difference, status):
    monkeypatch.delenv("MAX_TOLERANCE_DIFF_REAIS", raising=False)
    record = reconcile_credit(
        "credit", "529.982.247-25", "Credor", Decimal("1000") + Decimal(difference),
        Decimal("1000"),
    )
    assert record.status == status
    assert record.diffCents == int(Decimal(difference) * 100)
    assert record.documentNumber == "52998224725"
    assert summarize([record]).canGenerateTxt == (status != "CRITICAL_ERROR")


@pytest.mark.parametrize("value", ["-1", "NaN", "Infinity", "", "abc", "0.001", "1,00"])
def test_invalid_config_blocks(monkeypatch, value):
    monkeypatch.setenv("MAX_TOLERANCE_DIFF_REAIS", value)
    with pytest.raises(ValidationError, match="MAX_TOLERANCE_DIFF_REAIS"):
        validate_for_generation(_loaded([_prepared("A", "G")]))


def test_config_zero_and_custom_limit(monkeypatch):
    for value, expected in [("0", 0), ("0.05", 5), ("123.45", 12345)]:
        monkeypatch.setenv("MAX_TOLERANCE_DIFF_REAIS", value)
        assert tolerance_cents() == expected
    record = reconcile_credit("a", "", "", Decimal("1.01"), Decimal(1), limit_cents=0)
    assert record.status == "CRITICAL_ERROR"


def _divergent_row(pfmi="100.00", analytic="99.99", group="G", line="A"):
    row = _prepared(line, group, analytic_line=line, reference=line)
    for data in (row.values, row.originals):
        data["TOTAL_PFMI_GRUPO"] = Decimal(pfmi)
        data["VL_PRESENTE"] = Decimal(analytic)
        data["SEU_NUMERO"] = ""
    return row


def test_validation_uses_protected_sources_and_pfmi_value(monkeypatch):
    monkeypatch.setenv("MAX_TOLERANCE_DIFF_REAIS", "100")
    row = _divergent_row()
    row.values["VL_PRESENTE"] = Decimal("999.00")
    batch = validate_for_generation(_loaded([row]))
    assert batch.rows[0]["VL_PRESENTE"] == Decimal("100.00")
    summary = batch.reconciliation.to_dict()
    assert summary["canGenerateTxt"]
    assert summary["totalPfmiCents"] == 10000
    assert summary["totalAnalyticCents"] == 9999
    assert summary["totalDiffCents"] == 1
    assert summary["warningCount"] == 1
    assert summary["warnings"][0]["status"] == "TOLERATED_WARNING"
    assert row.originals["VL_PRESENTE"] == Decimal("99.99")
    assert row.values["VL_PRESENTE"] == Decimal("999.00")


def test_critical_not_offset_by_other_credit_or_manual_approval(monkeypatch):
    monkeypatch.setenv("MAX_TOLERANCE_DIFF_REAIS", "100")
    rows = [
        _divergent_row("200.01", "100", "G1", "1"),
        _divergent_row("100", "200.01", "G2", "2"),
    ]
    for row in rows:
        row.values["DIVERGENCIA_VALOR_APROVADA"] = "SIM"
        row.values["DIVERGENCIA_VALOR_JUSTIFICATIVA"] = "Revisado"
        row.values["VL_PRESENTE"] = row.originals["TOTAL_PFMI_GRUPO"]
    with pytest.raises(ValidationError) as error:
        validate_for_generation(_loaded(rows))
    summary = error.value.reconciliation
    assert not summary.canGenerateTxt
    assert summary.totalDiffCents == 0
    assert len(summary.errors) == 2


def test_tolerance_does_not_clear_document_errors(monkeypatch):
    monkeypatch.setenv("MAX_TOLERANCE_DIFF_REAIS", "100")
    row = _divergent_row()
    row.values["DOC_CEDENTE"] = "123"
    with pytest.raises(ValidationError) as error:
        validate_for_generation(_loaded([row]))
    assert not error.value.reconciliation.canGenerateTxt
    assert error.value.reconciliation.hasWarnings


@pytest.mark.parametrize("approved", [True, False])
def test_composition_counts_shared_credit_once_and_requires_approval(monkeypatch, approved):
    monkeypatch.setenv("MAX_TOLERANCE_DIFF_REAIS", "100")
    batch = _prepare(
        [_pay(PRINCIPAL, "60.00"), _pay(REPRESENTANTE, "39.99")],
        [_credit("REF-1", PRINCIPAL, acquisition="100.00")],
    )
    edits = _composition_divergence_edits(
        batch.rows, VL_NOMINAL=250, DIVERGENCIA_VALOR_APROVADA="",
        DIVERGENCIA_VALOR_JUSTIFICATIVA="", COMPOSICAO_APROVADA="SIM" if approved else "",
    )
    for row in batch.rows:
        row.values.update(edits[row.values["ID_LINHA"]])
    loaded = _loaded(batch.rows)
    if approved:
        result = validate_for_generation(loaded)
        summary = result.reconciliation
        assert [r["VL_PRESENTE"] for r in result.rows] == [Decimal("60"), Decimal("39.99")]
    else:
        with pytest.raises(ValidationError) as error:
            validate_for_generation(loaded)
        summary = error.value.reconciliation
        assert not summary.canGenerateTxt
        assert any("COMPOSICAO_APROVADA" in issue for issue in error.value.issues)
    assert summary.totalAnalyticCents == 10000
    assert summary.totalPfmiCents == 9999
    assert summary.warningCount == 1


@pytest.mark.parametrize("analytic_value", ["99.99", "200.00", "200.01"])
def test_excel_to_real_txt_and_structured_audit(
    source_files, tmp_path, monkeypatch, analytic_value,
):
    monkeypatch.setenv("MAX_TOLERANCE_DIFF_REAIS", "100.00")
    pfmi, analytic, due = source_files
    wb = load_workbook(analytic)
    wb.active["E2"] = Decimal(analytic_value)
    wb.save(analytic)
    wb.close()
    intermediate = tmp_path / "review.xlsx"
    prepare_workbook(
        pfmi, analytic, due, liquidation_date=date(2026, 9, 3), first_sequence=1,
        output_path=intermediate,
    )
    target = tmp_path / "payments.txt"
    if analytic_value == "200.01":
        # Explicitly selecting a critical row must not bypass reconciliation.
        wb = load_workbook(intermediate)
        sheet = wb["CREDITOS"]
        headers = {c.value: c.column for c in sheet[1]}
        sheet.cell(2, headers["INCLUIR_CNAB"], "SIM")
        wb.save(intermediate)
        wb.close()
        with pytest.raises(ValidationError) as error:
            generate_cnab(intermediate, output_path=target)
        assert error.value.reconciliation.errors
        assert not target.exists()
    else:
        result = generate_cnab(intermediate, output_path=target)
        assert result.total_present == Decimal("100.00")
        assert result.reconciliation.warningCount == 1
        assert result.warnings
        assert target.read_bytes().splitlines()[1][192:205] == b"0000000010000"
    entries = [json.loads(line) for line in
               (tmp_path / "logs" / "gerador_cnab_nox.jsonl").read_text().splitlines()]
    summary = entries[-1]["reconciliation"]
    assert summary["totalPfmiCents"] == 10000
    assert summary["totalAnalyticCents"] == int(Decimal(analytic_value) * 100)
    assert summary["canGenerateTxt"] == (analytic_value != "200.01")


def test_gui_shows_tolerated_warning(gui_application, tmp_path):
    _, app, _ = gui_application
    record = reconcile_credit(
        "G1", "52998224725", "Credor", Decimal("100"), Decimal("99.99"), limit_cents=10000,
    )
    result = GenerationResult(
        tmp_path / "synthetic.txt", 1, Decimal("150"), Decimal("100"), 1, 1, 1,
        (record.warningMessage,), reconciliation=summarize([record]),
    )
    app._generated(result)
    assert app.generation_result.kind == "warning"
    assert "TOLERATED_WARNING" in app.generation_result.issue_text.get("1.0", "end")
    assert "PFMI" in app.generation_result.issue_text.get("1.0", "end")
