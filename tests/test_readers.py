from __future__ import annotations

from pathlib import Path

import xlwt
from openpyxl import Workbook, load_workbook

from gerador_cnab_nox.readers import read_analytic, read_due_base, read_pfmi

from .conftest import VALID_CPF


def test_pfmi_preserves_blocks_and_unknown_rows(tmp_path: Path, source_files) -> None:
    pfmi, _, _ = source_files
    workbook = load_workbook(pfmi)
    sheet = workbook.active
    sheet.append([])
    sheet.append(["LINHA OPERACIONAL DESCONHECIDA", "X"])
    sheet.append([])
    sheet.append(["FUNDO SINTETICO NOX - BLOCO 2"])
    sheet.append(
        [
            "VALOR DA OPERAÇÃO",
            "FALÊNCIA",
            "CREDOR / CEDENTE",
            "BENEFICIÁRIOS / FAVORECIDOS",
            "BCO",
            "AG",
            "CTA",
            "CPF/ CNPJ",
            "VALOR",
        ]
    )
    sheet.append([25, "FALENCIA ALFA SA", "ACME LTDA", "OUTRO", "", "", "", VALID_CPF, 25])
    workbook.save(pfmi)
    workbook.close()
    data = read_pfmi(pfmi)
    assert [payment.block_id for payment in data.payments] == ["B0001", "B0001", "B0002"]
    assert len(data.unknown_rows) == 1
    assert data.unknown_rows[0].source_row > 0


def test_pfmi_summary_rows_are_visible_and_fund_title_is_kept(tmp_path: Path) -> None:
    path = tmp_path / "pfmi.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["DATA", "QUANTIDADE DE TED´S", "VALOR TOTAL A PAGAR"])
    sheet.append(["03/09/2026", 2, 15])
    sheet.append([])
    sheet.append(["FIDC NOX BLOCO SINTETICO"])
    headers = [
        "VALOR DA OPERAÇÃO",
        "FALÊNCIA",
        "CREDOR / CEDENTE",
        "BENEFICIÁRIOS / FAVORECIDOS",
        "CPF/CNPJ",
        "VALOR",
    ]
    sheet.append(headers)
    sheet.append([10, "FALENCIA A", "ACME LTDA", "BENEFICIARIO", VALID_CPF, 10])
    sheet.append(headers)  # repetição de paginação, não um novo bloco
    sheet.append([5, "FALENCIA A", "ACME LTDA", "BENEFICIARIO 2", VALID_CPF, 5])
    workbook.save(path)
    workbook.close()
    data = read_pfmi(path)
    assert data.payments[0].title == "FIDC NOX BLOCO SINTETICO"
    assert [payment.block_id for payment in data.payments] == ["B0001", "B0001"]
    assert not data.unknown_rows
    assert any(e.get("alerta") == "RESUMO_CONFERIDO" for e in data.evidence)


def test_pfmi_title_detection_does_not_hide_substring_match(tmp_path: Path) -> None:
    path = tmp_path / "pfmi.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["NOXIOUS PAYMENT", "X"])
    sheet.append(["FUNDO SINTETICO NOX"])
    sheet.append(
        [
            "VALOR DA OPERAÇÃO",
            "FALÊNCIA",
            "CREDOR / CEDENTE",
            "BENEFICIÁRIOS / FAVORECIDOS",
            "CPF/CNPJ",
            "VALOR",
        ]
    )
    sheet.append([10, "FALENCIA A", "ACME LTDA", "BENEFICIARIO", VALID_CPF, 10])
    workbook.save(path)
    workbook.close()

    data = read_pfmi(path)

    assert len(data.unknown_rows) == 1
    assert data.unknown_rows[0].source_row == 1


def test_analytic_reads_xlsx_csv_and_legacy_xls(tmp_path: Path) -> None:
    headers = ["Cpf", "Contato", "Campanha", "Valor Receber", "Valor Aquisicao", "Data Assinatura"]
    values = [52998224725.0, "ACME LTDA", "FALENCIA A", 20, 10, "01/09/2026"]

    xlsx = tmp_path / "a.xlsx"
    workbook = Workbook()
    workbook.active.append(headers)
    workbook.active.append(values)
    workbook.save(xlsx)
    workbook.close()

    csv_path = tmp_path / "a.csv"
    csv_path.write_text(";".join(headers) + "\n" + ";".join(map(str, values)), encoding="cp1252")

    xls = tmp_path / "a.xls"
    legacy = xlwt.Workbook()
    sheet = legacy.add_sheet("Analitico")
    for column, value in enumerate(headers):
        sheet.write(0, column, value)
    for column, value in enumerate(values):
        sheet.write(1, column, value)
    legacy.save(str(xls))

    for path in (xlsx, csv_path, xls):
        credits = read_analytic(path)
        assert len(credits) == 1
        assert credits[0].cedent_document == VALID_CPF
        assert credits[0].acquisition_value is not None


def test_due_base_does_not_silently_drop_invalid_date(tmp_path: Path) -> None:
    path = tmp_path / "v.xlsx"
    workbook = Workbook()
    workbook.active.append(["NOME_SACADO", "DOC_SACADO", "DATA_VENCIMENTO"])
    workbook.active.append(["SACADO", VALID_CPF, "invalida"])
    workbook.save(path)
    workbook.close()
    records = read_due_base(path)
    assert records[0].due_date is None


def test_extreme_operational_values_become_pending_instead_of_aborting(
    source_files,
) -> None:
    pfmi, analytic, due = source_files

    workbook = load_workbook(pfmi)
    workbook.active.cell(3, 1).value = "1e999999"
    workbook.save(pfmi)
    workbook.close()
    pfmi_data = read_pfmi(pfmi)
    assert any(row.reason == "VALOR DA OPERAÇÃO inválido" for row in pfmi_data.unknown_rows)

    workbook = load_workbook(analytic)
    workbook.active.cell(2, 4).value = "1e999999"
    workbook.active.cell(2, 5).value = "1e999999"
    workbook.active.cell(2, 6).value = 999_999_999
    workbook.active.cell(2, 6).number_format = "General"
    workbook.save(analytic)
    workbook.close()
    credit = read_analytic(analytic)[0]
    assert credit.nominal_value is None
    assert credit.acquisition_value is None
    assert credit.signature_date is None

    workbook = load_workbook(due)
    workbook.active.cell(2, 3).value = 999_999_999
    workbook.active.cell(2, 3).number_format = "General"
    workbook.save(due)
    workbook.close()
    assert read_due_base(due)[0].due_date is None


def _pfmi_book(path, tables):
    from .conftest import VALID_CPF

    book = Workbook()
    book.remove(book.active)
    headers = [
        "VALOR DA OPERAÇÃO",
        "FALÊNCIA",
        "CREDOR / CEDENTE",
        "BENEFICIÁRIOS / FAVORECIDOS",
        "CPF/CNPJ",
        "VALOR",
    ]
    for name, blocks in tables:
        sh = book.create_sheet(name)
        for block in blocks:
            sh.append(headers)
            for value, cedent in block:
                sh.append([value, "MOGIANO", cedent, "BENEFICIARIO", VALID_CPF, value])
            sh.append([])
    book.save(path)
    book.close()


def test_t05_t07_semantic_mirror_requires_order_and_group_partition(tmp_path):
    path = tmp_path / "mirror.xlsx"
    _pfmi_book(
        path,
        [
            ("renomeada", [[(10, "A"), (10, "A")], [(20, "B")]]),
            ("copia", [[(10, "A"), (10, "A"), (20, "B")]]),
        ],
    )
    data = read_pfmi(path)
    assert len(data.payments) == 3
    mirrors = [e for e in data.evidence if e["tipo"] == "ESPELHO_INTEGRAL"]
    assert len(mirrors) == 1 and mirrors[0]["pagamentos_nao_duplicados"] == 3
    assert len({p.block_id for p in data.payments}) == 2
    # Equal totals and content with reordered rows are insufficient.
    _pfmi_book(path, [("a", [[(10, "A"), (20, "B")]]), ("b", [[(20, "B"), (10, "A")]])])
    assert len(read_pfmi(path).payments) == 4
    # Equal consecutive credits in two physical blocks must not collapse in a flat sheet.
    _pfmi_book(path, [("a", [[(10, "A")], [(10, "A")]]), ("b", [[(10, "A"), (10, "A")]])])
    assert len(read_pfmi(path).payments) == 4
    # Partial subset is also not a mirror.
    _pfmi_book(path, [("a", [[(10, "A"), (20, "B")]]), ("b", [[(10, "A")]])])
    assert len(read_pfmi(path).payments) == 3


def test_t07_duplicate_files_preserved_and_signaled(tmp_path):
    from gerador_cnab_nox.models import PfmiInput
    from gerador_cnab_nox.readers import read_pfmis

    path = tmp_path / "p.xlsx"
    _pfmi_book(path, [("a", [[(10, "A")]])])
    data = read_pfmis([PfmiInput(path), PfmiInput(path)])
    assert len(data.payments) == 2
    assert [p.file_order for p in data.payments] == [1, 2]
    assert any(e["tipo"] == "ARQUIVO_OPERACIONAL_DUPLICADO" for e in data.evidence)


def test_t06_auxiliary_summary_empty_stale_legacy_and_unknown(tmp_path):
    path = tmp_path / "p.xlsx"
    _pfmi_book(path, [("a", [[(10, "A")]])])
    book = load_workbook(path)
    aux = book.create_sheet("qualquer_nome")
    aux.append(["DATA", "QUANTIDADE DE TED´S", "VALOR TOTAL A PAGAR"])
    aux.append([None, 99, 999])
    aux.append([])
    aux.append(["FALÊNCIA / SACADO", "CREDOR / CEDENTE", "VALOR"])
    aux.append(["MOGIANO", "A", 999])
    aux.append([None, None, "=SUM(C5:C5)"])
    aux.append([])
    aux.append(
        ["DOC_CEDENTE", "NOME_CEDENTE", "VL_NOMINAL", "VL_PRESENTE", "SEU_NUMERO", "NU_DOCUMENTO"]
    )
    aux.append([VALID_CPF, "A", 10, 10, 1, 1])
    aux.append([])
    aux.append(["NAO RECONHECIDA", "CONFERIR"])
    book.save(path)
    book.close()
    data = read_pfmi(path)
    assert len(data.payments) == 1
    assert len(data.unknown_rows) == 1
    assert data.unknown_rows[0].source_row == 11
    assert any(e.get("alerta") == "RESUMO_DIVERGENTE" for e in data.evidence)
    assert any(e["tipo"] == "LINHA_GERADOR_AUXILIAR" for e in data.evidence)


def test_t05_continuation_inherits_only_inside_block(source_files):
    pfmi, _, _ = source_files
    book = load_workbook(pfmi)
    sh = book.active
    sh.cell(4, 2).value = None
    sh.cell(4, 3).value = None
    sh.append([])
    sh.append([5, None, None, "OUTRO", "", "", "", VALID_CPF, 5])
    book.save(pfmi)
    book.close()
    data = read_pfmi(pfmi)
    assert len(data.payments) == 2
    assert data.payments[1].cedent_name == data.payments[0].cedent_name
    assert len(data.unknown_rows) == 1


def test_t10_two_contatos_references_and_multiline_csv(tmp_path):
    import csv

    headers = [
        "Prospect",
        "Cpf",
        "Contato",
        "Campanha",
        "Valor Receber",
        "Valor Aquisicao",
        "Data Assinatura",
        "Contato",
        "Id",
    ]
    rows = [
        ["REPETIDO", VALID_CPF, "PRINCIPAL", "MOGIANO", 20, 10, "01/09/2026", "SECUNDARIO", "ID-1"],
        [
            "REPETIDO",
            VALID_CPF,
            "PRINCIPAL\nDOIS",
            "MOGIANO",
            20,
            10,
            "01/09/2026",
            "SECUNDARIO",
            "ID-1",
        ],
    ]
    for extension in ("xlsx", "xls", "csv"):
        path = tmp_path / f"a.{extension}"
        if extension == "csv":
            with path.open("w", newline="", encoding="utf-8-sig") as f:
                csv.writer(f, delimiter="|").writerows([headers, *rows])
        elif extension == "xlsx":
            book = Workbook()
            for row in [headers, *rows]:
                book.active.append(row)
            book.save(path)
            book.close()
        else:
            book = xlwt.Workbook()
            sh = book.add_sheet("Origem")
            for r, row in enumerate([headers, *rows]):
                for c, value in enumerate(row):
                    sh.write(r, c, value)
            book.save(str(path))
        credits = read_analytic(path)
        assert [c.cedent_name for c in credits] == ["PRINCIPAL", "PRINCIPAL\nDOIS"]
        assert [c.reference for c in credits] == ["ID-1", "ID-1"]
        assert credits[0].credit_id != credits[1].credit_id
        assert credits[0].source_hash and credits[1].source_row == 3


def test_t06_known_headerless_projections_are_preserved_not_payments(tmp_path):
    path = tmp_path / "aux.xlsx"
    _pfmi_book(path, [("fonte", [[(10, "A"), (20, "B")]])])
    book = load_workbook(path)
    full = book.create_sheet("aux_12")
    partial = book.create_sheet("aux_9")
    dated = book.create_sheet("aux_5")
    dated.append([None, "DATA", None, None, None])
    dated.append([None, "01/09/2026", None, None, None])
    for i, (amount, name) in enumerate([(10, "A"), (20, "B")], 1):
        full.append(
            [
                None,
                amount,
                "MOGIANO",
                name,
                "BENEFICIARIO",
                "001",
                "0001",
                "0001",
                VALID_CPF,
                amount,
                None,
                f"=J{i}*-1",
            ]
        )
        partial.append(
            ["MOGIANO", name, "BENEFICIARIO", "001", "0001", "0001", VALID_CPF, amount, f"=H{i}*-1"]
        )
        dated.append(["MOGIANO", name, VALID_CPF, amount, None])
    book.save(path)
    book.close()
    data = read_pfmi(path)
    assert len(data.payments) == 2 and not data.unknown_rows
    assert sum(e["tipo"].startswith("PROJECAO_AUXILIAR") for e in data.evidence) == 3
    assert any("=J1*-1" in e.get("original", []) for e in data.evidence)
    # One extra unknown row invalidates the complete shape, instead of disappearing.
    book = load_workbook(path)
    book["aux_12"].append(["DESCONHECIDA"])
    book.save(path)
    book.close()
    assert read_pfmi(path).unknown_rows
