from dataclasses import replace
from datetime import date

import pytest

from gerador_cnab_nox.errors import InputFileError
from gerador_cnab_nox.failure_aliases import load_aliases, resolver, save_aliases, validate_aliases
from gerador_cnab_nox.matching import prepare_batch
from gerador_cnab_nox.models import PfmiData
from gerador_cnab_nox.service import generate_cnab, prepare_workbook
from gerador_cnab_nox.workbook import read_intermediate

from .test_matching import _credit, _due, _payment


def test_persistence_chains_and_removal(tmp_path):
    path = tmp_path / "aliases.json"
    assert load_aliases(path) == {}
    save_aliases({"apelido": "outro", "outro": "referência"}, path)
    assert resolver(load_aliases(path))("Apelido") == "REFERENCIA"
    save_aliases({}, path)
    assert load_aliases(path) == {}


def test_builtin_equivalence_can_be_saved_and_does_not_block_other_entries(tmp_path):
    path = tmp_path / "aliases.json"
    values = {"  keleti  ": "KELETI ENGENHARIA", "MOGIANO": "MOGIANO TRANSP GERAIS"}
    save_aliases(values, path)
    loaded = load_aliases(path)
    assert loaded["KELETI"] == "KELETI ENGENHARIA"
    save_aliases({**loaded, "OUTRA": "OUTRA ENGENHARIA"}, path)
    resolve = resolver(load_aliases(path))
    assert resolve("keleti") == "KELETI ENGENHARIA"
    assert resolve("MOGIANO") == "MOGIANO TRANSP GERAIS"
    assert resolve("OUTRA") == "OUTRA ENGENHARIA"


@pytest.mark.parametrize("values", [
    {"KELETI": "OUTRA", "KELETI ENGENHARIA": "TERCEIRA"},
    {"KELETI": "OUTRA", "OUTRA": "KELETI ENGENHARIA"},
])
def test_builtin_canonical_conflicts_and_cycles_remain_blocked(values):
    with pytest.raises(InputFileError):
        validate_aliases(values)


@pytest.mark.parametrize("values", [
    {"A": "B", "B": "A"}, {"A": "A"}, {"": "B"}, {"A": None},
    {"Á": "B", "a": "C"}, [],
])
def test_invalid_configuration_rejected(values):
    with pytest.raises(InputFileError):
        validate_aliases(values)


def test_corrupt_file_does_not_silently_disable_rules(tmp_path):
    path = tmp_path / "aliases.json"
    path.write_text("{broken", encoding="utf-8")
    with pytest.raises(InputFileError):
        load_aliases(path)


def test_aliases_resolve_all_sources_without_modifying_originals():
    payment = replace(_payment("B1", "100"), failure="APELIDO PFMI")
    credit = replace(_credit("1", "100"), campaign="APELIDO SISTEMA")
    due = replace(_due(), debtor_name="NOME BASE")
    aliases = {"APELIDO PFMI": "NOME BASE", "APELIDO SISTEMA": "NOME BASE"}
    def prepare(rules):
        return prepare_batch(PfmiData(payments=[payment]), [credit], [due],
                             liquidation_date=date(2026, 9, 10), first_sequence=1,
                             failure_aliases=rules)
    assert prepare({}).rows[0].values["INCLUIR_CNAB"] == "NAO"
    result = prepare(aliases).rows[0].values
    assert result["INCLUIR_CNAB"] == "SIM"
    assert result["FALENCIA_PFMI"] == "APELIDO PFMI"
    assert result["NOME_SACADO"] == "NOME BASE"
    assert result["FALENCIA_CANONICA"] == "NOME BASE"
    assert "SACADO_NAO_LOCALIZADO" not in result["PENDENCIAS"]
    assert credit.campaign == "APELIDO SISTEMA"
    assert prepare({}).rows[0].values["INCLUIR_CNAB"] == "NAO"


def test_aliases_do_not_merge_conflicting_due_documents():
    payment = _payment("B1", "100")
    aliases = {"OUTRO NOME": payment.failure}
    due = replace(_due(), debtor_name="OUTRO NOME", debtor_document="123")
    result = prepare_batch(PfmiData(payments=[payment]), [_credit("1", "100")],
                           [_due(), due], liquidation_date=date(2026, 9, 10),
                           first_sequence=1, failure_aliases=aliases)
    assert "CONFLITO_BASE_VENCIMENTOS" in result.rows[0].values["PENDENCIAS"]


def test_generation_uses_captured_rules_not_local_file(source_files, tmp_path):
    from openpyxl import load_workbook
    pfmi, analytic, base = source_files
    book = load_workbook(pfmi)
    for row in (3, 4):
        book.active.cell(row, 2, "APELIDO")
    book.save(pfmi)
    book.close()
    output = tmp_path / "intermediate.xlsx"
    prepare_workbook(pfmi, analytic, base, liquidation_date="03/09/2026",
                     first_sequence=1, output_path=output,
                     failure_aliases={"APELIDO": "FALENCIA ALFA SA"})
    read_intermediate(output)
    pfmi.unlink()
    analytic.unlink()
    base.unlink()
    result = generate_cnab(output, output_path=tmp_path / "out.txt")
    assert result.detail_count == 1
