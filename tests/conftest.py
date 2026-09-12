from __future__ import annotations

import os
from datetime import date
from pathlib import Path
from tkinter import TclError, Tk

import pytest
from openpyxl import Workbook

VALID_CPF = "52998224725"
VALID_CNPJ = "11222333000181"


@pytest.fixture(scope="session")
def _shared_tk_root():
    # One interpreter for the whole GUI suite: creating and destroying a Tk()
    # interpreter per test (hundreds of them) exhausts Windows GDI/USER
    # handles well before any single test leaks anything. Widgets are torn
    # down per test by the gui_application fixture; the interpreter itself
    # is not recreated.
    try:
        root = Tk()
    except TclError as exc:
        if os.environ.get("CNAB_NOX_REQUIRE_GUI") == "1":
            pytest.fail(f"Tk obrigatório indisponível: {exc}")
        pytest.skip(f"Tk indisponível nesta sessão: {exc}")
    root.withdraw()
    yield root
    root.destroy()


@pytest.fixture
def source_files(tmp_path: Path) -> tuple[Path, Path, Path]:
    pfmi = tmp_path / "pfmi_sintetica.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["FUNDO SINTETICO NOX"])
    sheet.append(
        [
            "VALOR DA OPERAÇÃO",
            "FALÊNCIA",
            "CREDOR / CEDENTE",
            "BENEFICIÁRIOS / FAVORECIDOS",
            "BCO",
            "AG",
            "CTA",
            "CPF/ CNPJ",
            "VALOR",
            "PIX",
        ]
    )
    sheet.append(
        [
            40,
            "FALENCIA ALFA SA",
            "ACME LTDA",
            "BENEFICIARIO UM",
            "",
            "",
            "",
            VALID_CPF,
            40,
        ]
    )
    sheet.append(
        [
            60,
            "FALENCIA ALFA SA",
            "ACME LTDA",
            "BENEFICIARIO DOIS",
            "",
            "",
            "",
            VALID_CPF,
            60,
        ]
    )
    workbook.save(pfmi)
    workbook.close()

    analytic = tmp_path / "analitico_sintetico.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "Cpf",
            "Contato",
            "Campanha",
            "Valor Receber",
            "Valor Aquisicao",
            "Data Assinatura",
            "ID",
        ]
    )
    sheet.append(
        [
            VALID_CPF,
            "ACME LIMITADA",
            "FALENCIA ALFA SOCIEDADE ANONIMA",
            150,
            100,
            date(2026, 9, 1),
            "P-1",
        ]
    )
    workbook.save(analytic)
    workbook.close()

    due = tmp_path / "vencimentos_sintetico.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["NOME_SACADO", "DOC_SACADO", "DATA_VENCIMENTO"])
    sheet.append(["FALENCIA ALFA S/A", VALID_CNPJ, date(2030, 1, 31)])
    workbook.save(due)
    workbook.close()
    return pfmi, analytic, due
