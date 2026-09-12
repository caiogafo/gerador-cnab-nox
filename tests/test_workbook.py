from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from gerador_cnab_nox.models import PreparedBatch
from gerador_cnab_nox.workbook import CREDIT_HEADERS, read_intermediate, write_intermediate

from .test_validation import _prepared


def test_formula_like_source_text_is_escaped_without_false_integrity_error(
    tmp_path: Path,
) -> None:
    row = _prepared("A", "G1")
    dangerous = '=HYPERLINK("https://invalid.example","X")'
    row.values["NOME_CEDENTE"] = dangerous
    row.originals["NOME_CEDENTE"] = dangerous
    row.originals["NOME_CEDENTE_PFMI"] = dangerous
    row.values["NOME_CEDENTE_PFMI"] = dangerous
    output = tmp_path / "seguro.xlsx"
    write_intermediate(PreparedBatch([row], [], date(2026, 9, 3), 1), output)

    workbook = load_workbook(output, data_only=False)
    headers = {cell.value: cell.column for cell in workbook["CREDITOS"][1]}
    stored = workbook["CREDITOS"].cell(2, headers["NOME_CEDENTE"]).value
    workbook.close()
    assert stored.startswith("'=")
    assert not stored.startswith("=")
    loaded = read_intermediate(output)
    assert not loaded.structural_issues
    assert set(CREDIT_HEADERS).issubset(loaded.rows[0].values)
