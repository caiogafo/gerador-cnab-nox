from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from gerador_cnab_nox.errors import ValidationError
from gerador_cnab_nox.normalize import header_key
from gerador_cnab_nox.service import generate_cnab, prepare_workbook
from gerador_cnab_nox.workbook import read_intermediate

from .conftest import VALID_CNPJ, VALID_CPF


def test_prepare_then_generate_uses_only_intermediate(tmp_path: Path, source_files) -> None:
    pfmi, analytic, due = source_files
    excel = tmp_path / "intermediario.xlsx"
    logs = tmp_path / "logs"
    result = prepare_workbook(
        pfmi,
        analytic,
        due,
        liquidation_date="03/09/2026",
        first_sequence="500",
        output_path=excel,
        log_directory=logs,
    )
    assert result.selected_count == 1
    workbook = load_workbook(excel)
    assert workbook.sheetnames == [
        "RESUMO",
        "CREDITOS",
        "PENDENCIAS_PFMI",
        "ORIGINAIS_CONTROLE",
        "MANIFESTO",
    ]
    assert workbook["ORIGINAIS_CONTROLE"].sheet_state == "hidden"
    assert workbook["MANIFESTO"].sheet_state == "hidden"
    workbook.close()

    # Prova de que o segundo passo não relê nenhuma das três fontes.
    pfmi.unlink()
    analytic.unlink()
    due.unlink()
    txt = tmp_path / "resultado.txt"
    generated = generate_cnab(excel, output_path=txt, log_directory=logs)
    assert generated.detail_count == 1
    assert generated.first_sequence == 500
    assert all(len(line) == 444 for line in txt.read_bytes().split(b"\r\n")[:-1])

    log_text = (logs / "gerador_cnab_nox.jsonl").read_text(encoding="utf-8")
    assert VALID_CPF not in log_text
    assert VALID_CNPJ not in log_text
    assert "ACME" not in log_text
    for line in log_text.splitlines():
        json.loads(line)


def test_log_failure_is_a_nonfatal_alert_after_outputs_are_created(
    tmp_path: Path, source_files
) -> None:
    invalid_log_directory = tmp_path / "arquivo_em_vez_de_diretorio"
    invalid_log_directory.write_text("x", encoding="utf-8")
    excel = tmp_path / "intermediario.xlsx"
    prepared = prepare_workbook(
        *source_files,
        liquidation_date="03/09/2026",
        first_sequence=1,
        output_path=excel,
        log_directory=invalid_log_directory,
    )
    assert excel.exists()
    assert any("LOG_AUDITORIA_NAO_GRAVADO" in item for item in prepared.warnings)

    txt = tmp_path / "resultado.txt"
    generated = generate_cnab(
        excel,
        output_path=txt,
        log_directory=invalid_log_directory,
    )
    assert txt.exists()
    assert any("LOG_AUDITORIA_NAO_GRAVADO" in item for item in generated.warnings)


def test_one_cent_edit_blocks_txt(tmp_path: Path, source_files) -> None:
    excel = tmp_path / "intermediario.xlsx"
    prepare_workbook(*source_files, output_path=excel)
    workbook = load_workbook(excel)
    sheet = workbook["CREDITOS"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers["VL_PRESENTE"]).value = 99.99
    workbook["RESUMO"]["B3"] = "03/09/2026"
    workbook["RESUMO"]["B4"] = 1
    workbook.save(excel)
    workbook.close()
    with pytest.raises(ValidationError, match="DIVERGENCIA_VALOR"):
        generate_cnab(excel, output_path=tmp_path / "nao_gerar.txt")
    assert not (tmp_path / "nao_gerar.txt").exists()


def test_name_divergence_needs_per_row_approval(tmp_path: Path, source_files) -> None:
    excel = tmp_path / "intermediario.xlsx"
    prepare_workbook(
        *source_files,
        liquidation_date="03/09/2026",
        first_sequence=1,
        output_path=excel,
    )
    workbook = load_workbook(excel)
    sheet = workbook["CREDITOS"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers["NOME_CEDENTE"]).value = "NOME DIFERENTE"
    sheet.cell(2, headers["APROVADO"]).value = ""
    workbook.save(excel)
    workbook.close()
    with pytest.raises(ValidationError, match="DIVERGENCIA_NOME"):
        generate_cnab(excel, output_path=tmp_path / "bloqueado.txt")

    workbook = load_workbook(excel)
    workbook["CREDITOS"].cell(2, headers["APROVADO"]).value = "SIM"
    workbook.save(excel)
    workbook.close()
    result = generate_cnab(excel, output_path=tmp_path / "aprovado.txt")
    assert result.warning_count >= 1
    assert any("NOME_CEDENTE" in warning for warning in result.warnings)


def test_structure_and_hidden_originals_are_protected(tmp_path: Path, source_files) -> None:
    excel = tmp_path / "intermediario.xlsx"
    prepare_workbook(
        *source_files,
        liquidation_date="03/09/2026",
        first_sequence=1,
        output_path=excel,
    )
    workbook = load_workbook(excel)
    originals = workbook["ORIGINAIS_CONTROLE"]
    originals.cell(2, 2).value = "ALTERADO"
    workbook.save(excel)
    workbook.close()
    loaded = read_intermediate(excel)
    assert any("ORIGINAIS_CONTROLE" in issue for issue in loaded.structural_issues)
    with pytest.raises(ValidationError):
        generate_cnab(excel, output_path=tmp_path / "bloqueado.txt")


def test_manifest_pfmi_pending_count_cannot_be_cleared(tmp_path: Path, source_files) -> None:
    pfmi, _, _ = source_files
    workbook = load_workbook(pfmi)
    workbook.active.append(["LINHA NAO RECONHECIDA"])
    workbook.save(pfmi)
    workbook.close()

    excel = tmp_path / "manifesto_pendente.xlsx"
    prepare_workbook(
        *source_files,
        liquidation_date="03/09/2026",
        first_sequence=1,
        output_path=excel,
    )
    workbook = load_workbook(excel)
    manifest = workbook["MANIFESTO"]
    manifest["B4"] = 0
    workbook.save(excel)
    workbook.close()

    loaded = read_intermediate(excel)
    assert any("pendências PFMI" in issue for issue in loaded.structural_issues)
    with pytest.raises(ValidationError):
        generate_cnab(excel, output_path=tmp_path / "bloqueado.txt")


def test_reordering_credit_rows_keeps_protected_origin_sequence(
    tmp_path: Path, source_files
) -> None:
    # Transforma o conjunto sintético em dois créditos (40 + 60 = PFMI 100).
    _, analytic, _ = source_files
    workbook = load_workbook(analytic)
    sheet = workbook.active
    sheet.cell(2, 4).value = 60
    sheet.cell(2, 5).value = 40
    second = [cell.value for cell in sheet[2]]
    second[3] = 90
    second[4] = 60
    second[6] = "P-2"
    sheet.append(second)
    workbook.save(analytic)
    workbook.close()

    excel = tmp_path / "intermediario.xlsx"
    prepare_workbook(
        *source_files,
        liquidation_date="03/09/2026",
        first_sequence=700,
        output_path=excel,
    )
    workbook = load_workbook(excel)
    sheet = workbook["CREDITOS"]
    assert sheet.max_row == 3
    first = [cell.value for cell in sheet[2]]
    second = [cell.value for cell in sheet[3]]
    for column, value in enumerate(second, start=1):
        sheet.cell(2, column).value = value
    for column, value in enumerate(first, start=1):
        sheet.cell(3, column).value = value
    # Security fix: two principal credits are never auto-selected; confirm
    # both manually (the reordering itself is the behaviour under test).
    include_column = {c.value: c.column for c in sheet[1]}["INCLUIR_CNAB"]
    sheet.cell(2, include_column).value = "SIM"
    sheet.cell(3, include_column).value = "SIM"
    workbook.save(excel)
    workbook.close()

    loaded = read_intermediate(excel)
    assert not loaded.structural_issues
    assert loaded.rows[0].values["REFERENCIA_ANALITICO"] == "P-1"
    txt = tmp_path / "reordenado.txt"
    generate_cnab(excel, output_path=txt)
    details = txt.read_bytes().split(b"\r\n")[1:3]
    assert details[0][37:62].strip() == b"700"
    assert details[0][192:205] == b"0000000004000"
    assert details[1][37:62].strip() == b"701"


@pytest.mark.parametrize("mutation", ["adicionar", "excluir", "duplicar_id"])
def test_structural_row_mutations_block_txt(tmp_path: Path, source_files, mutation: str) -> None:
    excel = tmp_path / f"{mutation}.xlsx"
    prepare_workbook(
        *source_files,
        liquidation_date="03/09/2026",
        first_sequence=1,
        output_path=excel,
    )
    workbook = load_workbook(excel)
    sheet = workbook["CREDITOS"]
    row = [cell.value for cell in sheet[2]]
    if mutation == "adicionar":
        row[0] = "ID-NOVO"
        sheet.append(row)
    elif mutation == "excluir":
        sheet.delete_rows(2)
    else:
        sheet.append(row)
    workbook.save(excel)
    workbook.close()
    with pytest.raises(ValidationError):
        generate_cnab(excel, output_path=tmp_path / "bloqueado.txt")


def test_missing_batch_inputs_only_block_second_step(tmp_path: Path, source_files) -> None:
    excel = tmp_path / "pendente.xlsx"
    prepare_workbook(
        *source_files, liquidation_date="invalida", first_sequence=0, output_path=excel
    )
    assert excel.exists()
    workbook = load_workbook(excel, data_only=True)
    summary = {
        workbook["RESUMO"].cell(row, 1).value: workbook["RESUMO"].cell(row, 2).value
        for row in range(1, workbook["RESUMO"].max_row + 1)
    }
    workbook.close()
    assert "DATA_LIQUIDACAO_INVALIDA" in summary["PENDENCIAS_INPUT_LOTE"]
    assert "PRIMEIRA_SEQUENCIA_INVALIDA" in summary["PENDENCIAS_INPUT_LOTE"]
    with pytest.raises(ValidationError, match="DATA_LIQUIDACAO"):
        generate_cnab(excel, output_path=tmp_path / "bloqueado.txt")


def test_absent_batch_inputs_are_explicit_in_excel(tmp_path: Path, source_files) -> None:
    excel = tmp_path / "ausentes.xlsx"
    prepare_workbook(*source_files, output_path=excel)
    workbook = load_workbook(excel, data_only=True)
    summary = {
        workbook["RESUMO"].cell(row, 1).value: workbook["RESUMO"].cell(row, 2).value
        for row in range(1, workbook["RESUMO"].max_row + 1)
    }
    workbook.close()
    assert "DATA_LIQUIDACAO_AUSENTE" in summary["PENDENCIAS_INPUT_LOTE"]
    assert "PRIMEIRA_SEQUENCIA_AUSENTE" in summary["PENDENCIAS_INPUT_LOTE"]


def test_sequence_outside_layout_is_explicit_in_preparation(tmp_path: Path, source_files) -> None:
    excel = tmp_path / "sequencia_fora.xlsx"
    prepare_workbook(
        *source_files,
        liquidation_date="03/09/2026",
        first_sequence=10_000_000_000,
        output_path=excel,
    )
    workbook = load_workbook(excel, data_only=True)
    summary = {
        workbook["RESUMO"].cell(row, 1).value: workbook["RESUMO"].cell(row, 2).value
        for row in range(1, workbook["RESUMO"].max_row + 1)
    }
    workbook.close()
    assert "FAIXA_SEQUENCIA_FORA_LAYOUT" in summary["PENDENCIAS_INPUT_LOTE"]


@pytest.mark.parametrize(
    ("beneficiary_value", "expected"),
    [
        (39.99, "DIVERGENCIA_VALOR_LINHA"),
        ("INVALIDO", "VALOR_BENEFICIARIO_INVALIDO"),
        (None, "VALOR_BENEFICIARIO_AUSENTE"),
    ],
)
def test_pfmi_source_value_issues_require_new_preparation(
    tmp_path: Path, source_files, beneficiary_value, expected: str
) -> None:
    pfmi, _, _ = source_files
    workbook = load_workbook(pfmi)
    sheet = workbook.active
    header_row = next(
        row
        for row in range(1, sheet.max_row + 1)
        if any(header_key(cell.value) == "VALORDAOPERACAO" for cell in sheet[row])
    )
    beneficiary_column = next(
        cell.column for cell in sheet[header_row] if header_key(cell.value) == "VALOR"
    )
    sheet.cell(header_row + 1, beneficiary_column).value = beneficiary_value
    workbook.save(pfmi)
    workbook.close()
    excel = tmp_path / "origem_pendente.xlsx"
    prepare_workbook(
        *source_files,
        liquidation_date="03/09/2026",
        first_sequence=1,
        output_path=excel,
    )
    with pytest.raises(ValidationError, match=expected):
        generate_cnab(excel, output_path=tmp_path / "bloqueado.txt")
    assert not (tmp_path / "bloqueado.txt").exists()


@pytest.mark.parametrize("field", ["ID_GRUPO", "REFERENCIA_ANALITICO", "TOTAL_PFMI_GRUPO"])
def test_provenance_controls_cannot_be_changed(tmp_path: Path, source_files, field: str) -> None:
    excel = tmp_path / f"controle_{field}.xlsx"
    prepare_workbook(
        *source_files,
        liquidation_date="03/09/2026",
        first_sequence=1,
        output_path=excel,
    )
    workbook = load_workbook(excel)
    sheet = workbook["CREDITOS"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers[field]).value = "ALTERADO"
    workbook.save(excel)
    workbook.close()
    with pytest.raises(ValidationError, match="[Cc]ontrole de proveniência"):
        generate_cnab(excel, output_path=tmp_path / "bloqueado.txt")
