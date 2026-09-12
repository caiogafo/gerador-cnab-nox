from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import load_workbook

import gerador_cnab_nox.gui as gui_module
from gerador_cnab_nox.gui import CESSAO_LABEL, NORMAL_LABEL, Application
from gerador_cnab_nox.models import CESSAO
from tests.test_gui import (
    _choose_modality,
    _select_pfmi,
    _set_commission,
    _wait_for_idle,
    _write_sources,
)
from tests.test_gui import gui_application as gui_application


def _select_file(monkeypatch, button, value: str) -> None:
    monkeypatch.setattr(
        gui_module.filedialog,
        "askopenfilename",
        lambda **_kwargs: value,
    )
    button.invoke()


def _correct_cession_emissions(intermediate: Path, emission: date) -> int:
    workbook = load_workbook(intermediate)
    sheet = workbook["CREDITOS"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    corrected = 0
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, headers["MODALIDADE"]).value == CESSAO:
            assert not sheet.cell(row, headers["DT_EMISSAO_TITULO"]).value
            sheet.cell(row, headers["DT_EMISSAO_TITULO"]).value = emission
            corrected += 1
    workbook.save(intermediate)
    workbook.close()
    return corrected


def _assert_status_text_is_current(app: Application) -> None:
    panel = app.results[app.current_step]
    assert panel.title.cget("text") in app.status.get()
    assert panel.body.cget("text") in app.status.get()
    assert str(panel.issue_text.cget("state")) == "disabled"


def test_cession_only_real_selectors_block_then_generate_from_corrected_excel(
    tmp_path: Path,
    monkeypatch,
    gui_application,
) -> None:
    root, app, dialogs = gui_application
    groups = [
        {
            "cedent": "EMPRESA CESSAO",
            "failure": "FALENCIA CESSAO",
            "operation": 115,
            "nominal": 200,
            "acquisition": 100,
            "cession": True,
        }
    ]
    pfmis, analytic, due = _write_sources(tmp_path, groups)
    intermediate = tmp_path / "cessao_intermediario.xlsx"
    cancelled_output = tmp_path / "cancelado.txt"
    blocked_output = tmp_path / "bloqueado_sem_emissao.txt"
    final_output = tmp_path / "cessao_final.txt"

    monkeypatch.setattr(
        gui_module.filedialog,
        "askopenfilenames",
        lambda **_kwargs: (str(pfmis[0]),),
    )
    app.add_pfmi_button.invoke()
    root.update()
    _select_pfmi(root, app, 0)
    _choose_modality(root, app, CESSAO_LABEL)
    assert app.pfmi_commission.get() == "15"

    _select_file(monkeypatch, app.analytic_select_button, "")
    assert app.analytic.get() == ""
    _select_file(monkeypatch, app.analytic_select_button, str(analytic))
    assert app.analytic.get() == str(analytic)
    _select_file(monkeypatch, app.due_select_button, "")
    assert app.due.get() == ""
    _select_file(monkeypatch, app.due_select_button, str(due))
    assert app.due.get() == str(due)

    app.liquidation.set("03/09/2026")
    app.sequence.set("810")
    monkeypatch.setattr(
        gui_module.filedialog,
        "asksaveasfilename",
        lambda **_kwargs: "",
    )
    app.prepare_button.invoke()
    _wait_for_idle(root, app)
    root.update()
    assert not intermediate.exists()

    monkeypatch.setattr(
        gui_module.filedialog,
        "asksaveasfilename",
        lambda **_kwargs: str(intermediate),
    )
    app.prepare_button.invoke()
    _wait_for_idle(root, app)
    root.update()
    assert intermediate.exists()
    assert "operações: 1" in app.status.get()
    assert "pendentes: 1" in app.status.get()
    _assert_status_text_is_current(app)

    monkeypatch.setattr(
        gui_module.filedialog,
        "asksaveasfilename",
        lambda **_kwargs: "",
    )
    app.generate_button.invoke()
    _wait_for_idle(root, app)
    root.update()
    assert not cancelled_output.exists()
    assert not dialogs["error"]

    monkeypatch.setattr(
        gui_module.filedialog,
        "asksaveasfilename",
        lambda **_kwargs: str(blocked_output),
    )
    app.generate_button.invoke()
    _wait_for_idle(root, app)
    root.update()
    assert not blocked_output.exists()
    assert "DT_EMISSAO_TITULO" in app.generation_result.details
    assert "inválida" in app.generation_result.details
    _assert_status_text_is_current(app)

    assert _correct_cession_emissions(intermediate, date(2026, 9, 2)) == 1
    for source in (*pfmis, analytic, due):
        source.unlink()

    _set_commission(root, app, "0")
    app.analytic.set(str(tmp_path / "analitico_indisponivel.xlsx"))
    app.due.set(str(tmp_path / "base_indisponivel.xlsx"))
    app.liquidation.set("31/12/2099")
    app.sequence.set("9999")
    app.intermediate.set("")
    _select_file(monkeypatch, app.intermediate_select_button, "")
    assert app.intermediate.get() == ""
    _select_file(monkeypatch, app.intermediate_select_button, str(intermediate))
    assert app.intermediate.get() == str(intermediate)

    errors_before_success = len(dialogs["error"])
    monkeypatch.setattr(
        gui_module.filedialog,
        "asksaveasfilename",
        lambda **_kwargs: str(final_output),
    )
    app.generate_button.invoke()
    _wait_for_idle(root, app)
    root.update()

    assert final_output.exists()
    records = final_output.read_bytes().split(b"\r\n")[:-1]
    assert len(records) == 3
    assert all(len(record) == 444 for record in records)
    assert records[1][37:62].strip() == b"810"
    assert "Títulos: 1" in app.status.get()
    assert "Sequência: 810–810" in app.status.get()
    assert len(dialogs["error"]) == errors_before_success
    _assert_status_text_is_current(app)


def test_mixed_batch_real_buttons_block_then_ignore_changed_gui_and_missing_sources(
    tmp_path: Path,
    monkeypatch,
    gui_application,
) -> None:
    root, app, dialogs = gui_application
    groups = [
        {
            "cedent": "EMPRESA NORMAL MISTA",
            "failure": "FALENCIA NORMAL MISTA",
            "operation": 90,
            "nominal": 130,
            "acquisition": 90,
            "cession": False,
        },
        {
            "cedent": "EMPRESA CESSAO MISTA",
            "failure": "FALENCIA CESSAO MISTA",
            "operation": 115,
            "nominal": 200,
            "acquisition": 100,
            "cession": True,
        },
    ]
    pfmis, analytic, due = _write_sources(tmp_path, groups)
    intermediate = tmp_path / "misto_intermediario.xlsx"
    blocked_output = tmp_path / "misto_bloqueado.txt"
    final_output = tmp_path / "misto_final.txt"

    monkeypatch.setattr(
        gui_module.filedialog,
        "askopenfilenames",
        lambda **_kwargs: tuple(str(path) for path in pfmis),
    )
    app.add_pfmi_button.invoke()
    root.update()
    _select_pfmi(root, app, 1)
    _choose_modality(root, app, CESSAO_LABEL)
    assert app.pfmi_commission.get() == "15"

    _select_file(monkeypatch, app.analytic_select_button, str(analytic))
    _select_file(monkeypatch, app.due_select_button, str(due))
    app.liquidation.set("03/09/2026")
    app.sequence.set("910")
    destinations = iter((str(intermediate), str(blocked_output)))
    monkeypatch.setattr(
        gui_module.filedialog,
        "asksaveasfilename",
        lambda **_kwargs: next(destinations),
    )

    app.prepare_button.invoke()
    _wait_for_idle(root, app)
    root.update()
    assert intermediate.exists()
    assert "operações: 2" in app.status.get()
    assert "prontas: 1" in app.status.get()
    assert "pendentes: 1" in app.status.get()
    app.generate_button.invoke()
    _wait_for_idle(root, app)
    root.update()
    assert not blocked_output.exists()
    assert "DT_EMISSAO_TITULO" in app.generation_result.details

    assert _correct_cession_emissions(intermediate, date(2026, 9, 2)) == 1
    for source in (*pfmis, analytic, due):
        source.unlink()

    _select_pfmi(root, app, 0)
    _choose_modality(root, app, CESSAO_LABEL)
    _set_commission(root, app, "99")
    _select_pfmi(root, app, 1)
    _choose_modality(root, app, NORMAL_LABEL)
    app.analytic.set(str(tmp_path / "outro_analitico.xlsx"))
    app.due.set(str(tmp_path / "outra_base.xlsx"))
    app.liquidation.set("01/01/2040")
    app.sequence.set("1")

    errors_before_success = len(dialogs["error"])
    monkeypatch.setattr(
        gui_module.filedialog,
        "asksaveasfilename",
        lambda **_kwargs: str(final_output),
    )
    app.generate_button.invoke()
    _wait_for_idle(root, app)
    root.update()

    assert final_output.exists()
    records = final_output.read_bytes().split(b"\r\n")[:-1]
    assert len(records) == 4
    assert all(len(record) == 444 for record in records)
    assert [record[37:62].strip() for record in records[1:3]] == [b"910", b"911"]
    assert "Títulos: 2" in app.status.get()
    assert "Sequência: 910–911" in app.status.get()
    assert len(dialogs["error"]) == errors_before_success
    _assert_status_text_is_current(app)


def test_module_entrypoint_builds_real_application_and_runs_event_loop(
    monkeypatch, gui_application
):
    import runpy
    from tkinter import Tk

    roots = []

    def timed_root():
        root = Tk()
        roots.append(root)
        root.after(150, root.quit)
        return root

    monkeypatch.setattr(gui_module, "Tk", timed_root)
    runpy.run_module("gerador_cnab_nox", run_name="__main__")
    assert len(roots) == 1
    roots[0].destroy()
