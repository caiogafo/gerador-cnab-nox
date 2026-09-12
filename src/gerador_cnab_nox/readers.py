from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from collections.abc import Iterable, Sequence
from dataclasses import replace
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .errors import InputFileError
from .models import AnalyticCredit, DueRecord, PfmiData, PfmiPayment, PfmiUnknownRow
from .normalize import (
    MAX_AGGREGATE,
    header_key,
    money,
    normalize_failure,
    normalize_name,
    parse_date,
    strip_accents,
)

PFMI_ALIASES = {
    "operation": {"VALORDAOPERACAO", "VALOROPERACAO"},
    "failure": {"FALENCIA", "CAMPANHA", "FALENCIASACADO"},
    "cedent": {"CREDORCEDENTE", "CREDOR", "CEDENTE"},
    "beneficiary": {"BENEFICIARIOSFAVORECIDOS", "BENEFICIARIO", "FAVORECIDO"},
    "beneficiary_doc": {"CPFCNPJ", "DOCUMENTO"},
    "beneficiary_value": {"VALOR"},
}

ANALYTIC_ALIASES = {
    "cedent_doc": {"CPF", "CPFCNPJ", "DOCUMENTO", "CPFCNPJCONTATO"},
    "cedent_name": {"CONTATO", "NOMECEDENTE", "CEDENTE", "NOMECONTATO"},
    "campaign": {"CAMPANHA", "FALENCIA"},
    "nominal": {"VALORRECEBER", "VALORNOMINAL"},
    "present": {"VALORAQUISICAO", "VALORPRESENTE"},
    "signature": {"DATAASSINATURA", "DATADEASSINATURA"},
    "reference": {"ID", "IDPROSPECT", "PROSPECT", "CODIGO", "CODIGOPROSPECT"},
}

DUE_ALIASES = {
    "name": {"NOMESACADO"},
    "document": {"DOCSACADO", "DOCUMENTOSACADO", "CPFCNPJSACADO"},
    "date": {"DATAVENCIMENTO", "VENCIMENTO"},
}


def _require_file(path: str | Path, extensions: set[str]) -> Path:
    candidate = Path(path)
    if not candidate.is_file():
        raise InputFileError(f"Arquivo não encontrado: {candidate}")
    if candidate.suffix.lower() not in extensions:
        accepted = ", ".join(sorted(extensions))
        raise InputFileError(f"Formato não suportado em {candidate.name}. Use: {accepted}")
    return candidate


def _resolve_headers(
    values: Sequence[Any], aliases: dict[str, set[str]], required: set[str]
) -> dict[str, int] | None:
    keys = [header_key(value) for value in values]
    result: dict[str, int] = {}
    for field, accepted in aliases.items():
        for index, key in enumerate(keys):
            if key in accepted:
                result[field] = index
                break
    if "reference" in aliases:
        for preferred in ("ID", "IDPROSPECT", "PROSPECT", "CODIGO", "CODIGOPROSPECT"):
            if preferred in keys:
                result["reference"] = keys.index(preferred)
                break
    return result if required.issubset(result) else None


def _cell(row: Sequence[Any], index: int | None) -> Any:
    if index is None or index >= len(row):
        return None
    return row[index]


def _nonempty_summary(row: Sequence[Any]) -> str:
    populated = [f"C{index + 1}" for index, value in enumerate(row) if str(value or "").strip()]
    return ",".join(populated)


def read_pfmi(path: str | Path) -> PfmiData:
    source = _require_file(path, {".xlsx"})
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
        original_book = load_workbook(source, read_only=True, data_only=False)
        raw_sheets = {sh.title: list(sh.values) for sh in original_book}
        original_book.close()
        result = PfmiData()
        digest = hashlib.sha256(source.read_bytes()).hexdigest()
        signatures: dict[str, str] = {}
        all_signatures = []
        found_header = False
        block_counter = 0
        for sheet in workbook.worksheets:
            rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
            payments, unknown, evidence, signature, found, block_counter = _pfmi_sheet(
                rows, sheet.title, source.name, block_counter, raw_sheets[sheet.title]
            )
            found_header |= found
            result.evidence.extend(
                {
                    "tipo": "LINHA_ORIGINAL_PFMI",
                    "aba": sheet.title,
                    "linha": i,
                    "original": list(row),
                }
                for i, row in enumerate(raw_sheets[sheet.title], 1)
                if any(v is not None for v in row)
            )
            result.unknown_rows.extend(unknown)
            result.evidence.extend(evidence)
            if payments:
                if signature in signatures:
                    result.evidence.append(
                        {
                            "tipo": "ESPELHO_INTEGRAL",
                            "aba": sheet.title,
                            "principal": signatures[signature],
                            "sha256_operacional": signature,
                            "pagamentos_nao_duplicados": len(payments),
                            "criterio": "conteudo integral ordenado e limites de blocos",
                        }
                    )
                else:
                    signatures[signature] = sheet.title
                    all_signatures.append(signature)
                    result.payments.extend(replace(item, source_hash=digest) for item in payments)
        for name, raw_rows in raw_sheets.items():
            auxiliary = _headerless_auxiliary(raw_rows, result.payments)
            if auxiliary:
                result.unknown_rows = [r for r in result.unknown_rows if r.source_sheet != name]
                result.evidence.append(
                    {
                        "tipo": auxiliary,
                        "aba": name,
                        "alerta": "AUXILIAR_LEGADO_PRESERVADO_CONFERIR",
                        "linhas": len(raw_rows),
                    }
                )
        workbook.close()
    except InputFileError:
        raise
    except Exception as exc:
        raise InputFileError(f"Não foi possível ler a PFMI: {source.name}") from exc
    if not found_header:
        raise InputFileError("PFMI sem os cabeçalhos mínimos necessários")
    result.operational_digest = hashlib.sha256("|".join(all_signatures).encode()).hexdigest()
    total_valid = all(p.beneficiary_value is not None for p in result.payments)
    total = sum((p.beneficiary_value or 0 for p in result.payments), 0)
    summaries = [e for e in result.evidence if e["tipo"] == "RESUMO_DECLARADO"]
    for evidence in summaries:
        evidence.update(
            pagamentos_calculados=len(result.payments),
            total_calculado=str(total) if total_valid else None,
        )
        evidence["alerta"] = (
            "RESUMO_DIVERGENTE"
            if evidence["quantidade"] != len(result.payments) or evidence["total"] != str(total)
            else "RESUMO_CONFERIDO"
        )
    if not summaries:
        result.evidence.append({"tipo": "RESUMO_AUSENTE", "alerta": "RESUMO_AUSENTE"})
    result.sources.append(
        {
            "arquivo": source.name,
            "sha256": digest,
            "pagamentos": len(result.payments),
            "total": str(total) if total_valid else None,
        }
    )
    return result


def _pfmi_sheet(rows, sheet_name, file_name, block_counter, raw_rows=None):
    payments, unknown, evidence, signature = [], [], [], []
    mapping = None
    aux = None
    current_block, title, previous_cedent, previous_failure = "", "", "", ""
    local_block = 0
    semantic_group = 0
    previous_group = None
    found = False
    last_table = None
    raw_rows = rows if raw_rows is None else raw_rows
    for number, row in enumerate(rows, 1):
        raw_row = raw_rows[number - 1] if number <= len(raw_rows) else row
        # A sum footer is only recognized after a real table and with no other data.
        filled_raw = [v for v in raw_row if v is not None and str(v).strip()]
        if (
            last_table
            and filled_raw
            and all(
                re.fullmatch(r"=SUM\([A-Z]+[0-9]+:[A-Z]+[0-9]+\)", str(v), re.I)
                or (
                    last_table == "mini"
                    and re.fullmatch(
                        r"=[A-Z]+[0-9]+\+\[[0-9]+\][A-Za-z0-9_]+!\$?[A-Z]+\$?[0-9]+", str(v)
                    )
                )
                or header_key(v) in {"TOTAL", "TOTALGERAL"}
                for v in filled_raw
            )
        ):
            evidence.append(
                {
                    "tipo": "TOTAL_AUXILIAR",
                    "aba": sheet_name,
                    "linha": number,
                    "original": list(raw_row),
                }
            )
            continue
        keys = [header_key(v) for v in row]
        nonempty = [str(v).strip() for v in row if v is not None and str(v).strip()]
        header = _resolve_headers(
            row,
            PFMI_ALIASES,
            {"operation", "failure", "cedent", "beneficiary", "beneficiary_value"},
        )
        if header:
            found = True
            last_table = "operational"
            if mapping is None:
                block_counter += 1
                local_block += 1
                current_block = f"B{block_counter:04d}"
                previous_cedent = previous_failure = ""
            mapping, aux = header, None
            continue
        summary = _resolve_headers(
            row,
            {
                "date": {"DATA"},
                "count": {"QUANTIDADEDETEDS", "QUANTIDADEDETED"},
                "total": {"VALORTOTALAPAGAR"},
            },
            {"date", "count"},
        )
        mini = _resolve_headers(row, PFMI_ALIASES, {"failure", "cedent", "beneficiary_value"})
        legacy_keys = _legacy_headers(keys)
        legacy = {"NOMECEDENTE", "VLNOMINAL", "VLPRESENTE", "SEUNUMERO", "NUDOCUMENTO"}.issubset(
            keys
        ) and bool({"DOCCEDENTE", "CNPJCEDENTE"}.intersection(keys))
        legacy = legacy or legacy_keys is not None
        if summary or legacy or mini:
            mapping = None
            aux = (
                ("summary", summary)
                if summary
                else ("mini", mini)
                if mini
                else ("legacy", legacy_keys or keys)
            )
            last_table = aux[0]
            evidence.append(
                {
                    "tipo": "CABECALHO_AUXILIAR",
                    "aba": sheet_name,
                    "linha": number,
                    "formato": aux[0],
                }
            )
            continue
        if not nonempty:
            mapping = None
            # Auxiliary empty summaries are evidence too, never a payment.
            if aux and aux[0] == "summary":
                evidence.append(
                    {
                        "tipo": "RESUMO_DECLARADO",
                        "aba": sheet_name,
                        "linha": number,
                        "quantidade": None,
                        "total": None,
                    }
                )
            aux = None
            continue
        if (
            mapping is None
            and len(nonempty) == 1
            and re.fullmatch(
                r"PFMI(?: FIDC NOX)?\s*-\s*\d{2}/\d{2}/\d{4}(?:\s*-\s*[\d.,]+)?", nonempty[0], re.I
            )
        ):
            evidence.append(
                {
                    "tipo": "TITULO_RESUMO_AUXILIAR",
                    "aba": sheet_name,
                    "linha": number,
                    "original": list(raw_row),
                }
            )
            aux = None
            continue
        if _is_ignorable_title(nonempty):
            title = " | ".join(nonempty)
            evidence.append({"tipo": "TITULO_FUNDO", "aba": sheet_name, "linha": number})
            continue
        if aux:
            kind, cols = aux
            if kind == "summary":
                try:
                    total = str(money(_cell(row, cols.get("total")), maximum=MAX_AGGREGATE))
                except ValueError:
                    total = None
                evidence.append(
                    {
                        "tipo": "RESUMO_DECLARADO",
                        "aba": sheet_name,
                        "linha": number,
                        "quantidade": _cell(row, cols["count"]),
                        "total": total,
                        "original": list(row),
                    }
                )
                aux = None
                continue
            if kind == "mini" and all(
                _cell(row, cols[k]) not in (None, "") for k in ("failure", "cedent")
            ):
                evidence.append(
                    {
                        "tipo": "RESUMO_CREDITOS_AUXILIAR",
                        "aba": sheet_name,
                        "linha": number,
                        "original": list(raw_row),
                        "alerta": "CONFERIR_RESUMO_AUXILIAR",
                    }
                )
                continue
            # A legacy row must contain its identifying and monetary columns.
            if kind == "legacy" and all(
                _cell(row, cols.index(k)) not in (None, "")
                for k in ("NOMECEDENTE", "VLNOMINAL", "VLPRESENTE")
            ):
                evidence.append(
                    {
                        "tipo": "LINHA_GERADOR_AUXILIAR",
                        "aba": sheet_name,
                        "linha": number,
                        "original": list(row),
                    }
                )
                continue
        if mapping is None and header_key("".join(nonempty)) == "DATA" and len(nonempty) == 1:
            has_mini = any(
                _resolve_headers(r, PFMI_ALIASES, {"failure", "cedent", "beneficiary_value"})
                for r in rows[number : number + 6]
            )
            if has_mini:
                aux = ("date", {})
                evidence.append(
                    {"tipo": "DATA_RESUMO_AUXILIAR", "aba": sheet_name, "linha": number}
                )
                continue
        if aux and aux[0] == "date" and len(nonempty) == 1:
            try:
                parse_date(nonempty[0])
            except ValueError:
                pass
            else:
                evidence.append(
                    {
                        "tipo": "DATA_RESUMO_AUXILIAR",
                        "aba": sheet_name,
                        "linha": number,
                        "original": list(raw_row),
                    }
                )
                aux = None
                continue
        if mapping is None:
            unknown.append(
                PfmiUnknownRow(
                    current_block or "FORA_BLOCO",
                    number,
                    "Linha preenchida fora de bloco não reconhecida",
                    _nonempty_summary(row),
                    sheet_name,
                    file_name,
                    raw_values=row,
                )
            )
            signature.append(["UNKNOWN", list(row)])
            continue
        cedent = _text(_cell(row, mapping.get("cedent")))
        failure = _text(_cell(row, mapping.get("failure")))
        beneficiary = _text(_cell(row, mapping.get("beneficiary")))
        raw_operation = _cell(row, mapping.get("operation"))
        raw_beneficiary = _cell(row, mapping.get("beneficiary_value"))
        if beneficiary and raw_beneficiary not in (None, ""):
            cedent = cedent or previous_cedent
            failure = failure or previous_failure
        if not cedent or not failure or not beneficiary:
            unknown.append(
                PfmiUnknownRow(
                    current_block,
                    number,
                    "Linha preenchida não reconhecida como pagamento",
                    _nonempty_summary(row),
                    sheet_name,
                    file_name,
                    raw_values=row,
                )
            )
            signature.append(["UNKNOWN", local_block, list(row)])
            continue
        issues = []
        try:
            operation = money(raw_operation, maximum=MAX_AGGREGATE)
            if operation < 0:
                issues.append("VALOR_OPERACAO_NEGATIVO")
        except ValueError:
            operation = None
            issues.append("VALOR_OPERACAO_INVALIDO")
        try:
            beneficiary_value = money(raw_beneficiary, maximum=MAX_AGGREGATE)
            if beneficiary_value < 0:
                issues.append("VALOR_BENEFICIARIO_NEGATIVO")
        except ValueError:
            beneficiary_value = None
            issues.append(
                "VALOR_BENEFICIARIO_AUSENTE"
                if raw_beneficiary in (None, "")
                else "VALOR_BENEFICIARIO_INVALIDO"
            )
        if (
            operation is not None
            and beneficiary_value is not None
            and operation != beneficiary_value
        ):
            issues.append("DIVERGENCIA_VALOR_LINHA")
        previous_cedent, previous_failure = cedent, failure
        payments.append(
            PfmiPayment(
                current_block,
                number,
                title,
                failure,
                cedent,
                beneficiary,
                _text(_cell(row, mapping.get("beneficiary_doc"))),
                operation,
                beneficiary_value,
                tuple(issues),
                source_sheet=sheet_name,
                source_file=file_name,
                raw_values=row,
            )
        )
        group_key = (local_block, normalize_name(cedent), normalize_failure(failure))
        if group_key != previous_group:
            semantic_group += 1
            previous_group = group_key
        signature.append(
            [
                "PAYMENT",
                semantic_group,
                *[
                    _cell(row, mapping.get(k))
                    for k in (
                        "operation",
                        "failure",
                        "cedent",
                        "beneficiary",
                        "beneficiary_doc",
                        "beneficiary_value",
                    )
                ],
            ]
        )
        if issues:
            unknown.append(
                PfmiUnknownRow(
                    current_block,
                    number,
                    "VALOR DA OPERAÇÃO inválido" if operation is None else "; ".join(issues),
                    _nonempty_summary(row),
                    sheet_name,
                    file_name,
                    raw_values=row,
                )
            )
    digest = hashlib.sha256(
        json.dumps(signature, default=str, ensure_ascii=False).encode()
    ).hexdigest()
    return payments, unknown, evidence, digest, found, block_counter


def _legacy_headers(keys):
    prefixes = (
        "CNPJCEDENTE",
        "NOMECEDENTE",
        "SEUNUMERO",
        "NUDOCUMENTO",
        "DTVENCIMENTO",
        "VLNOMINAL",
        "NUCPFCNPJSACADO",
        "NMSACADO",
        "VLPRESENTE",
        "IDENTIFICACAOCPFCNPJSACADO",
        "ENDERECO",
        "CEP",
        "TPTITULO",
        "DTEMISSAOTITULO",
        "COOBRIGACAO",
        "IDENTIFICACAOCPFCNPJCEDENTE",
        "NFE",
        "VALORPAGOTITULO",
        "INDEXADOR",
        "TAXADEJUROSINDEXADOR",
        "MOVIMENTO",
    )
    if len(keys) == 21 and all(
        k.startswith(prefix) for k, prefix in zip(keys, prefixes, strict=True)
    ):
        return list(prefixes)
    return None


def _headerless_auxiliary(rows, payments):
    """Recognize only demonstrated legacy projections, anchored to ordered identities.

    Partial columns never become payments nor prove operational mirrors. Unknown
    rows outside these complete shapes remain blocking. Values/formulas are kept raw.
    """
    nonempty = [
        (i, tuple(row))
        for i, row in enumerate(rows, 1)
        if any(v is not None and str(v).strip() for v in row)
    ]
    if not nonempty:
        return None
    by_sheet = {}
    for payment in payments:
        by_sheet.setdefault(payment.source_sheet, []).append(payment)
    for source in by_sheet.values():
        for columns, failure_col, cedent_col, formula_col, value_col in (
            (12, 2, 3, 11, 9),
            (9, 0, 1, 8, 7),
        ):
            if len(nonempty) != len(source):
                continue
            matches = True
            for (number, row), pay in zip(nonempty, source, strict=True):
                formula = _cell(row, formula_col)
                reference = f"{'J' if columns == 12 else 'H'}{number}"
                pattern = rf"={reference}\*-(?:[0-9]+(?:\.[0-9]+)?)"
                if (
                    len(row) != columns
                    or not re.fullmatch(pattern, str(formula), re.I)
                    or normalize_failure(_cell(row, failure_col)) != normalize_failure(pay.failure)
                    or normalize_name(_cell(row, cedent_col)) != normalize_name(pay.cedent_name)
                ):
                    matches = False
                    break
                try:
                    money(_cell(row, value_col))
                except ValueError:
                    matches = False
                    break
            if matches:
                return "PROJECAO_AUXILIAR_LEGADA_" + str(columns)
        # Known five-column date-led projection. It has no formula or payment header.
        if len(nonempty) == len(source) + 2:
            first, second = nonempty[:2]
            if len(first[1]) == 5 and [header_key(v) for v in first[1]] == ["", "DATA", "", "", ""]:
                try:
                    parse_date(_cell(second[1], 1))
                    valid = all(
                        len(row) == 5
                        and normalize_failure(row[0]) == normalize_failure(pay.failure)
                        and normalize_name(row[1]) == normalize_name(pay.cedent_name)
                        and money(row[3]) is not None
                        for (_, row), pay in zip(nonempty[2:], source, strict=True)
                    )
                except ValueError:
                    valid = False
                if valid:
                    return "PROJECAO_AUXILIAR_LEGADA_DATA_5"
    return None


def read_pfmis(inputs) -> PfmiData:
    if not inputs:
        raise InputFileError("Selecione ao menos uma PFMI")
    combined = PfmiData()
    seen = {}
    for order, item in enumerate(inputs, 1):
        data = read_pfmi(item.path)
        file_id = f"F{order:04d}"
        combined.payments.extend(
            replace(
                p,
                file_id=file_id,
                file_order=order,
                modality=item.modalidade,
                commission_text=item.comissao_texto,
                emissao_nova_cessao_texto=item.emissao_nova_cessao_texto,
            )
            for p in data.payments
        )
        combined.unknown_rows.extend(replace(p, file_id=file_id) for p in data.unknown_rows)
        combined.evidence.extend({**e, "id_arquivo": file_id} for e in data.evidence)
        if data.operational_digest in seen:
            combined.evidence.append(
                {
                    "tipo": "ARQUIVO_OPERACIONAL_DUPLICADO",
                    "id_arquivo": file_id,
                    "primeiro": seen[data.operational_digest],
                    "alerta": "ARQUIVO_OPERACIONAL_DUPLICADO_REVISAR",
                }
            )
        seen.setdefault(data.operational_digest, file_id)
        combined.sources.extend(
            {
                **s,
                "id_arquivo": file_id,
                "ordem": order,
                "modalidade": item.modalidade,
                "comissao_texto": item.comissao_texto,
                "emissao_nova_cessao_texto": item.emissao_nova_cessao_texto,
            }
            for s in data.sources
        )
    return combined


def _is_ignorable_title(values: Sequence[str]) -> bool:
    if len(values) > 3:
        return False
    normalized = strip_accents(" ".join(values)).upper()
    tokens = re.findall(r"[A-Z0-9]+", normalized)
    return bool(tokens) and tokens[0] in {"FUNDO", "FIDC", "NOX"}


def read_analytic(path: str | Path) -> list[AnalyticCredit]:
    source = _require_file(path, {".xls", ".xlsx", ".csv"})
    try:
        rows = list(_analytic_rows(source))
    except InputFileError:
        raise
    except Exception as exc:
        raise InputFileError(f"Não foi possível ler o Analítico: {source.name}") from exc

    header_position, mapping = _find_header(
        rows,
        ANALYTIC_ALIASES,
        {"cedent_doc", "cedent_name", "campaign", "nominal", "present", "signature"},
    )
    snapshot_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    source_sheet = _analytic_sheet_name(source)
    credits: list[AnalyticCredit] = []
    for physical_row, row in enumerate(rows[header_position + 1 :], start=header_position + 2):
        if not any(str(value or "").strip() for value in row):
            continue
        reference_value = _cell(row, mapping.get("reference"))
        try:
            nominal = money(_cell(row, mapping["nominal"]), allow_blank=True)
        except ValueError:
            nominal = None
        try:
            present = money(_cell(row, mapping["present"]), allow_blank=True)
        except ValueError:
            present = None
        try:
            signature = parse_date(_cell(row, mapping["signature"]), allow_blank=True)
        except ValueError:
            signature = None
        credits.append(
            AnalyticCredit(
                source_row=physical_row,
                reference=str(reference_value or f"LINHA-{physical_row}").strip(),
                cedent_document=_text(_cell(row, mapping["cedent_doc"])),
                cedent_name=str(_cell(row, mapping["cedent_name"]) or "").strip(),
                campaign=str(_cell(row, mapping["campaign"]) or "").strip(),
                nominal_value=nominal,
                acquisition_value=present,
                signature_date=signature,
                source_hash=snapshot_hash,
                source_sheet=source_sheet,
                raw_values={key: _cell(row, col) for key, col in mapping.items()},
            )
        )
    return credits


def _analytic_sheet_name(path):
    if path.suffix.lower() == ".csv":
        return "REGISTROS_LOGICOS"
    if path.suffix.lower() == ".xls":
        import xlrd

        book = xlrd.open_workbook(path, on_demand=True)
        name = book.sheet_names()[0]
        book.release_resources()
        return name
    book = load_workbook(path, read_only=True)
    name = book.active.title
    book.close()
    return name


def _analytic_rows(path: Path) -> Iterable[tuple[Any, ...]]:
    if path.suffix.lower() == ".csv":
        yield from _csv_rows(path)
    elif path.suffix.lower() == ".xls":
        import xlrd

        book = xlrd.open_workbook(path, on_demand=True)
        sheet = book.sheet_by_index(0)
        for row_number in range(sheet.nrows):
            values: list[Any] = []
            for column in range(sheet.ncols):
                cell = sheet.cell(row_number, column)
                value: Any = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        value = xlrd.xldate_as_datetime(value, book.datemode).date()
                    except (ValueError, OverflowError):
                        # Preserva o serial inválido para virar pendência de campo,
                        # em vez de classificar o arquivo inteiro como ilegível.
                        value = cell.value
                values.append(value)
            yield tuple(values)
        book.release_resources()
    else:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        yield from (tuple(row) for row in sheet.iter_rows(values_only=True))
        workbook.close()


def _csv_rows(path: Path) -> Iterable[tuple[str, ...]]:
    raw = path.read_bytes()
    text: str | None = None
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise InputFileError("CSV do Analítico não está em UTF-8 ou Windows-1252")
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";"
    yield from (tuple(row) for row in csv.reader(io.StringIO(text, newline=""), dialect))


def read_due_base(path: str | Path) -> list[DueRecord]:
    source = _require_file(path, {".xlsx"})
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
        sheet_name = workbook.active.title
        source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
        rows = [tuple(row) for row in workbook.active.iter_rows(values_only=True)]
        workbook.close()
    except Exception as exc:
        raise InputFileError(f"Não foi possível ler a Base de Vencimentos: {source.name}") from exc
    header_position, mapping = _find_header(rows, DUE_ALIASES, {"name", "document", "date"})
    result: list[DueRecord] = []
    for physical_row, row in enumerate(rows[header_position + 1 :], start=header_position + 2):
        if not any(str(value or "").strip() for value in row):
            continue
        try:
            due_date = parse_date(_cell(row, mapping["date"]), allow_blank=True)
        except ValueError:
            due_date = None
        result.append(
            DueRecord(
                source_row=physical_row,
                debtor_name=str(_cell(row, mapping["name"]) or "").strip(),
                debtor_document=_text(_cell(row, mapping["document"])),
                due_date=due_date,
                source_hash=source_hash,
                source_sheet=sheet_name,
                raw_values={key: _cell(row, col) for key, col in mapping.items()},
            )
        )
    return result


def _find_header(
    rows: Sequence[Sequence[Any]], aliases: dict[str, set[str]], required: set[str]
) -> tuple[int, dict[str, int]]:
    for position, row in enumerate(rows[:30]):
        mapping = _resolve_headers(row, aliases, required)
        if mapping:
            return position, mapping
    missing = ", ".join(sorted(required))
    raise InputFileError(f"Arquivo sem os cabeçalhos mínimos: {missing}")


def _text(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value or "").strip()
    match = re.fullmatch(r"(\d+)\.0+", text)
    return match.group(1) if match else text
