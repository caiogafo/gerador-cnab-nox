from __future__ import annotations

import hashlib
import json
import os
import tempfile
from dataclasses import asdict, dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

from .errors import InputFileError
from .models import CONTROL_FIELDS, FINAL_FIELDS, IMMUTABLE_FIELDS, PreparedBatch, PreparedRow
from .normalize import excel_safe

SCHEMA_VERSION = "CNAB-NOX-V1-2"
# Keep editable finals adjacent to the review decisions; provenance remains visible.
CREDIT_HEADERS = tuple(
    dict.fromkeys(
        (
            "ID_LINHA",
            "INCLUIR_CNAB",
            "APROVADO",
            "STATUS",
            "PENDENCIAS",
            "ALERTAS",
            "ARQUIVO_PFMI",
            "MODALIDADE",
            "NOME_CEDENTE_ANALITICO",
            "AQUISICAO_ORIGINAL",
            "NOMINAL_ORIGINAL",
            "PERCENTUAL_USADO",
            "COMISSAO_CALCULADA",
            "AQUISICAO_ESPERADA",
            *FINAL_FIELDS,
            *CONTROL_FIELDS,
        )
    )
)
ORIGINAL_HEADERS = tuple(dict.fromkeys((*CONTROL_FIELDS, *FINAL_FIELDS)))
PENDING_HEADERS = (
    "ID_BLOCO",
    "LINHA_PFMI",
    "MOTIVO",
    "COLUNAS_PREENCHIDAS",
    "ACAO",
    "ID_ARQUIVO",
    "ARQUIVO_PFMI",
    "ABA_PFMI",
    "ORIGINAL",
)


@dataclass
class LoadedIntermediate:
    rows: list[PreparedRow]
    liquidation_date: Any
    first_sequence: Any
    expected_ids: list[str]
    unknown_count: int
    structural_issues: list[str]


def write_intermediate(batch: PreparedBatch, output_path: str | Path) -> Path:
    destination = Path(output_path)
    if destination.suffix.lower() != ".xlsx":
        raise InputFileError("A saída intermediária deve usar a extensão .xlsx")
    if destination.exists():
        raise InputFileError(f"A saída já existe e não será sobrescrita: {destination.name}")
    destination.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    summary = workbook.active
    summary.title = "RESUMO"
    _write_summary(summary, batch)
    credits = workbook.create_sheet("CREDITOS")
    _write_credits(credits, batch.rows)
    pending = workbook.create_sheet("PENDENCIAS_PFMI")
    _write_pending(pending, batch)
    originals = workbook.create_sheet("ORIGINAIS_CONTROLE")
    _write_originals(originals, batch.rows)
    manifest = workbook.create_sheet("MANIFESTO")
    _write_manifest(manifest, batch, summary, pending)
    originals.sheet_state = "hidden"
    manifest.sheet_state = "hidden"
    for sheet in workbook:
        sheet.protection.sheet = True
        sheet.protection.selectLockedCells = False
        sheet.protection.selectUnlockedCells = False
        sheet.protection.autoFilter = False
    for row in credits.iter_rows(min_row=2):
        for cell in row:
            field = credits.cell(1, cell.column).value
            cell.protection = Protection(
                locked=field not in {*FINAL_FIELDS, "INCLUIR_CNAB", "APROVADO"}
            )
    summary["B3"].protection = Protection(locked=False)
    summary["B4"].protection = Protection(locked=False)

    _atomic_save(workbook, destination)
    workbook.close()
    return destination


def _write_summary(sheet, batch: PreparedBatch) -> None:
    selected = [row for row in batch.rows if row.values.get("INCLUIR_CNAB") == "SIM"]
    total_nominal = sum(
        (row.values.get("VL_NOMINAL") or Decimal("0.00") for row in selected),
        Decimal("0.00"),
    )
    total_present = sum(
        (row.values.get("VL_PRESENTE") or Decimal("0.00") for row in selected),
        Decimal("0.00"),
    )
    rows = (
        ("GERADOR CNAB NOX — DOUBLE CHECK", ""),
        ("VERSAO_ESQUEMA", SCHEMA_VERSION),
        ("DATA_LIQUIDACAO", batch.liquidation_date),
        ("PRIMEIRA_SEQUENCIA", batch.first_sequence),
        ("PENDENCIAS_INPUT_LOTE", "; ".join(batch.warnings) or "NENHUMA"),
        ("CANDIDATOS", len(batch.rows)),
        ("SELECIONADOS_PREVISTOS", len(selected)),
        ("TOTAL_NOMINAL_PREVISTO", total_nominal),
        ("TOTAL_PRESENTE_PREVISTO", total_present),
        ("PENDENCIAS_PFMI", len(batch.unknown_rows)),
        ("INSTRUCAO_1", "Revise todos os campos finais na aba CREDITOS."),
        ("INSTRUCAO_2", "Use INCLUIR_CNAB=SIM/NAO; vazio bloqueia a geração."),
        ("INSTRUCAO_3", "APROVADO=SIM libera somente divergência material de nome."),
        ("INSTRUCAO_4", "Não inclua, exclua ou duplique linhas."),
    )
    for row in rows:
        sheet.append([_xlsx_value(excel_safe(value)) for value in row])
    _append_summary_details(sheet, batch)
    sheet.column_dimensions["A"].width = 38
    sheet.column_dimensions["B"].width = 78
    sheet["A1"].font = Font(bold=True, color="FFFFFF", size=14)
    sheet["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    sheet.merge_cells("A1:B1")
    sheet.row_dimensions[1].height = 28
    for cell in (sheet["B3"],):
        cell.number_format = "DD/MM/YYYY"
        cell.fill = PatternFill("solid", fgColor="FFF2CC")
    sheet["B4"].fill = PatternFill("solid", fgColor="FFF2CC")
    sheet["B8"].number_format = "#,##0.00"
    sheet["B9"].number_format = "#,##0.00"
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.value = _xlsx_value(cell.value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[row[0].row].height = 42
        if "TOTAL_" in str(row[0].value):
            row[1].number_format = "#,##0.00"
    sheet.freeze_panes = "A2"


def _append_summary_details(sheet, batch):
    sheet.append(("PAGAMENTOS_BANCARIOS", batch.payment_count))
    sheet.append(("GRUPOS_FISICOS", batch.group_count))
    selected = [r for r in batch.rows if r.values.get("INCLUIR_CNAB") == "SIM"]
    sequence = batch.first_sequence
    sheet.append(
        (
            "FAIXA_SEQUENCIAL_PREVISTA",
            f"{sequence} a {sequence + len(selected) - 1}"
            if isinstance(sequence, int) and selected
            else "PENDENTE",
        )
    )
    groups = {r.values.get("ID_GRUPO"): r.values.get("TOTAL_PFMI_GRUPO") for r in batch.rows}
    sheet.append(
        ("TOTAL_PFMI_VALIDO", sum((v for v in groups.values() if v is not None), Decimal(0)))
    )
    sheet.append(("GRUPOS_COM_TOTAL_INCOMPLETO", sum(v is None for v in groups.values())))
    sheet.append(
        ("LEGENDA", "Amarelo: editável; azul: referência protegida. Filtros na aba CREDITOS.")
    )
    sheet.append(
        (
            "ATUALIZACAO_RESUMO",
            "Totais, status e diferenças são da preparação. Gerar CNAB revalida os finais salvos.",
        )
    )
    sheet.append(
        (
            "CORRIGIR_PARAMETROS",
            "Modalidade/taxa/ordem: alterar por PFMI e preparar novamente. Data/sequência: B3/B4.",
        )
    )
    sheet.append(
        (
            "RESOLVER_PENDENCIAS",
            "Base e finais: corrigir em CREDITOS. PFMI/linha sem origem/taxa inválida: "
            "corrigir origem e preparar novamente.",
        )
    )
    for source in batch.sources:
        if source.get("id_arquivo"):
            _append_file_summary(sheet, batch, source)
        elif source.get("tipo") in {"ANALITICO", "BASE"}:
            sheet.append(("FONTE_" + source["tipo"], "SHA256: " + source["sha256"]))
        elif source.get("tipo") == "EQUIVALENCIAS_FALENCIAS":
            for alternative, canonical in sorted(source["confirmadas"].items()):
                sheet.append(("EQUIVALENCIA_FALENCIA_CONFIRMADA", f"{alternative} = {canonical}"))
    for evidence in batch.evidence:
        if evidence.get("tipo") == "LINHA_ORIGINAL_PFMI":
            continue
        visible = " | ".join(
            f"{k}: {v}"
            for k, v in evidence.items()
            if k
            in {
                "tipo",
                "id_arquivo",
                "aba",
                "principal",
                "linha",
                "alerta",
                "quantidade",
                "total",
                "pagamentos_nao_duplicados",
            }
        )
        sheet.append(
            (
                "RESUMO_ESPELHO_ALERTA",
                excel_safe(visible),
            )
        )


def _append_file_summary(sheet, batch, source):
    file_id = source["id_arquivo"]
    rows = [r.values for r in batch.rows if r.values["ID_ARQUIVO"] == file_id]
    selected = [r for r in rows if r["INCLUIR_CNAB"] == "SIM"]
    groups = {r["ID_GRUPO"]: r["TOTAL_PFMI_GRUPO"] for r in rows}
    counts = {r["ID_GRUPO"]: r["PAGAMENTOS_GRUPO"] for r in rows}
    values = [
        ("ARQUIVO", source.get("arquivo", source.get("nome", ""))),
        ("MODALIDADE", source["modalidade"]),
        (
            "COMISSAO_INFORMADA",
            source["comissao_texto"] if source["modalidade"] != "NORMAL" else "IGNORADA EM NORMAL",
        ),
        ("PAGAMENTOS", sum(counts.values())),
        ("GRUPOS", len(groups)),
        ("CANDIDATOS", len(rows)),
        ("SELECIONADOS_PREVISTOS", len(selected)),
        ("TOTAL_PFMI_VALIDO", sum((v for v in groups.values() if v is not None), Decimal(0))),
        ("GRUPOS_INCOMPLETOS", sum(v is None for v in groups.values())),
        (
            "TOTAL_NOMINAL_PREVISTO",
            sum((r["VL_NOMINAL"] or Decimal(0) for r in selected), Decimal(0)),
        ),
        (
            "TOTAL_PRESENTE_PREVISTO",
            sum((r["VL_PRESENTE"] or Decimal(0) for r in selected), Decimal(0)),
        ),
    ]
    for key, value in values:
        sheet.append((f"{file_id}_{key}", _xlsx_value(excel_safe(value))))


def _write_credits(sheet, rows: list[PreparedRow]) -> None:
    sheet.append(CREDIT_HEADERS)
    for prepared in rows:
        sheet.append(
            [_xlsx_value(excel_safe(prepared.values.get(field))) for field in CREDIT_HEADERS]
        )
    _style_table(sheet, len(CREDIT_HEADERS))
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = sheet.dimensions
    yes_no = DataValidation(type="list", formula1='"SIM,NAO"', allow_blank=True)
    approved = DataValidation(type="list", formula1='"SIM,NAO"', allow_blank=True)
    sheet.add_data_validation(yes_no)
    sheet.add_data_validation(approved)
    header_map = {cell.value: cell.column for cell in sheet[1]}
    if sheet.max_row >= 2:
        yes_no.add(
            f"{sheet.cell(1, header_map['INCLUIR_CNAB']).column_letter}2:"
            f"{sheet.cell(1, header_map['INCLUIR_CNAB']).column_letter}{sheet.max_row}"
        )
        approved.add(
            f"{sheet.cell(1, header_map['APROVADO']).column_letter}2:"
            f"{sheet.cell(1, header_map['APROVADO']).column_letter}{sheet.max_row}"
        )
        status_letter = sheet.cell(1, header_map["STATUS"]).column_letter
        range_ref = f"{status_letter}2:{status_letter}{sheet.max_row}"
        sheet.conditional_formatting.add(
            range_ref,
            FormulaRule(
                formula=[f'${status_letter}2<>"OK"'],
                fill=PatternFill("solid", fgColor="FFF2CC"),
            ),
        )
    for field in (
        "VL_NOMINAL",
        "VL_PRESENTE",
        "VALOR_PAGO_TITULO",
        "TOTAL_PFMI_GRUPO",
        "AQUISICAO_ORIGINAL",
        "NOMINAL_ORIGINAL",
        "BASE_COMISSAO",
        "COMISSAO_CALCULADA",
        "AQUISICAO_ESPERADA",
        "DIFERENCA_PRESENTE",
        "DIFERENCA_NOMINAL",
    ):
        column = header_map[field]
        for cell in sheet.iter_cols(min_col=column, max_col=column, min_row=2):
            for item in cell:
                item.number_format = "#,##0.00"
    for field in ("DT_VENCIMENTO", "DT_EMISSAO_TITULO", "ASSINATURA_ORIGINAL"):
        column = header_map[field]
        for cell in sheet.iter_cols(min_col=column, max_col=column, min_row=2):
            for item in cell:
                item.number_format = "DD/MM/YYYY"
    widths = {
        "STATUS": 30,
        "PENDENCIAS": 55,
        "ALERTAS": 45,
        "NOME_CEDENTE": 35,
        "NOME_SACADO": 35,
        "NOME_CEDENTE_PFMI": 35,
        "NOME_CEDENTE_ANALITICO": 35,
        "ARQUIVO_PFMI": 32,
        "FALENCIA_PFMI": 35,
        "DOC_BENEFICIARIO_PFMI": 28,
        "PENDENCIAS_FONTE_PFMI": 38,
        "ORIGENS_CNPJ_CONEXCRED": 65,
    }
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            field = sheet.cell(1, cell.column).value
            editable = field in {*FINAL_FIELDS, "INCLUIR_CNAB", "APROVADO"}
            cell.fill = PatternFill("solid", fgColor="FFF2CC" if editable else "EAF2F8")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if field.startswith(("DOC_", "ID_", "HASH_")) or field in {
                "CEP",
                "NFE",
                "PERCENTUAL_USADO",
                "COMISSAO_TEXTO",
            }:
                cell.number_format = "@"
        sheet.row_dimensions[row[0].row].height = 44
    for field, width in widths.items():
        sheet.column_dimensions[sheet.cell(1, header_map[field]).column_letter].width = width


def _write_pending(sheet, batch: PreparedBatch) -> None:
    headers = PENDING_HEADERS
    sheet.append(headers)
    for row in batch.unknown_rows:
        sheet.append(
            (
                row.block_id,
                row.source_row,
                excel_safe(row.reason),
                row.nonempty_columns,
                "Corrigir a PFMI e executar Preparar CNAB novamente",
                row.file_id,
                excel_safe(row.source_file),
                excel_safe(row.source_sheet),
                json.dumps(row.raw_values, default=str, ensure_ascii=False),
            )
        )
    _style_table(sheet, len(headers))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.column_dimensions["C"].width = 55
    sheet.column_dimensions["E"].width = 60


def _write_originals(sheet, rows: list[PreparedRow]) -> None:
    sheet.append(ORIGINAL_HEADERS)
    for row in rows:
        sheet.append(
            [_xlsx_value(excel_safe(row.originals.get(field))) for field in ORIGINAL_HEADERS]
        )
    _style_table(sheet, len(ORIGINAL_HEADERS))


def _sheet_digest(sheet, editable=()):
    content = [
        [_canonical(None if cell.coordinate in editable else cell.value) for cell in row]
        for row in sheet.iter_rows()
    ]
    return hashlib.sha256(json.dumps(content, ensure_ascii=False).encode()).hexdigest()


def _manifest_digest(rows):
    return hashlib.sha256(
        json.dumps([[_canonical(v) for v in row] for row in rows], ensure_ascii=False).encode()
    ).hexdigest()


def _write_manifest(sheet, batch, summary, pending):
    ids = [str(row.values["ID_LINHA"]) for row in batch.rows]
    originals = {str(row.originals["ID_LINHA"]): row.originals for row in batch.rows}
    rows = [
        ("CHAVE", "VALOR"),
        ("VERSAO_ESQUEMA", SCHEMA_VERSION),
        ("QUANTIDADE_IDS", len(ids)),
        ("PENDENCIAS_PFMI", len(batch.unknown_rows)),
        ("DIGEST_IDS", _ids_digest(ids, len(batch.unknown_rows))),
        ("DIGEST_ORIGINAIS", _originals_digest(ids, originals)),
        ("DIGEST_RESUMO", _sheet_digest(summary, ("B3", "B4"))),
        ("DIGEST_PENDENCIAS", _sheet_digest(pending)),
        ("PAGAMENTOS", batch.payment_count),
        ("GRUPOS", batch.group_count),
    ]
    rows.extend(("ID_ESPERADO", line_id) for line_id in ids)
    for key, items in (
        ("FONTE", batch.sources),
        ("EVIDENCIA", batch.evidence),
        ("PAGAMENTO_ORIGINAL", batch.payments_snapshot),
        ("PENDENCIA_ORIGINAL", [asdict(r) for r in batch.unknown_rows]),
    ):
        rows.extend((key, json.dumps(item, default=str, ensure_ascii=False)) for item in items)
    for row in rows:
        sheet.append(row)
    sheet.append(("DIGEST_MANIFESTO", _manifest_digest(rows)))
    _style_table(sheet, 2)
    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 85
    for row in sheet.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[row[0].row].height = 36


def _style_table(sheet, columns: int) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 44
    for column in range(1, columns + 1):
        letter = sheet.cell(1, column).column_letter
        if sheet.column_dimensions[letter].width == 13.0:
            sheet.column_dimensions[letter].width = 20
    for row in sheet.iter_rows(min_row=2):
        sheet.row_dimensions[row[0].row].height = 44
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if str(sheet.cell(1, cell.column).value).startswith(("DOC_", "ID_", "HASH_")):
                cell.number_format = "@"


def _xlsx_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        # Excel guarantees 15 significant decimal digits. Preserve larger source
        # aggregates as exact text; money() parses it without a binary round trip.
        return format(value, "f") if abs(value) > Decimal("9999999999999.99") else float(value)
    return value


def _atomic_save(workbook: Workbook, destination: Path) -> None:
    descriptor, temporary = tempfile.mkstemp(
        prefix=f".{destination.stem}-", suffix=".xlsx.tmp", dir=destination.parent
    )
    os.close(descriptor)
    temp_path = Path(temporary)
    try:
        workbook.save(temp_path)
        os.link(temp_path, destination)  # atomic no-clobber publication
    finally:
        temp_path.unlink(missing_ok=True)


def _ids_digest(ids: list[str], unknown_count: int) -> str:
    content = f"{SCHEMA_VERSION}|PENDENCIAS_PFMI={unknown_count}|" + "|".join(ids)
    return hashlib.sha256(content.encode("utf-8")).hexdigest()


def _originals_digest(ids: list[str], originals: dict[str, dict[str, Any]]) -> str:
    parts = [SCHEMA_VERSION]
    for line_id in ids:
        parts.append(line_id)
        row = originals.get(line_id, {})
        for header in ORIGINAL_HEADERS:
            parts.append(header)
            parts.append(_canonical(row.get(header)))
    return hashlib.sha256("|".join(parts).encode("utf-8")).hexdigest()


def _canonical(value: Any) -> str:
    value = excel_safe(_xlsx_value(value))
    if value is None or value == "":
        return "n:"
    if isinstance(value, datetime):
        return "d:" + value.date().isoformat()
    if isinstance(value, date):
        return "d:" + value.isoformat()
    if isinstance(value, bool):
        return "b:" + str(int(value))
    if isinstance(value, (int, float, Decimal)):
        return "m:" + format(Decimal(str(value)).normalize(), "f")
    return "s:" + str(value)


def read_intermediate(path: str | Path) -> LoadedIntermediate:
    source = Path(path)
    if not source.is_file() or source.suffix.lower() != ".xlsx":
        raise InputFileError("Selecione um Excel intermediário .xlsx existente")
    try:
        workbook = load_workbook(source, data_only=False, read_only=False)
    except Exception as exc:
        raise InputFileError(f"Não foi possível ler o Excel intermediário: {source.name}") from exc
    required_sheets = {
        "RESUMO",
        "CREDITOS",
        "PENDENCIAS_PFMI",
        "ORIGINAIS_CONTROLE",
        "MANIFESTO",
    }
    missing_sheets = required_sheets.difference(workbook.sheetnames)
    if missing_sheets:
        workbook.close()
        raise InputFileError("Excel intermediário incompleto: " + ", ".join(sorted(missing_sheets)))

    structural: list[str] = []
    summary = _key_value_sheet(workbook["RESUMO"])
    manifest = _key_value_sheet(workbook["MANIFESTO"])
    if summary.get("VERSAO_ESQUEMA") != SCHEMA_VERSION:
        structural.append("Versão do esquema incompatível no RESUMO")
    if manifest.get("VERSAO_ESQUEMA") != SCHEMA_VERSION:
        structural.append("Versão do esquema incompatível no MANIFESTO")
    if (
        summary.get("VERSAO_ESQUEMA") == "CNAB-NOX-V1-1"
        or manifest.get("VERSAO_ESQUEMA") == "CNAB-NOX-V1-1"
    ):
        structural.append(
            "Este Excel foi preparado por uma versão anterior. "
            "Prepare novamente as fontes na versão revisada."
        )
    manifest_rows = list(workbook["MANIFESTO"].values)
    if (
        not manifest_rows
        or manifest_rows[-1][0] != "DIGEST_MANIFESTO"
        or manifest_rows[-1][1] != _manifest_digest(manifest_rows[:-1])
    ):
        structural.append("Conteúdo/estrutura do MANIFESTO foi alterado")
    expected_ids = [str(row[1]) for row in manifest_rows if row[0] == "ID_ESPERADO"]
    if _safe_int(manifest.get("QUANTIDADE_IDS")) != len(expected_ids) or len(
        set(expected_ids)
    ) != len(expected_ids):
        structural.append("Contagem/IDs do MANIFESTO inconsistentes")
    unknown_count = _safe_int(manifest.get("PENDENCIAS_PFMI"))
    if manifest.get("DIGEST_IDS") != _ids_digest(expected_ids, unknown_count):
        structural.append("Manifesto de IDs ou pendências PFMI foi alterado")
    if manifest.get("DIGEST_RESUMO") != _sheet_digest(workbook["RESUMO"], ("B3", "B4")):
        structural.append("Parâmetros/controles do RESUMO foram alterados")
    if manifest.get("DIGEST_PENDENCIAS") != _sheet_digest(workbook["PENDENCIAS_PFMI"]):
        structural.append("Conteúdo da aba PENDENCIAS_PFMI foi alterado")
    if set(workbook.sheetnames) != required_sheets:
        structural.append("Estrutura das cinco abas foi alterada")

    originals_by_id = _rows_by_id(workbook["ORIGINAIS_CONTROLE"], ORIGINAL_HEADERS, structural)
    current = _rows_by_id(workbook["CREDITOS"], CREDIT_HEADERS, structural)
    current_ids = list(current)
    if len(current_ids) != len(expected_ids) or set(current_ids) != set(expected_ids):
        structural.append("Linhas/IDs da aba CREDITOS foram incluídos, removidos ou alterados")
    if set(originals_by_id) != set(expected_ids):
        structural.append("A aba ORIGINAIS_CONTROLE foi alterada")
    if manifest.get("DIGEST_ORIGINAIS") != _originals_digest(expected_ids, originals_by_id):
        structural.append("O conteúdo da aba ORIGINAIS_CONTROLE foi alterado")

    rows: list[PreparedRow] = []
    for line_id in expected_ids:
        if line_id not in current:
            continue
        for field in IMMUTABLE_FIELDS:
            if _canonical(current[line_id].get(field)) != _canonical(
                originals_by_id.get(line_id, {}).get(field)
            ):
                structural.append(f"Controle de proveniência {field} foi alterado")
        rows.append(
            PreparedRow(
                values=current[line_id],
                originals=originals_by_id.get(line_id, {}),
            )
        )
    loaded = LoadedIntermediate(
        rows=rows,
        liquidation_date=summary.get("DATA_LIQUIDACAO"),
        first_sequence=summary.get("PRIMEIRA_SEQUENCIA"),
        expected_ids=expected_ids,
        unknown_count=unknown_count,
        structural_issues=structural,
    )
    workbook.close()
    return loaded


def _key_value_sheet(sheet) -> dict[str, Any]:
    return {
        str(sheet.cell(row, 1).value or "").strip(): sheet.cell(row, 2).value
        for row in range(1, sheet.max_row + 1)
        if str(sheet.cell(row, 1).value or "").strip()
    }


def _rows_by_id(sheet, required_headers: tuple[str, ...], issues: list[str]) -> dict[str, dict]:
    headers = [str(cell.value or "").strip() for cell in sheet[1]]
    if len(headers) != len(required_headers) or len(set(headers)) != len(headers):
        issues.append(f"Estrutura de colunas da aba {sheet.title} foi alterada")
    missing = [header for header in required_headers if header not in headers]
    if missing:
        issues.append(f"Aba {sheet.title} sem colunas obrigatórias: {', '.join(missing)}")
        return {}
    positions = {header: headers.index(header) + 1 for header in required_headers}
    result: dict[str, dict] = {}
    for row_number in range(2, sheet.max_row + 1):
        populated = any(
            sheet.cell(row_number, column).value not in (None, "") for column in positions.values()
        )
        if not populated:
            issues.append(f"Aba {sheet.title} possui linha vazia inserida")
            continue
        values = {header: sheet.cell(row_number, col).value for header, col in positions.items()}
        line_id = str(values.get("ID_LINHA") or "").strip()
        if not line_id:
            issues.append(f"Aba {sheet.title} possui linha sem ID_LINHA")
            continue
        if line_id in result:
            issues.append(f"Aba {sheet.title} possui ID_LINHA duplicado: {line_id}")
        result[line_id] = values
    return result


def _safe_int(value: Any) -> int:
    try:
        return max(0, int(value))
    except (TypeError, ValueError, OverflowError):
        return 0
