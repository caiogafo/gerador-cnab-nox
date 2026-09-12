"""Reproducible visible Tk smoke with synthetic inputs, real commands and optional PNGs.

Needs a graphical DISPLAY (or xvfb-run) and the dev environment. The optional
capture Python must have Pillow; capture is limited to this application's window.
"""

from __future__ import annotations

import argparse
import json
import subprocess
import sys
import time
from datetime import date
from pathlib import Path
from tkinter import Tk
from unittest.mock import patch

# QA helpers are deliberately outside the distributed runtime package.
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook

from gerador_cnab_nox.gui import CESSAO_LABEL, Application
from tests.test_gui import (
    _choose_modality,
    _select_pfmi,
    _set_commission,
    _wait_for_idle,
    _write_sources,
)


def run(output, capture_python=None, scale=1.0, window="1040x700"):

    output.mkdir(parents=True, exist_ok=False)
    groups = [
        dict(
            cedent="EMPRESA NORMAL",
            failure="FALENCIA NORMAL",
            operation=90,
            nominal=130,
            acquisition=90,
            cession=False,
        ),
        dict(
            cedent="EMPRESA ZERO",
            failure="FALENCIA ZERO",
            operation=80,
            nominal=200,
            acquisition=80,
            cession=True,
        ),
        dict(
            cedent="EMPRESA FRACAO",
            failure="FALENCIA FRACAO",
            operation=100.5,
            nominal=200,
            acquisition=100,
            cession=True,
        ),
    ]
    pfmis, analytic, due = _write_sources(output, groups)
    root = Tk()
    root.tk.call("tk", "scaling", (96 / 72) * scale)
    app = Application(root)

    # Visible automated runs must not consume unrelated clicks/keys from the desktop.
    # invoke(), variable edits and virtual selection events still exercise real widgets.
    def protect(widget):
        widget.bindtags(("SyntheticCapture", *widget.bindtags()))
        for child in widget.winfo_children():
            protect(child)

    for event in ("<ButtonPress>", "<KeyPress>"):
        root.bind_class("SyntheticCapture", event, lambda _event: "break")
    protect(root)
    root.geometry(window)
    root.update()
    snapshots = []
    result_states = []

    def capture(label):
        root.lift()
        painted = time.monotonic() + 0.2
        while time.monotonic() < painted:
            root.update()
            time.sleep(0.01)
        assert len(app._ordered_pfmi_inputs()) == 3, "O lote sintético deve manter três PFMIs"
        geometry = [root.winfo_rootx(), root.winfo_rooty(), root.winfo_width(), root.winfo_height()]
        bounds = {
            name: dict(
                x=widget.winfo_rootx() - geometry[0],
                y=widget.winfo_rooty() - geometry[1],
                width=widget.winfo_width(),
                height=widget.winfo_height(),
                mapped=bool(widget.winfo_ismapped()),
            )
            for name, widget in [
                ("prepare", app.prepare_button),
                ("generate", app.generate_button),
                ("open_excel", app.open_excel_button),
                ("open_txt_folder", app.open_generated_folder_button),
                ("result", app.results.get(app.current_step, app.review_result).body),
                ("result_title", app.results.get(app.current_step, app.review_result).title),
                ("viewport", app.body.canvas),
            ]
        }
        if label in {"02_excel_pendente", "03_emissao_bloqueada", "04_txt_gerado"}:
            title, viewport = bounds["result_title"], bounds["viewport"]
            assert viewport["y"] <= title["y"]
            assert title["y"] + title["height"] <= viewport["y"] + viewport["height"]
        result_states.append(
            dict(step=app.current_step, kind=app.results[app.current_step].kind)
            if app.current_step in app.results
            else dict(step=1, kind="info")
        )
        snapshots.append(
            dict(etapa=label, geometry=geometry, widgets=bounds, status=app.status.get())
        )
        if capture_python:
            target = output / f"{label}.png"
            capture_script = Path(__file__).with_name("capture_tk_window.py")
            subprocess.run(
                [
                    str(capture_python),
                    str(capture_script),
                    str(root.winfo_id()),
                    str(geometry[2]),
                    str(geometry[3]),
                    str(target),
                ],
                check=True,
            )

    try:
        with (
            patch(
                "gerador_cnab_nox.gui.filedialog.askopenfilenames",
                return_value=tuple(map(str, pfmis)),
            ),
        ):
            app.add_pfmi_button.invoke()
            root.update()
            for index, rate in [(1, "0"), (2, "0,5")]:
                _select_pfmi(root, app, index)
                _choose_modality(root, app, CESSAO_LABEL)
                _set_commission(root, app, rate)
            app.analytic.set(str(analytic))
            app.due.set(str(due))
            app.liquidation.set("03/09/2026")
            app.sequence.set("500")
            app.body.canvas.yview_moveto(0)
            capture("01_lista_mista")
            app.body.canvas.yview_moveto(1)
            capture("01b_configuracao")
            app.continue_button.invoke()
            capture("02_antes_preparar")
            excel = output / "intermediario.xlsx"
            with patch(
                "gerador_cnab_nox.gui.filedialog.asksaveasfilename", return_value=str(excel)
            ):
                app.prepare_button.invoke()
                _wait_for_idle(root, app)
            assert excel.exists()
            capture("02_excel_pendente")
            app.continue_generate_button.invoke()
            capture("03_excel_revisado")
            blocked = output / "nao_deve_existir.txt"
            with patch(
                "gerador_cnab_nox.gui.filedialog.asksaveasfilename", return_value=str(blocked)
            ):
                app.generate_button.invoke()
                _wait_for_idle(root, app)
            assert not blocked.exists() and app.generation_result.kind == "error"
            capture("03_emissao_bloqueada")
            book = load_workbook(excel)
            sheet = book["CREDITOS"]
            h = {c.value: c.column for c in sheet[1]}
            corrections = []
            for row in range(2, sheet.max_row + 1):
                if sheet.cell(row, h["MODALIDADE"]).value == "CESSAO_DA_CESSAO":
                    sheet.cell(row, h["DT_EMISSAO_TITULO"]).value = date(2026, 9, 2)
                    corrections.append(
                        dict(
                            linha=row,
                            campo="DT_EMISSAO_TITULO",
                            valor="2026-09-02",
                            origem="correcao_manual_sintetica",
                        )
                    )
            book.save(excel)
            book.close()
            for source in [*pfmis, analytic, due]:
                source.unlink()
            app.sequence.set("999999")
            app.liquidation.set("INVALIDA")
            _set_commission(root, app, "99")
            txt = output / "resultado.txt"
            with patch("gerador_cnab_nox.gui.filedialog.asksaveasfilename", return_value=str(txt)):
                app.generate_button.invoke()
                _wait_for_idle(root, app)
            assert txt.exists(), app.generation_result.details or app.status.get()
            capture("04_txt_gerado")
            records = txt.read_bytes().split(b"\r\n")[:-1]
            assert len(records) == 5 and all(len(r) == 444 for r in records)
            assert sum(int(r[192:205]) for r in records[1:-1]) == 27050
            assert [int(r[37:62]) for r in records[1:-1]] == [500, 501, 502]
            report = dict(
                snapshots=snapshots,
                correcoes=corrections,
                result_states=result_states,
                scale=scale,
                window=window,
                runtime=dict(
                    python=sys.version.split()[0],
                    tk=str(root.tk.call("package", "require", "Tk")),
                    fontsystem=str(root.tk.call("tk::pkgconfig", "get", "fontsystem")),
                    tk_scaling=float(root.tk.call("tk", "scaling")),
                    body_font=app.fonts["body"].actual(),
                    body_metrics=app.fonts["body"].metrics(),
                ),
                titulos=3,
                nominal_centavos=53000,
                aquisicao_centavos=27050,
                fontes_inacessiveis=True,
            )
            (output / "evidencia_sintetica.json").write_text(
                json.dumps(report, ensure_ascii=False, indent=2)
            )
            print(
                json.dumps(
                    dict(
                        titulos=3,
                        aquisicao_centavos=27050,
                        correcoes_emissao=len(corrections),
                        snapshots=len(snapshots),
                        fontes_inacessiveis=True,
                        widgets=snapshots[-1]["widgets"],
                    )
                )
            )
    finally:
        root.destroy()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--capture-python", type=Path)
    parser.add_argument(
        "--scale", type=float, default=1.0,
        help="Escala lógica do Tk; com Xft, ajuste também o DPI das fontes no processo de QA.",
    )
    parser.add_argument("--window", default="1040x700")
    args = parser.parse_args()
    run(args.output, args.capture_python, args.scale, args.window)
