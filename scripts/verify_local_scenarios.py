"""Durable synthetic fixtures for a local, data-free blocking-scenario audit."""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from gerador_cnab_nox.errors import ValidationError
from gerador_cnab_nox.models import CESSAO, NORMAL, PfmiInput
from gerador_cnab_nox.readers import read_pfmi
from gerador_cnab_nox.service import generate_cnab, prepare_workbook
from gerador_cnab_nox.workbook import read_intermediate
from scripts.local_audit_compare import parse_cnab
from tests.test_gui import _write_sources


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def save(path, value):
    with Path(path).open("x", encoding="utf-8") as handle:
        json.dump(value, handle, ensure_ascii=False, indent=2, default=str)
    Path(path).chmod(0o600)


def require(condition, message):
    if not condition:
        raise RuntimeError(message)


def change(path, row, fields):
    book = load_workbook(path)
    sh = book["CREDITOS"]
    h = {c.value: c.column for c in sh[1]}
    edits = []
    for field, value in fields.items():
        cell = sh.cell(row, h[field])
        edits.append(dict(celula=cell.coordinate, campo=field, antes=cell.value, depois=value))
        cell.value = value
    book.save(path)
    book.close()
    return edits


def group(name="EMPRESA ALFA", failure="FALENCIA ALFA", **overrides):
    return dict(
        cedent=name,
        failure=failure,
        operation=overrides.get("operation", 100),
        nominal=overrides.get("nominal", 150),
        acquisition=overrides.get("acquisition", 100),
        cession=overrides.get("cession", False),
    )


def setup(output, name, groups=None):
    folder = output / name
    folder.mkdir()
    ps, a, b = _write_sources(folder, groups or [group()])
    return folder, ps, a, b


def prepare(folder, ps, a, b, modes=None):
    path = folder / "intermediario.xlsx"
    inputs = [
        PfmiInput(p, mode, "15") for p, mode in zip(ps, modes or [NORMAL] * len(ps), strict=True)
    ]
    prepare_workbook(
        inputs,
        a,
        b,
        liquidation_date="03/09/2026",
        first_sequence=500,
        output_path=path,
        log_directory=folder / "logs",
    )
    return path


def expect_blocked(path, target, token):
    try:
        generate_cnab(path, output_path=target, log_directory=target.parent / "logs")
    except ValidationError as exc:
        require(
            any(token in issue for issue in exc.issues), "Bloqueio não comprovou causa esperada"
        )
        require(not target.exists(), "Bloqueio deixou TXT")
        return dict(bloqueado=True, causa=token, quantidade_pendencias=len(exc.issues))
    raise RuntimeError("TXT indevido em cenário bloqueante")


def synthetic(output):
    output.mkdir(parents=True, exist_ok=False, mode=0o700)
    report = []
    cases = [
        ("centavo", {"VL_PRESENTE": 100.01}, "DIVERGENCIA_VALOR"),
        ("documento", {"DOC_CEDENTE": "123"}, "DOC_CEDENTE"),
        ("emissao_ausente", {"DT_EMISSAO_TITULO": None}, "DT_EMISSAO_TITULO"),
        ("emissao_futura", {"DT_EMISSAO_TITULO": date(2026, 9, 4)}, "posterior"),
        ("controle_alterado", {"TOTAL_PFMI_GRUPO": 101}, "TOTAL_PFMI_GRUPO"),
    ]
    for name, edit, token in cases:
        folder, ps, a, b = setup(output, name)
        path = prepare(folder, ps, a, b)
        edits = change(path, 2, {**edit, "APROVADO": "SIM"})
        result = expect_blocked(path, folder / "nao_deve_existir.txt", token)
        save(folder / "evidencia.json", dict(correcoes=edits, resultado=result))
        report.append(dict(cenario=name, **result))

    folder, ps, a, b = setup(
        output, "centavos_opostos", [group(), group("EMPRESA BETA", "FALENCIA BETA")]
    )
    path = prepare(folder, ps, a, b)
    edits = change(path, 2, {"VL_PRESENTE": 100.01}) + change(path, 3, {"VL_PRESENTE": 99.99})
    result = expect_blocked(path, folder / "nao_deve_existir.txt", "DIVERGENCIA_VALOR")
    save(folder / "evidencia.json", dict(correcoes=edits, resultado=result))
    report.append(dict(cenario="centavos_opostos", **result))

    folder, ps, a, b = setup(output, "alternativas_ambiguas")
    book = load_workbook(a)
    book.active.cell(2, 2).value = "ALTERNATIVA UM"
    duplicate = [c.value for c in book.active[2]]
    duplicate[1], duplicate[6] = "ALTERNATIVA DOIS", "CREDITO-2"
    book.active.append(duplicate)
    book.save(a)
    book.close()
    path = prepare(folder, ps, a, b)
    rows = read_intermediate(path).rows
    require(
        len(rows) == 2 and all(r.values["INCLUIR_CNAB"] == "NAO" for r in rows),
        "Alternativas foram escolhidas automaticamente",
    )
    result = expect_blocked(path, folder / "nao_deve_existir.txt", "nenhum crédito selecionado")
    report.append(dict(cenario="alternativas_ambiguas", candidatos=2, **result))

    folder, ps, a, b = setup(output, "reuso_credito", [group(), group()])
    book = load_workbook(a)
    book.active.delete_rows(3)
    book.save(a)
    book.close()
    path = prepare(folder, ps, a, b)
    edits = change(path, 2, {"INCLUIR_CNAB": "SIM"}) + change(path, 3, {"INCLUIR_CNAB": "SIM"})
    result = expect_blocked(path, folder / "nao_deve_existir.txt", "mais de uma vez")
    save(folder / "evidencia.json", dict(correcoes=edits, resultado=result))
    report.append(dict(cenario="reuso_credito", **result))

    folder, ps, a, b = setup(output, "espelho_parcial")
    book = load_workbook(ps[0])
    sh = book.active
    sh.cell(3, 1).value = 40
    sh.cell(3, 9).value = 40
    other = [c.value for c in sh[3]]
    other[0] = other[8] = 60
    sh.append(other)
    mirror = book.copy_worksheet(sh)
    mirror.cell(4, 1).value = 61
    mirror.cell(4, 9).value = 61
    book.save(ps[0])
    book.close()
    pdata = read_pfmi(ps[0])
    require(
        len(pdata.payments) == 4
        and not any(e["tipo"] == "ESPELHO_INTEGRAL" for e in pdata.evidence),
        "Espelho parcial descartou pagamentos",
    )
    path = prepare(folder, ps, a, b)
    result = expect_blocked(path, folder / "nao_deve_existir.txt", "nenhum crédito selecionado")
    report.append(dict(cenario="espelho_parcial", pagamentos=4, **result))

    folder, ps, a, b = setup(
        output,
        "comissao_por_credito",
        [group(operation=0.04, nominal=0.10, acquisition=0, cession=True)],
    )
    book = load_workbook(a)
    duplicate = [c.value for c in book.active[2]]
    duplicate[6] = "CREDITO-2"
    book.active.append(duplicate)
    book.save(a)
    book.close()
    path = prepare(folder, ps, a, b, [CESSAO])
    loaded = read_intermediate(path)
    require(
        len(loaded.rows) == 2
        and all(str(r.values["COMISSAO_CALCULADA"]) == "0.02" for r in loaded.rows),
        "Arredondamento por crédito incorreto",
    )
    blocked = expect_blocked(path, folder / "sem_emissao.txt", "DT_EMISSAO_TITULO")
    edits = change(path, 2, {"DT_EMISSAO_TITULO": date(2026, 9, 2)}) + change(
        path, 3, {"DT_EMISSAO_TITULO": date(2026, 9, 2)}
    )
    # Make only our synthetic sources inaccessible; no real file is renamed/deleted.
    for p in [*ps, a, b]:
        p.rename(p.with_suffix(p.suffix + ".indisponivel"))
    result = generate_cnab(path, output_path=folder / "valido.txt", log_directory=folder / "logs")
    decoded = parse_cnab(result.output_path.read_bytes())
    require(
        [str(d["VL_PRESENTE"]) for d in decoded["details"]] == ["0.02", "0.02"],
        "Centavos do TXT divergem da expectativa independente",
    )
    save(
        folder / "evidencia.json",
        dict(
            correcoes=edits,
            bloqueio_inicial=blocked,
            aquisicao_centavos=4,
            nominal_centavos=20,
            comissao_agregada_incorreta_centavos=3,
            fontes_inacessiveis=True,
        ),
    )
    report.append(
        dict(
            cenario="comissao_por_credito",
            gerado=True,
            titulos=2,
            aquisicao_centavos=4,
            fontes_inacessiveis=True,
        )
    )

    long_name = "ÁCME, CIA. " + "NOME " * 9
    folder, ps, a, b = setup(output, "nome_longo_pontuacao", [group(name=long_name)])
    path = prepare(folder, ps, a, b)
    require(
        read_intermediate(path).rows[0].values["NOME_CEDENTE"] == long_name.strip(),
        "Excel perdeu nome completo",
    )
    result = generate_cnab(path, output_path=folder / "valido.txt", log_directory=folder / "logs")
    detail = result.output_path.read_bytes().split(b"\r\n")[1]
    expected = ("ACME, CIA. " + "NOME " * 9)[:40].encode("cp1252")
    require(detail[334:374] == expected, "Nome técnico/pontuação diverge")
    require(any("NOME_TRUNCADO" in w for w in result.warnings), "Truncamento sem alerta")
    report.append(
        dict(
            cenario="nome_longo_pontuacao",
            gerado=True,
            nome_completo_preservado=True,
            pontuacao_preservada=True,
        )
    )
    save(output / "relatorio_sintetico.json", report)
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    result = synthetic(args.output)
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
